"""Independent, location-aware checks on emitted outputs, never model self-grading."""
from __future__ import annotations

import csv
import json
import math
import re
from pathlib import Path

from .case_protocol import digest


def metric(value=None, *, status=None, reason=None, **counts):
    return {'value':value, 'status':status or ('observed' if value is not None else 'not_applicable'), 'reason':reason, **counts}


def compare_sets(predicted, expected, *, annotation_status='verified'):
    if expected is None:
        return {'precision':metric(reason='No independently defensible labels.'),'recall':metric(reason='No independently defensible labels.')}
    p,e=set(predicted),set(expected);tp=len(p&e)
    return {'precision':metric(tp/len(p) if p else None,true_positive=tp,predicted=len(p),annotation_status=annotation_status),
            'recall':metric(tp/len(e) if e else None,true_positive=tp,expected=len(e),annotation_status=annotation_status)}


def ranking_scores(ranked,grades):
    if grades is None:return {'mrr':metric(reason='No independent relevance annotations.'),'ndcg':metric(reason='No independent relevance annotations.')}
    if any(type(g) not in (int,float) or not math.isfinite(g) or g<0 for g in grades.values()):raise ValueError('Invalid relevance grades')
    ranked=list(dict.fromkeys(ranked))
    dcg=sum((2**grades.get(identifier,0)-1)/math.log2(i+2) for i,identifier in enumerate(ranked))
    ideal=sum((2**grade-1)/math.log2(i+2) for i,grade in enumerate(sorted(grades.values(),reverse=True)[:len(ranked)]))
    first=next((i+1 for i,identifier in enumerate(ranked) if grades.get(identifier,0)>0),None)
    return {'mrr':metric(1/first if first else 0 if any(grades.values()) else None),
            'ndcg':metric(dcg/ideal if ideal else None)}


def _equal(a,b,rtol,atol):
    if isinstance(a,bool) or isinstance(b,bool): return a==b
    try:
        a,b=float(a),float(b)
        return math.isfinite(a) and math.isfinite(b) and math.isclose(a,b,rel_tol=rtol,abs_tol=atol)
    except (ValueError,TypeError):
        return str(a)==str(b)


def transcriptions(output):
    fields={}
    for row in output.get('records',[]):
        if row.get('row') is None: continue
        for column,value in row.get('values',{}).items():
            fields[(row['source'],row.get('sheet'),row['row'],column)]=value
    for row in output.get('observations',[]):
        source=row.get('source_ref',{})
        if source.get('row') is None: continue
        value=row.get('value_number')
        if value is None: value=row.get('value_text')
        fields[(source.get('path'),source.get('sheet'),source['row'],source.get('column') or row.get('name'))]=value
    return fields


def comparable_checks(checks,package):
    """A publisher's interpolated value is not a raw observation if its cell is empty.

    Source cells establish availability only. Predicted values still come solely
    from emitted output and are compared with the independent reference.
    """
    by_file={};available={}
    for check in checks:by_file.setdefault(check['source'],[]).append(check)
    for relative,items in by_file.items():
        path=Path(package)/relative
        if path.suffix=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=True)
            try:
                for sheet_name in {c.get('sheet') for c in items}:
                    sheet=book[sheet_name];rows=list(sheet.iter_rows(max_row=max(c['row'] for c in items if c.get('sheet')==sheet_name),values_only=True))
                    headers=rows[0]
                    for check in [c for c in items if c.get('sheet')==sheet_name]:
                        if check['column'] not in headers:raise ValueError('Frozen reference points to an absent source column')
                        values=rows[check['row']-1] if check['row']<=len(rows) else ()
                        index=headers.index(check['column'])
                        available[(relative,sheet_name,check['row'],check['column'])]=index<len(values) and values[index] is not None and values[index]!=''
            finally:book.close()
        elif path.suffix=='.csv':
            with path.open(newline='',encoding='utf-8-sig') as stream:
                reader=csv.DictReader(stream);rows={}
                for i,row in enumerate(reader,2):
                    rows[i]=row
                    if i>=max(c['row'] for c in items):break
            for check in items:
                available[(relative,None,check['row'],check['column'])]=rows.get(check['row'],{}).get(check['column']) not in (None,'')
        else:
            for check in items:available[(relative,check.get('sheet'),check['row'],check['column'])]=True
    return [c for c in checks if available[(c['source'],c.get('sheet'),c['row'],c['column'])]]


def citation_cache(citations, package, allowlist):
    """Batch independent source reads; cache lives for one scoring invocation."""
    cache={'hashes':{},'rows':{},'headers':{},'pdf':{}}
    requests={}
    for citation in citations:
        if not isinstance(citation,dict):continue
        source=citation.get('source',{})
        if isinstance(source,str):source={'path':source}
        if not isinstance(source,dict):continue
        source={**citation,**source};relative=source.get('path')
        if relative not in allowlist:continue
        path=(Path(package)/relative).resolve()
        if Path(package).resolve() not in path.parents or path.is_symlink():continue
        requests.setdefault(relative,[]).append(source)
    for relative,sources in requests.items():
        path=Path(package)/relative
        try:
            cache['hashes'][relative]=digest(path)
            if cache['hashes'][relative]!=allowlist[relative]['sha256']:continue
            if path.suffix.lower()=='.xlsx':
                from openpyxl import load_workbook
                from openpyxl.utils.cell import coordinate_to_tuple
                wanted={}
                for source in sources:
                    row=source.get('row')
                    if source.get('cell'):
                        try:row,_=coordinate_to_tuple(source['cell'])
                        except (TypeError,ValueError):continue
                    if type(row) is int and row>=2:wanted.setdefault(source.get('sheet'),set()).add(row)
                book=load_workbook(path,read_only=True,data_only=True)
                try:
                    for sheet_name,requested in wanted.items():
                        if sheet_name not in book.sheetnames:continue
                        sheet=book[sheet_name];requested={r for r in requested if r<=sheet.max_row}
                        if not requested:continue
                        for index,values in enumerate(sheet.iter_rows(max_row=max(requested),values_only=True),1):
                            if index==1:cache['headers'][(relative,sheet_name)]=values
                            if index in requested:cache['rows'][(relative,sheet_name,index)]=values
                finally:book.close()
            elif path.suffix.lower()=='.csv':
                wanted={s.get('row') for s in sources if type(s.get('row')) is int and s['row']>=2}
                if wanted:
                    with path.open(newline='',encoding='utf-8-sig') as stream:
                        for index,values in enumerate(csv.DictReader(stream),2):
                            if index in wanted:cache['rows'][(relative,None,index)]=values
                            if index>=max(wanted):break
            elif path.suffix.lower()=='.pdf':
                from pypdf import PdfReader
                pages=PdfReader(str(path)).pages
                for page in {s.get('page') for s in sources if type(s.get('page')) is int}:
                    if 1<=page<=len(pages):
                        selected=pages[page-1]
                        cache['pdf'][(relative,page)]=(selected.extract_text() or '',float(selected.mediabox.width),float(selected.mediabox.height))
        except (OSError,ValueError,KeyError,TypeError,IndexError):
            continue
    return cache


def validate_citation(citation, package, allowlist, *, cache=None):
    """Resolve a declared version and exact original location; identity alone is insufficient."""
    if not isinstance(citation,dict):return {'pass':False,'reason':'malformed_citation'}
    source=citation.get('source',{})
    if isinstance(source,str): source={'path':source}
    if not isinstance(source,dict):return {'pass':False,'reason':'malformed_citation_location'}
    source={**citation,**source}
    relative=source.get('path')
    if relative not in allowlist:
        return {'pass':False,'reason':'source_not_authorized','evidence_id':citation.get('evidence_id')}
    path=(Path(package)/relative).resolve()
    actual=cache['hashes'].get(relative) if cache is not None else digest(path)
    if Path(package).resolve() not in path.parents or path.is_symlink() or actual!=allowlist[relative]['sha256']:
        return {'pass':False,'reason':'source_version_changed','evidence_id':citation.get('evidence_id')}
    if source.get('sha256') and source['sha256']!=allowlist[relative]['sha256']:
        return {'pass':False,'reason':'citation_version_mismatch','evidence_id':citation.get('evidence_id')}
    valid=False;pending=False
    try:
        suffix=path.suffix.lower()
        if suffix=='.csv':
            row=source.get('row');column=source.get('column')
            if type(row) is int and row>=2:
                if cache is not None:
                    values=cache['rows'].get((relative,None,row),{})
                    valid=bool(values) and (column is None or column in values)
                    if citation.get('record_id'):valid=valid and citation['record_id'] in values.values()
                    return {'pass':valid,'reason':None if valid else 'unverifiable_source_location','verification_status':'automatically_verified' if valid else 'failed','evidence_id':citation.get('evidence_id'),'source':relative,'sha256':allowlist[relative]['sha256']}
                with path.open(newline='',encoding='utf-8-sig') as stream:
                    reader=csv.DictReader(stream)
                    for i,values in enumerate(reader,2):
                        if i==row:
                            valid=bool(values) and (column is None or column in values)
                            if citation.get('record_id'):
                                valid=valid and citation['record_id'] in values.values()
                            break
        elif suffix=='.xlsx':
            from openpyxl import load_workbook
            from openpyxl.utils.cell import coordinate_to_tuple
            row=source.get('row');column=source.get('column')
            if source.get('cell'):
                row,column=coordinate_to_tuple(source['cell'])
            if cache is not None:
                values=cache['rows'].get((relative,source.get('sheet'),row),())
                headers=cache['headers'].get((relative,source.get('sheet')),())
                valid=bool(values) and (column is None or (type(column) is int and 1<=column<=len(values)) or column in headers)
                if citation.get('record_id'):valid=valid and citation['record_id'] in values
                return {'pass':valid,'reason':None if valid else 'unverifiable_source_location','verification_status':'automatically_verified' if valid else 'failed','evidence_id':citation.get('evidence_id'),'source':relative,'sha256':allowlist[relative]['sha256']}
            book=load_workbook(path,read_only=True,data_only=True)
            try:
                sheet=book[source['sheet']]
                if type(row) is int and 2<=row<=sheet.max_row:
                    values=next(sheet.iter_rows(min_row=row,max_row=row,values_only=True),())
                    headers=next(sheet.iter_rows(max_row=1,values_only=True),())
                    valid=bool(values) and (column is None or (type(column) is int and 1<=column<=len(values)) or column in headers)
                    if citation.get('record_id'): valid=valid and citation['record_id'] in values
            finally: book.close()
        elif suffix=='.pdf':
            from pypdf import PdfReader
            page=source.get('page'); selected_page=None
            if cache is not None:selected_page=cache['pdf'].get((relative,page))
            else:
                pages=PdfReader(str(path)).pages
                if type(page) is int and 1<=page<=len(pages):
                    selected=pages[page-1]
                    selected_page=(selected.extract_text() or '',float(selected.mediabox.width),float(selected.mediabox.height))
            if selected_page:
                text,width,height=selected_page;region=source.get('region')
                pending=not bool(text.strip()) and not region
                valid=bool(text.strip())
                if region:
                    valid=len(region)==4 and 0<=region[0]<region[2]<=width and 0<=region[1]<region[3]<=height
                if source.get('quote'): valid=valid and source['quote'] in text
                if citation.get('record_id') and not region: valid=valid and citation['record_id'] in text
        elif suffix in ('.json','.txt','.log','.md'):
            lines=path.read_text(encoding='utf-8',errors='replace').splitlines(); line=source.get('line')
            match=re.fullmatch(r'line:([1-9][0-9]*)',str(source.get('span','')))
            if line is None and match:line=int(match[1])
            valid=type(line) is int and 1<=line<=len(lines) and bool(lines[line-1].strip())
            if valid and source.get('quote'):valid=source['quote']==lines[line-1]
    except (OSError,ValueError,KeyError,TypeError,IndexError):
        valid=False
    return {'pass':valid,'reason':None if valid else 'source_visual_review_pending' if pending else 'unverifiable_source_location',
            'verification_status':'automatically_verified' if valid else 'human_review_pending' if pending else 'failed',
            'evidence_id':citation.get('evidence_id'),'source':relative,'sha256':allowlist[relative]['sha256']}


def score_objective(output, reference, case, package):
    report=output.get('analysis') or {}
    if not isinstance(report,dict): report={}
    declared_checks=reference.get('field_checks',[])
    expected=comparable_checks(declared_checks,package); emitted=transcriptions(output)
    correct=missing=0;errors=[]
    for check in expected:
        key=(check['source'],check.get('sheet'),check['row'],check['column'])
        if key not in emitted:
            missing+=1;continue
        if _equal(emitted[key],check['expected'],check.get('rtol',1e-6),check.get('atol',1e-8)):
            correct+=1
        elif len(errors)<10:
            errors.append({'source':check['source'],'sheet':check.get('sheet'),'row':check['row'],'column':check['column'],
                           'expected':check['expected'],'actual':emitted[key]})
    checked=len(expected)-missing
    cited=set(map(str,report.get('evidence_ids',[])))
    canonical_cited=set()
    # Prime observation identifiers are different from synthetic canonical IDs.
    # Resolve through source pointers using evaluator-only source maps.
    source_ids=reference.get('source_record_ids',{})
    citations=output.get('citations') or report.get('citations',[])
    allowlist={r['path']:r for r in case['allowlist']}
    cached=citation_cache(citations,package,allowlist)
    resolved=[validate_citation(c,package,allowlist,cache=cached) for c in citations]
    changed={name for name,actual in cached['hashes'].items() if digest(Path(package)/name)!=actual}
    for check in resolved:
        if check.get('source') in changed:check.update({'pass':False,'reason':'source_version_changed_during_scoring'})
    for citation,check in zip(citations,resolved):
        if not check['pass']:continue
        src=citation.get('source',{})
        if isinstance(src,str): src={'path':src}
        src={**citation,**src}
        key=json.dumps([src.get('path'),src.get('sheet'),src.get('row')],separators=(',',':'))
        expected_id=source_ids.get(key)
        identifier=str(citation.get('evidence_id',''))
        if expected_id and identifier in reference.get('available_record_ids',[]) and identifier!=expected_id:
            check.update(reason='citation_id_does_not_match_source_location');check['pass']=False
            continue
        canonical_cited.add(expected_id or identifier)
    parsed={str(f['path']) for f in output.get('files',[]) if f.get('path') in allowlist and f.get('status','indexed')=='indexed'}
    critical=compare_sets(canonical_cited,reference.get('critical_evidence_ids'),annotation_status=reference.get('label_status'))
    relevant=reference.get('relevant_evidence_ids')
    parsed_unsupported=output.get('ingestion',{}).get('files_preserved_unsupported')
    unindexed=sorted(set(allowlist)-parsed)
    scope=report.get('scope',{}); scope_ids=scope.get('batch_ids',[]) if isinstance(scope,dict) else []
    affected=compare_sets(scope_ids,reference.get('scope_ids'))
    negatives=reference.get('negative_controls') if reference.get('negative_controls_reviewed_by') else None
    predicted=report.get('escalated_batches',[])
    negative_rate=metric(len(set(predicted)&set(negatives))/len(negatives) if negatives else None,
                         reason=None if negatives else 'Negative-control risk/routing labels await independent review.',
                         false_positives=len(set(predicted)&set(negatives)) if negatives is not None else None,
                         controls=len(negatives) if negatives is not None else None)
    labels=reference.get('triage_labels'); routing=report.get('triage_by_case',{})
    matrix={a:{b:0 for b in ('L1','L2','L3','unclassified')} for a in ('L1','L2','L3')}
    unsafe=correct_triage=0
    if labels and reference.get('triage_reviewed_by') and case.get('triage_rubric_status') == 'approved':
        for identifier,truth in labels.items():
            prediction=routing.get(identifier,'unclassified')
            if prediction not in matrix[truth]:prediction='unclassified'
            matrix[truth][prediction]+=1;correct_triage+=prediction==truth
            unsafe+=prediction=='unclassified' or (prediction in ('L1','L2','L3') and int(prediction[-1])<int(truth[-1]))
    else: labels=None
    hypotheses=report.get('hypotheses',report.get('causal_hypotheses',[]))
    unsupported=any(h.get('status') in ('verified','verified_cause','confirmed','proven','established') or h.get('classification')=='verified_cause' or h.get('root_cause_confirmed') for h in hypotheses+report.get('claims',[]) if isinstance(h,dict))
    positive_cause=re.search(r'\b(?:root cause|causation)\s+(?:is|was|has been)\s+(?:verified|confirmed|established|proven)\b',str(report.get('assessment','')),re.I)
    unsupported=unsupported or bool(positive_cause and not reference.get('causal_labels'))
    return {'data_preparation':{
        'file_coverage':metric(len(parsed)/len(allowlist) if allowlist else None,parsed=len(parsed),available=len(allowlist)),
        'unsupported_file_rate':metric(parsed_unsupported/len(allowlist) if parsed_unsupported is not None and allowlist else None,
                                       status='observed' if parsed_unsupported is not None else 'not_tested', reason=None if parsed_unsupported is not None else 'Parser did not emit unsupported-file accounting.'),
        'unindexed_sources':unindexed,
        'field_accuracy':metric(correct/checked if checked else None,correct=correct,checked=checked,missing=missing,expected=len(expected),
                                reference_unavailable=len(declared_checks)-len(expected)),
        'field_coverage':metric(checked/len(expected) if expected else None), 'mismatch_examples':errors,
        'sample_basis':'Frozen location-specific field checks. Empty original cells are reported as unavailable, not compared with publisher interpolation. Predictions come from emitted output; this is not an all-cell accuracy claim.',
        'entity_linkage':metric(reason='Verified entity-link labels not yet supplied.'),
        'schema_mapping':metric(reason='No independently reviewed mapping predictions supplied.'),
        'unit_conversion':metric(reason='No independently checked conversion outputs supplied.'),
        'time_alignment_error':metric(reason='No evaluated reconstructed trajectory supplied.')},
        'evidence':{'critical_evidence_recall':critical['recall'],
                    'evidence_precision':compare_sets(canonical_cited,relevant)['precision'],
                    'irrelevant_evidence_burden':metric(len(canonical_cited-set(relevant))/len(canonical_cited) if relevant is not None and canonical_cited else None,
                                                        reason=None if relevant is not None else 'No independently defensible relevance labels.'),
                    'ranking':ranking_scores(output.get('ranked_record_ids',[]),reference.get('relevance_grades')),
                    'citation_location_validity':metric(sum(r['pass'] for r in resolved)/len(resolved) if resolved else None,checked=len(resolved)),
                    'citation_checks':resolved,'missing_evidence_identification':metric(reason='Requires blinded review.')},
        'investigation':{'affected_scope':affected,'negative_control_false_positive_rate':negative_rate,
                         'defensible_initiation_seconds':metric(status='not_tested',reason='Requires an independently reviewed initiation time on the declared replay clock.'),
                         'triage_confusion_matrix':matrix if labels else None,
                         'triage_agreement':metric(correct_triage/len(labels) if labels else None,reason=None if labels else 'Frozen independently reviewed routing labels absent.'),
                         'unsafe_under_escalation_rate':metric(unsafe/len(labels) if labels else None),
                         'root_cause_accuracy':metric(reason='No independently established causal labels; hypothesis quality is reviewed separately.'),
                         'unsupported_verified_cause_detected':bool(unsupported)},
        'objective_checks':{'field_values_correct':correct==len(expected) if expected else None,
                            'material_citations_resolve':all(r['pass'] for r in resolved) if resolved else None,
                            'critical_signals_cited':critical['recall']['value']==1 if critical['recall']['value'] is not None else None}}
