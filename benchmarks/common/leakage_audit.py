from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path


FORBIDDEN = re.compile(r"benchmark_ground_truth|public_ground_truth|source[_-]to[_-]record|hidden[_-]label|recall|root[_-]cause|public[_-]lot[_-]mapping", re.I)


def audit_package(package: Path, *, as_of: str | None = None) -> tuple[bool, list[str]]:
    failures: list[str] = []
    package = package.resolve()
    for path in package.rglob("*"):
        if path.is_symlink():
            failures.append(f"symlink in agent package: {path}")
        if path.is_file():
            if FORBIDDEN.search(str(path.relative_to(package))):
                failures.append(f"disclosing filename: {path.name}")
            if path.suffix == ".csv" and as_of:
                import csv
                with path.open(newline="", encoding="utf-8-sig") as stream:
                    for row in csv.DictReader(stream):
                        if row.get("available_at") and row["available_at"] > as_of:
                            failures.append(f"future row in {path.name}")
                            break
                        local_release = row.get("Released local (UTC-07)")
                        if local_release and datetime.fromisoformat(local_release) > datetime.fromisoformat(as_of.replace("Z", "+00:00")):
                            failures.append(f"future locally timed row in {path.name}")
                            break
                        for field, value in row.items():
                            if field.endswith("_at") or field.endswith("_time"):
                                if value and re.match(r"^\d{4}-\d\d-\d\dT", value) and value > as_of:
                                    failures.append(f"future field {field} in {path.name}")
                                    break
            if as_of and path.suffix in {".csv", ".txt", ".json"}:
                body = path.read_text(encoding="utf-8", errors="replace")
                if re.search(r"true_incident|public_lot_mapping|final_outcome|fda warning letter|voluntary nationwide recall", body, re.I):
                    failures.append(f"hidden/public outcome phrase in {path.name}")
            if as_of and path.suffix == ".pdf":
                from pypdf import PdfReader
                body = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
                if re.search(r"true_incident|public_lot_mapping|final_outcome|fda warning letter|voluntary nationwide recall", body, re.I):
                    failures.append(f"hidden/public outcome phrase in {path.name}")
            if as_of and path.suffix == ".xlsx":
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=True)
                for sheet in book:
                    for row in sheet.iter_rows(values_only=True):
                        if any(isinstance(value, str) and re.search(r"true_incident|public_lot_mapping|final_outcome|fda warning letter|voluntary nationwide recall", value, re.I) for value in row):
                            failures.append(f"hidden/public outcome phrase in {path.name}/{sheet.title}")
                            break
                book.close()
    return not failures, failures


def print_audit(package: Path, *, as_of: str | None = None) -> bool:
    passed, failures = audit_package(package, as_of=as_of)
    print("LEAKAGE AUDIT: PASS" if passed else "LEAKAGE AUDIT: FAIL")
    for failure in failures[:20]:
        print(" -", failure)
    return passed
