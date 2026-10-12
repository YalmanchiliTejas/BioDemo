from __future__ import annotations

import bisect
import csv
import json
import math
import time as clock
from collections import defaultdict
from pathlib import Path


ROOT = Path(__file__).resolve().parent


def _number(value: object) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _load_stream(sheet) -> dict:
    iterator = sheet.iter_rows(values_only=True)
    header = tuple(str(value) for value in next(iterator))
    times: list[float] = []
    rows: list[tuple] = []
    backward_steps = 0
    for source_row, values in enumerate(iterator, 2):
        time = _number(values[0] if values else None)
        if time is None:
            continue
        if times and time < times[-1]:
            backward_steps += 1
        times.append(time)
        rows.append(values)
    if backward_steps:
        ordered = sorted(zip(times, rows), key=lambda item: item[0])
        times = [item[0] for item in ordered]
        rows = [item[1] for item in ordered]
    return {"header": header, "times": times, "rows": rows, "sheet": sheet.title,
            "backward_steps_reordered": backward_steps}


def _interpolate(stream: dict, column: int, time: float) -> float | None:
    times = stream["times"]
    if not times or time < times[0]-1e-8 or time > times[-1]+1e-8:
        return None
    index = bisect.bisect_left(times, time)
    if index < len(times) and abs(times[index]-time) < 1e-8:
        return _number(stream["rows"][index][column])
    if index == 0 or index == len(times):
        return None
    before, after = index-1, index
    gap = times[after]-times[before]
    if gap <= 0 or gap > 0.1:  # never bridge a gap greater than six minutes.
        return None
    left = _number(stream["rows"][before][column])
    right = _number(stream["rows"][after][column])
    if left is None or right is None:
        return None
    return left+(right-left)*(time-times[before])/gap


def build(output_root: Path = ROOT / "tasks/normalized") -> dict:
    """Fixed-format baseline: one-minute interpolation from raw streams only."""
    started = clock.monotonic()
    from openpyxl import load_workbook
    output_root.mkdir(parents=True, exist_ok=True)
    manifest = {"version": "cho-normalization-v1", "method": "linear interpolation on a one-minute grid; no gap over 0.1 h bridged",
                "source_role": "original_workbooks_only", "cultivations": []}
    for workbook_path in sorted((ROOT / "raw/workbooks").glob("*.xlsx")):
        book = load_workbook(workbook_path, read_only=True, data_only=True)
        for reactor in ("R01", "R02"):
            streams = [_load_stream(sheet) for sheet in book if sheet.title.startswith(reactor+"_")]
            inline = [stream for stream in streams if "_offline_" not in stream["sheet"]]
            offline = [stream for stream in streams if "_offline_" in stream["sheet"]]
            variables = []
            for stream in inline:
                for index, column in enumerate(stream["header"][1:], 1):
                    if column == "Status":
                        continue
                    variables.append((column, stream, index))
            names = [name for name, _, _ in variables]
            if len(names) != len(set(names)):
                raise ValueError(f"duplicate measurement name in {workbook_path.name}/{reactor}")
            latest = max(stream["times"][-1] for stream in inline if stream["times"])
            prefix = f"{workbook_path.stem}_{reactor}"
            inline_path = output_root / f"{prefix}_inline.csv"
            with inline_path.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream)
                writer.writerow(["Time (h)", *names])
                for minute in range(math.floor(latest*60)+1):
                    time = minute/60
                    writer.writerow([f"{time:.10f}", *[
                        "" if (value := _interpolate(source, index, time)) is None else f"{value:.15g}"
                        for _, source, index in variables]])
            offline_times = sorted({time for source in offline for time in source["times"]})
            offline_variables = [(column, source, index) for source in offline
                                 for index, column in enumerate(source["header"][1:], 1)]
            offline_path = output_root / f"{prefix}_offline.csv"
            with offline_path.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream)
                writer.writerow(["Time (h)", *[column for column, _, _ in offline_variables]])
                for time in offline_times:
                    writer.writerow([f"{time:.10f}", *[
                        "" if (value := _interpolate(source, index, time)) is None else f"{value:.15g}"
                        for _, source, index in offline_variables]])
            manifest["cultivations"].append({"experiment": workbook_path.stem, "reactor": reactor,
                                             "inline_file": inline_path.name, "offline_file": offline_path.name,
                                             "source_workbook": workbook_path.name,
                                             "reordered_backward_steps": {source["sheet"]: source["backward_steps_reordered"]
                                                                          for source in streams if source["backward_steps_reordered"]},
                                             "variables": {name: {"sheet": source["sheet"], "column": index+1}
                                                           for name, source, index in variables+offline_variables}})
        book.close()
    manifest["runtime_seconds"] = round(clock.monotonic() - started, 3)
    (output_root / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest


if __name__ == "__main__":
    print("cultivations", len(build()["cultivations"]))
