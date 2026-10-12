from __future__ import annotations

import csv
import json
import math
import re
import time
from pathlib import Path
from typing import Any

from benchmarks.common.agent_adapter import InvestigationInput


class ConventionalExtraction:
    """Format parser with no autonomous investigation or domain conclusions."""

    def run_investigation(self, request: InvestigationInput) -> dict[str, Any]:
        started = time.monotonic()
        capture_records = request.configuration.get("capture_records", True)
        summarize_numeric = request.configuration.get("summarize_numeric", False)
        records: list[dict[str, Any]] = []
        files: list[dict[str, Any]] = []
        for path in sorted(request.input_package.rglob("*")):
            if not path.is_file() or path.name == "package_manifest.json":
                continue
            relative = str(path.relative_to(request.input_package))
            if path.suffix.lower() == ".csv":
                with path.open(newline="", encoding="utf-8-sig") as stream:
                    reader = csv.DictReader(stream)
                    columns = reader.fieldnames or []
                    count = 0
                    for line, row in enumerate(reader, 2):
                        if not any(value for value in row.values()):
                            continue
                        if capture_records:
                            records.append({"values": row, "source": relative, "row": line,
                                            "record_id": row.get("_evidence_id") or row.get("record_id")})
                        count += 1
                files.append({"path": relative, "type": "csv", "columns": columns, "records": count})
            elif path.suffix.lower() == ".xlsx":
                try:
                    from openpyxl import load_workbook
                except ImportError as exc:
                    raise RuntimeError("Install benchmark dependencies: pip install openpyxl reportlab pypdf") from exc
                workbook = load_workbook(path, read_only=True, data_only=True)
                for sheet in workbook:
                    iterator = sheet.iter_rows(values_only=True)
                    header = next(iterator, ())
                    columns = [str(cell) if cell is not None else f"unnamed_{i}" for i, cell in enumerate(header)]
                    count = 0
                    numeric: dict[str, dict[str, Any]] = {}
                    for line, cells in enumerate(iterator, 2):
                        if not any(cell is not None for cell in cells):
                            continue
                        row = {columns[i]: value for i, value in enumerate(cells) if i < len(columns) and value is not None}
                        if capture_records:
                            records.append({"values": row, "source": relative, "sheet": sheet.title, "row": line,
                                            "record_id": row.get("_evidence_id") or row.get("record_id")})
                        if summarize_numeric:
                            for index, cell in enumerate(cells[:len(columns)]):
                                if not isinstance(cell, (int, float)) or not math.isfinite(float(cell)):
                                    continue
                                name = columns[index]
                                summary = numeric.setdefault(name, {"count": 0, "first": float(cell), "first_row": line,
                                                                 "min": float(cell), "min_row": line,
                                                                 "max": float(cell), "max_row": line,
                                                                 "last": float(cell), "last_row": line, "sum": 0.0})
                                value = float(cell)
                                summary["count"] += 1
                                summary["sum"] += value
                                summary["last"] = value
                                summary["last_row"] = line
                                if value < summary["min"]:
                                    summary["min"], summary["min_row"] = value, line
                                if value > summary["max"]:
                                    summary["max"], summary["max_row"] = value, line
                        count += 1
                    file = {"path": relative, "sheet": sheet.title, "type": "xlsx",
                            "columns": columns, "records": count}
                    if summarize_numeric:
                        file["numeric_summary"] = {name: {**{key: value for key, value in row.items() if key != "sum"},
                                                          "mean": row["sum"]/row["count"]}
                                                   for name, row in numeric.items()}
                    files.append(file)
                workbook.close()
            elif path.suffix.lower() == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                for page_number, page in enumerate(reader.pages, 1):
                    content = page.extract_text() or ""
                    for line_number, line in enumerate(content.splitlines(), 1):
                        if not line.strip():
                            continue
                        records.append({"values": {"text": line}, "source": relative,
                                        "page": page_number, "line": line_number,
                                        "record_id": next(iter(re.findall(r"(?:DEV|CMP)-\d+", line)), None)})
                files.append({"path": relative, "type": "pdf", "pages": len(reader.pages)})
            elif path.suffix.lower() == ".txt":
                for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
                    if line.strip():
                        records.append({"values": {"text": line}, "source": relative, "line": line_number})
                files.append({"path": relative, "type": "txt"})
        return {"mode": "conventional", "objective": request.objective, "as_of": request.as_of,
                "files": files, "records": records, "hypotheses": [], "actions": [],
                "usage": {"runtime_seconds": round(time.monotonic() - started, 3),
                          "tool_calls": len(files), "input_tokens": 0, "output_tokens": 0,
                          "inference_cost_usd": 0, "manual_corrections": 0,
                          "human_clarification_requests": 0, "recursion_depth": 0}}
