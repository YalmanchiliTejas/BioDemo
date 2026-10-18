import copy
import importlib.util
import json
from pathlib import Path
import pytest
from src.workflows.claim_verification import check_claims


def claim():
    return {'claims':[{'claim_id':'C','classification':'observed','statement':'Viability is 38.47 percent.',
        'evidence_ids':['viability'],'contradicting_evidence_ids':[],'counterevidence_search':'No independent replicate supplied.',
        'metric_ids':[],'temporal_basis':'Source sampling coordinate; onset unknown.',
        'uncertainty':'Measurement uncertainty unavailable.','missing_information':'Independent viability measurement.'}]}


def rows():
    return [{'id':'viability','value_number':38.47,'unit':'%','source_ref':{'path':'source.xlsx','cell':'F9'}},
            {'id':'unrelated','value_number':67.0,'unit':'h','source_ref':{'path':'other.xlsx','cell':'A9'}}]


def test_unrelated_citation_cannot_support_claim_number():
    report=claim();report['claims'][0]['statement']='Viability is 67 percent.'
    result=check_claims(report,rows())
    assert any(i['code']=='claim_numeric_literals_unlinked' for i in result['issues'])
    assert check_claims(claim(),rows())['status']=='technical_checks_passed'
    assert check_claims(claim(),rows())['semantic_support']=='independent_review_pending'


def test_substantive_missing_information_arrays_are_valid_but_empty_arrays_are_not():
    report=claim();report['claims'][0]['missing_information']=['Independent measurement','Approved operating strategy']
    assert check_claims(report,rows())['status']=='technical_checks_passed'
    report['claims'][0]['missing_information']=[]
    assert any(i['code']=='missing_information_missing' for i in check_claims(report,rows())['issues'])


@pytest.mark.parametrize('change,code',[
    ({'classification':'verified_cause'},'unsupported_verified_cause'),
    ({'confidence':.99},'uncalibrated_probability'),
    ({'temporal_basis':''},'temporal_basis_missing'),
    ({'classification':'derived'},'derived_calculation_missing'),
    ({'statement':'The measurement proves a failure.'},'causal_language_in_observation'),
    ({'counterevidence_search':''},'counterevidence_search_missing')])
def test_misleading_claims_are_flagged(change,code):
    report=claim();report['claims'][0].update(change)
    assert code in {i['code'] for i in check_claims(report,rows())['issues']}


def test_hypotheses_remain_hypotheses_with_ambiguous_strategy_and_truncated_logs():
    report=claim();report['claims'][0].update(classification='hypothesis',
        statement='Feed delivery may differ; a truncated log does not establish physical interruption.',
        uncertainty='Operating intent and contradictory sensor readings remain unresolved.',
        missing_information='Pump-weight measurements, recipe records and clock validation.')
    result=check_claims(report,rows())
    assert result['status']=='technical_checks_passed'
    assert result['claims'][0]['classification']=='hypothesis'
    assert result['claims'][0]['semantic_support']=='independent_review_pending'


def publication_module():
    spec=importlib.util.spec_from_file_location('publication_builder',Path(__file__).parents[1]/'scripts/build_publication.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    return module


def test_publication_includes_failures_and_unknown_cost_without_mutating_records(tmp_path):
    module=publication_module()
    for identifier,status in [('a','completed'),('b','failed')]:
        p=tmp_path/'results/evaluation'/identifier;p.mkdir(parents=True)
        record={'version':'manufacturing-evaluation-v2','run_id':identifier,'case_id':'CHO-04','benchmark_id':'cho',
            'condition':'raw','mode':'full_rlm','attempted':True,'status':status,
            'case':{'task_spec_sha256':'case'},'usage':{'inference_cost_usd':0},'wall_clock_seconds':100}
        (p/'run.json').write_text(json.dumps(record))
    before=(tmp_path/'results/evaluation/a/run.json').read_bytes()
    result=module.load_results(tmp_path)
    assert result['cohorts'][0]['attempts']==2
    assert result['cohorts'][0]['recorded_execution_statuses']=={'completed':1,'failed':1}
    assert result['runs'][0]['usage']['inference_cost_usd'] is None
    assert result['runs'][0]['usage']['recorded_inference_cost_usd']==0
    assert (tmp_path/'results/evaluation/a/run.json').read_bytes()==before
    assert result['runs'][0]['publication_status']=='evaluation_unavailable'


def test_canonical_citation_overlap_ignores_generated_ids():
    module=publication_module();source={'path':'file.xlsx','sheet':'R1','cell':'A2','row':2,'column':'DO'}
    a={'citations':[{'evidence_id':'first-generated-id','source':source}]}
    b={'citations':[{'evidence_id':'different-generated-id','source':source}]}
    assert module.canonical_citations(a)==module.canonical_citations(b)


def test_lossless_evidence_aliases_reduce_context_without_changing_locations():
    from benchmarks.baselines.investigation import alias_packet,translate_identifiers
    packet={'preparation':{'evidence_ids':['long-evidence-id'],
        'source_versions':[{'source_id':'source-id','path':'source.xlsx'}],
        'screening':{'coverage_units':[{'unit_id':'unit-id','evidence_ids':['long-evidence-id']}],
                    'priority_candidates':[{'candidate_id':'candidate-id','evidence_ids':['long-evidence-id']}]}},
        'source_evidence':[{'id':'long-evidence-id','source_ref':{'source_id':'source-id','path':'source.xlsx','cell':'A2'}}]}
    compressed,aliases=alias_packet(packet)
    assert compressed['source_evidence'][0]['id']=='E0001'
    assert translate_identifiers(compressed,aliases)==packet
    assert compressed['source_evidence'][0]['source_ref']['cell']=='A2'


def test_conventional_investigation_checks_real_sources_without_model(tmp_path):
    from benchmarks.baselines.investigation import run_automated
    from benchmarks.common.agent_adapter import InvestigationInput
    package=tmp_path/'input';package.mkdir()
    (package/'measurements.csv').write_text('batch_id,Time (h),Viability (%)\nB1,0,90\nB1,1,40\n')
    output=run_automated(InvestigationInput(package,'Investigate descriptive changes.',None,
        {'evaluation_time':'2030-01-01T00:00:00Z','state_directory':tmp_path/'state'}))
    assert output['usage']['model_calls']==0
    assert output['analysis']['verification']['claim_verification']['status']=='technical_checks_passed'
    assert output['analysis']['report_status']=='complete_for_available_evidence'
    assert output['analysis']['verification']['causality_verified'] is False


def test_failed_validation_usage_is_preserved(tmp_path):
    from benchmarks.common.usage import recover_usage
    (tmp_path/'single-pass-response.json').write_text(json.dumps({'model':'recorded-model','usage':{'input_tokens':10,'output_tokens':20}}))
    usage=recover_usage(tmp_path)
    assert usage['model_calls']==1 and usage['output_tokens']==20
    assert usage['inference_cost_usd'] is None


def test_investigation_baseline_is_eligible_for_the_same_quality_comparison():
    from benchmarks.common.human_work import aggregate_effort,template
    run={'run_id':'conventional-fixture','mode':'conventional','attempted':True,
         'qualification':{'status':'pending_review','qualified':None}}
    result=aggregate_effort([run],template(['conventional-fixture']))['groups']['conventional']
    assert result['pending_review_cases']==1
    assert result['status']=='pending_measurement'
    assert result['quality_qualified_rate'] is None


def test_blinded_review_packet_preserves_claims_calculations_and_scope():
    from benchmarks.common.case_runner import blinded_packet
    report={'claims':claim()['claims'],'calculated_metrics':[{'metric_id':'M'}],
            'investigation_review':{'coverage':[{'unit_id':'U'}],'candidate_reviews':[{'candidate_id':'P'}]}}
    packet=blinded_packet({'case':{'objective':'Investigate','required_deliverables':[],'source_version':'frozen'},
                          'output':{'analysis':report}})
    assert packet['source_version']=='frozen'
    assert packet['report']['claims']==report['claims']
    assert packet['report']['calculated_metrics']==report['calculated_metrics']
    assert packet['report']['coverage']==[{'unit_id':'U'}]
    assert 'mode' not in packet


def test_child_usage_attributions_are_not_double_counted(tmp_path):
    from src.prime_runtime.telemetry import session_usage
    parent=[{'type':'message','id':'p','message':{'role':'assistant','usage':{'input':30,'output':15}}},
            {'type':'child_usage_attributed','targetId':'p','childUsage':{'input':20,'output':10}}]
    child=[{'type':'message','id':'c','message':{'role':'assistant','usage':{'input':20,'output':10}}}]
    (tmp_path/'parent.jsonl').write_text('\n'.join(map(json.dumps,parent)))
    (tmp_path/'child.jsonl').write_text('\n'.join(map(json.dumps,child)))
    usage=session_usage(tmp_path)
    assert usage['tokens_recorded']['input']==30
    assert usage['tokens_recorded']['output']==15
    assert usage['model_calls_recorded']==2


def test_batched_citation_validation_reads_each_sheet_once_and_matches_reference(tmp_path,monkeypatch):
    import openpyxl
    from benchmarks.common.objective_scores import citation_cache,validate_citation
    from benchmarks.common.case_protocol import digest
    book=openpyxl.Workbook();sheet=book.active;sheet.title='Results'
    sheet.append(['record_id','value'])
    for i in range(200):sheet.append([f'R{i}',i])
    book.save(tmp_path/'source.xlsx');book.close()
    allowed={'source.xlsx':{'sha256':digest(tmp_path/'source.xlsx')}}
    citations=[{'source':{'path':'source.xlsx','sheet':'Results','row':r,'column':'value'}} for r in (2,50,100,201,202)]
    expected=[validate_citation(c,tmp_path,allowed) for c in citations]
    actual_load=openpyxl.load_workbook;loads=[]
    def tracked(*args,**kwargs):loads.append(args[0]);return actual_load(*args,**kwargs)
    monkeypatch.setattr(openpyxl,'load_workbook',tracked)
    cache=citation_cache(citations,tmp_path,allowed)
    actual=[validate_citation(c,tmp_path,allowed,cache=cache) for c in citations]
    assert actual==expected
    assert len(loads)==1
    (tmp_path/'source.xlsx').write_bytes(b'changed')
    fresh=citation_cache(citations,tmp_path,allowed)
    assert not validate_citation(citations[0],tmp_path,allowed,cache=fresh)['pass']
