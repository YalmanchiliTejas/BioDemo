from __future__ import annotations

from io import BytesIO

from .domain import DocumentSection


class OpenDocumentExtractor:
    """Open-source extraction for PDF, DOCX, text, and scanned image inputs."""

    def extract(self, content: bytes, media_type: str) -> tuple[str, tuple[DocumentSection, ...]]:
        if media_type == "application/pdf":
            return self._pdf(content)
        if media_type in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        }:
            return self._docx(content)
        if media_type.startswith("image/"):
            return self._image(content)
        if media_type.startswith("text/") or media_type in {"application/json", "application/xml"}:
            text = content.decode("utf-8")
            return text, (DocumentSection("body", "Body", text),)
        raise ValueError(f"unsupported document media type: {media_type}")

    def _pdf(self, content: bytes) -> tuple[str, tuple[DocumentSection, ...]]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("Install the 'documents' extra to extract PDF files") from exc
        reader = PdfReader(BytesIO(content))
        sections = tuple(
            DocumentSection(f"page-{number}", f"Page {number}", page.extract_text() or "", number)
            for number, page in enumerate(reader.pages, start=1)
        )
        if any(not section.text.strip() for section in sections):
            sections = self._ocr_pdf(content, sections)
        return "\n\n".join(section.text for section in sections), sections

    def _ocr_pdf(
        self, content: bytes, extracted: tuple[DocumentSection, ...]
    ) -> tuple[DocumentSection, ...]:
        try:
            import pypdfium2
            import pytesseract
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("Install the 'documents' extra to OCR scanned PDF pages") from exc
        pdf = pypdfium2.PdfDocument(content)
        return tuple(
            section if section.text.strip() else DocumentSection(
                section.section_id, section.title,
                pytesseract.image_to_string(pdf[index].render(scale=2).to_pil()), section.page,
            )
            for index, section in enumerate(extracted)
        )

    def _docx(self, content: bytes) -> tuple[str, tuple[DocumentSection, ...]]:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("Install the 'documents' extra to extract DOCX files") from exc
        paragraphs = [paragraph.text for paragraph in Document(BytesIO(content)).paragraphs if paragraph.text.strip()]
        text = "\n".join(paragraphs)
        return text, (DocumentSection("body", "Body", text),)

    def _image(self, content: bytes) -> tuple[str, tuple[DocumentSection, ...]]:
        try:
            import pytesseract
            from PIL import Image
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("Install the 'documents' extra to OCR scanned images") from exc
        text = pytesseract.image_to_string(Image.open(BytesIO(content)))
        return text, (DocumentSection("scan", "Scanned document", text, 1),)
