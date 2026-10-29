from __future__ import annotations

import hashlib
import json
import math
from itertools import zip_longest
from pathlib import Path


from benchmarks.common.artifact_store import uses_assets

ROOT = Path(__file__).resolve().parent


@uses_assets(ROOT / "raw", ROOT / "reference/aligned")
def validate() -> dict:
    """Verify that the labeled source-format perturbation preserves observations."""
    from openpyxl import load_workbook

    manifest = json.loads((ROOT / "reference/perturbation_manifest.json").read_text(encoding="utf-8"))
    original = ROOT / "raw/workbooks" / manifest["source"]
    variant = ROOT / "tasks/perturbed/EXP001_batch_perturbed.xlsx"
    if hashlib.sha256(original.read_bytes()).hexdigest() != manifest["source_sha256"]:
        raise ValueError("original workbook checksum differs from perturbation manifest")
    if hashlib.sha256(variant.read_bytes()).hexdigest() != manifest["output_sha256"]:
        raise ValueError("perturbed workbook checksum differs from perturbation manifest")
    base = load_workbook(original, read_only=True, data_only=True)
    changed = load_workbook(variant, read_only=True, data_only=True)
    expected_titles = set(base.sheetnames) | {"Instrument service notes"}
    if set(changed.sheetnames) != expected_titles:
        raise ValueError("worksheet set changed unexpectedly")
    if changed.sheetnames[0] != "Instrument service notes":
        raise ValueError("irrelevant worksheet was not inserted at the expected position")
    expected_duplicate = next(item for item in manifest["changes"] if item["type"] == "duplicate_record")
    checked = 0
    intentional = {"renamed_header": 0, "unit_conversions": 0, "timestamp_format": 0,
                   "missing_optional_status": 0, "duplicate_record": 0}
    failures = []
    for title in base.sheetnames:
        source_rows = base[title].iter_rows(values_only=True)
        variant_rows = changed[title].iter_rows(values_only=True)
        for row_number, (source_row, variant_row) in enumerate(zip_longest(source_rows, variant_rows), 1):
            if source_row is None:
                if title == expected_duplicate["sheet"] and row_number == expected_duplicate["new_row"]:
                    expected = list(base[title].iter_rows(min_row=expected_duplicate["source_row"],
                                                           max_row=expected_duplicate["source_row"], values_only=True))[0]
                    if tuple(variant_row) == tuple(expected):
                        intentional["duplicate_record"] += 1
                        continue
                failures.append(f"unexpected added row: {title}:{row_number}")
                continue
            if variant_row is None:
                failures.append(f"missing row: {title}:{row_number}")
                continue
            for column, source_value in enumerate(source_row, 1):
                target = variant_row[column-1] if column <= len(variant_row) else None
                checked += 1
                if title == "R01_inline_ProcessParameters" and row_number == 1 and column == 4:
                    ok = target == "Temp [°C]"
                    intentional["renamed_header"] += int(ok)
                elif title == "R02_inline_ProcessParameters" and column == 4:
                    if row_number == 1:
                        ok = target == "Temperature (°F)"
                    elif isinstance(source_value, (int, float)):
                        ok = isinstance(target, (int, float)) and math.isclose(
                            float(target), float(source_value)*9/5+32, rel_tol=1e-12, abs_tol=1e-10)
                        intentional["unit_conversions"] += int(ok)
                    else:
                        ok = target == source_value
                elif title == "R01_online_Feed" and column == 1 and row_number > 1 and isinstance(source_value, (int, float)):
                    ok = isinstance(target, str) and math.isclose(float(target), float(source_value),
                                                                   rel_tol=1e-12, abs_tol=1e-12)
                    intentional["timestamp_format"] += int(ok)
                elif title == "R01_inline_Turbidity" and row_number == 10 and column == 4:
                    ok = target is None
                    intentional["missing_optional_status"] += int(ok)
                else:
                    ok = target == source_value
                if not ok and len(failures) < 20:
                    failures.append(f"unexpected cell change: {title}:{row_number}:{column}")
    base.close()
    changed.close()
    result = {"pass": not failures and all(intentional.values()), "cells_checked": checked,
              "intentional_change_counts": intentional, "unexpected_changes": failures,
              "source_sha256": manifest["source_sha256"], "variant_sha256": manifest["output_sha256"]}
    destination = ROOT / "reports/perturbation_validation.json"
    destination.write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


if __name__ == "__main__":
    outcome = validate()
    print(json.dumps({key: outcome[key] for key in ("pass", "cells_checked", "intentional_change_counts", "unexpected_changes")}, indent=2))
    raise SystemExit(0 if outcome["pass"] else 1)
