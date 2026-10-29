from __future__ import annotations

import csv
import json
import re
import math
from pathlib import Path


from benchmarks.common.artifact_store import uses_assets

ROOT = Path(__file__).resolve().parent


@uses_assets(ROOT / "raw", ROOT / "reference/aligned")
def build_reference() -> dict:
    """Structural map inferred from exact source headers; no biological labels."""
    from openpyxl import load_workbook
    mappings = []
    for workbook_path in sorted((ROOT / "raw/workbooks").glob("*.xlsx")):
        book = load_workbook(workbook_path, read_only=True, data_only=True)
        name = workbook_path.stem
        for reactor in ("R01", "R02"):
            sheet_columns = {}
            for sheet in book:
                if not sheet.title.startswith(reactor + "_"):
                    continue
                first = next(sheet.iter_rows(max_row=1, values_only=True), ())
                for index, value in enumerate(first, 1):
                    if value and value != "Time (h)":
                        sheet_columns.setdefault(str(value), []).append({"sheet": sheet.title, "column": index})
            for kind in ("inline", "offline"):
                csv_path = ROOT / "reference/aligned" / f"{name}_{reactor}_{kind}.csv"
                if not csv_path.exists():
                    continue
                with csv_path.open(newline="", encoding="utf-8-sig") as stream:
                    columns = next(csv.reader(stream), [])
                for column_index, column in enumerate(columns, 1):
                    if column == "Time (h)":
                        continue
                    candidates = sheet_columns.get(column, [])
                    mappings.append({"aligned_csv": csv_path.name, "aligned_column": column,
                                     "aligned_column_index": column_index,
                                     "workbook": workbook_path.name,
                                     "source": candidates[0] if len(candidates) == 1 else None,
                                     "candidate_count": len(candidates)})
        book.close()
    result = {"method": "exact header matching within workbook and reactor", "mappings": mappings,
              "unique_matches": sum(item["source"] is not None for item in mappings),
              "total_columns": len(mappings)}
    destination = ROOT / "reference/structural_map.json"
    destination.write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


@uses_assets(ROOT / "reference/aligned")
def build_objective_summaries() -> dict:
    """Numeric descriptive references only; these do not encode causal relevance."""
    cultivations = {}
    for path in sorted((ROOT / "reference/aligned").glob("*.csv")):
        parts = re.match(r"(EXP\d+)_(batch|fed-batch)_(R\d+)_(inline|offline)\.csv", path.name)
        if not parts:
            raise ValueError(f"unrecognized aligned filename: {path.name}")
        experiment, mode, reactor, kind = parts.groups()
        key = f"{experiment}_{reactor}"
        cultivation = cultivations.setdefault(key, {"experiment": experiment, "reactor": reactor,
                                                    "process_mode": mode, "sources": {}, "measurements": {}})
        stats = {}
        count_rows = 0
        with path.open(newline="", encoding="utf-8-sig") as stream:
            for row in csv.DictReader(stream):
                count_rows += 1
                for column, value in row.items():
                    if column == "Time (h)" or value in (None, ""):
                        continue
                    try:
                        number = float(value)
                    except ValueError:
                        continue
                    if not math.isfinite(number):
                        continue
                    item = stats.setdefault(column, {"count": 0, "sum": 0.0, "min": number, "max": number,
                                                     "first": number, "last": number})
                    item["count"] += 1
                    item["sum"] += number
                    item["min"] = min(item["min"], number)
                    item["max"] = max(item["max"], number)
                    item["last"] = number
        cultivation["sources"][kind] = {"file": path.name, "rows": count_rows}
        for column, item in stats.items():
            cultivation["measurements"][column] = {"kind": kind, "count": item["count"],
                                                     "min": item["min"], "max": item["max"],
                                                     "mean": item["sum"]/item["count"],
                                                     "first": item["first"], "last": item["last"]}
    output = {"method": "descriptive numeric summaries of publication-aligned CSVs",
              "ground_truth_type": "alignment_reference_not_biological_causality",
              "cultivations": cultivations}
    (ROOT / "reference/objective_summaries.json").write_text(json.dumps(output, indent=2), encoding="utf-8")
    return output


if __name__ == "__main__":
    result = build_reference()
    print(result["unique_matches"], result["total_columns"])
    objective = build_objective_summaries()
    print("cultivations", len(objective["cultivations"]))
