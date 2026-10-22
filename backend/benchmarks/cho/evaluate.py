from __future__ import annotations

from pathlib import Path

import csv
import json
import math


def evaluate(output: dict, reference_root: Path, *, rtol: float = 1e-6,
             atol: float = 1e-8) -> dict:
    """Score structural coverage and independently compare inoculation-time values."""
    workbook_files = [file for file in output["files"] if file["type"] == "xlsx"]
    source_files = {file["path"] for file in workbook_files}
    root = reference_root.parents[1]
    mapping = json.loads((root / "reference/structural_map.json").read_text(encoding="utf-8"))["mappings"]
    from openpyxl import load_workbook
    raw_first = {}
    for workbook in sorted(source_files):
        book = load_workbook(root / "raw/workbooks" / workbook, read_only=True, data_only=True)
        for sheet in book:
            first_rows = list(sheet.iter_rows(max_row=2, values_only=True))
            if len(first_rows) == 2:
                raw_first[(workbook, sheet.title)] = first_rows[1]
        book.close()
    aligned_first = {}
    for path in reference_root.glob("*.csv"):
        with path.open(newline="", encoding="utf-8-sig") as stream:
            first = next(csv.reader(stream), [])
            second = next(csv.reader(stream), [])
            aligned_first[path.name] = second
    checked = matched = missing = 0
    examples = []
    for item in mapping:
        source = item["source"]
        if source is None:
            missing += 1; continue
        raw = raw_first.get((item["workbook"], source["sheet"]), ())
        aligned = aligned_first.get(item["aligned_csv"], ())
        ri = source["column"] - 1
        ci = item["aligned_column_index"] - 1
        if ri >= len(raw) or ci >= len(aligned):
            missing += 1; continue
        a, b = raw[ri], aligned[ci]
        if a is None or b == "":
            missing += 1; continue
        checked += 1
        try:
            equal = math.isclose(float(a), float(b), rel_tol=rtol, abs_tol=atol)
        except (TypeError, ValueError):
            equal = str(a) == str(b)
        if equal:
            matched += 1
        elif len(examples) < 5:
            examples.append({"csv": item["aligned_csv"], "column": item["aligned_column"], "raw": str(a), "aligned": b})
    return {"workbook_files_parsed": len(source_files), "worksheets_parsed": len(workbook_files),
            "source_rows_extracted": sum(file["records"] for file in workbook_files),
            "aligned_reference_files": len(list(reference_root.glob("*.csv"))),
            "unique_structural_column_mappings": sum(item["source"] is not None for item in mapping),
            "total_aligned_columns_excluding_time": len(mapping),
            "time_zero_values_checked": checked, "time_zero_matches": matched,
            "time_zero_match_fraction": matched/checked if checked else None,
            "time_zero_unavailable_values": missing, "time_zero_mismatch_examples": examples,
            "full_trajectory_reconstruction_scored": False,
            "reason": "Published alignment can transform time and measurements; only directly comparable inoculation-time values are checked."}
