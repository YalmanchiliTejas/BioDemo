"""Repackage the public CHO cultivation dataset as fictional CMO deliveries.

Source: Zenodo record 20829178 (DOI 10.5281/zenodo.20829178), CC BY 4.0. The
measurements are real laboratory data. The CMO names, file layouts, sample IDs
and certificates are invented for demonstration; injected defects are listed in
each package's evaluation key and are not source measurements.

Packages (one per condition, same structure as the other generated families):

* calibration  EXP001-003  "September" delivery from Northfield Biologics
* returning    EXP004-006  same layouts, new runs
* changed      EXP007-009  renamed/reordered/added columns, one CoA missing,
                           one conflicting duplicate lab row (injected)
* unseen       EXP010-012  fed-batch runs from a second CMO in a different form

Usage: python -m benchmarks.cmo_ingestion.cho_demo [--archive PATH] [--output DIR]
Keys live in ``<package>/evaluation/`` beside ``raw/`` and are never read by
the pipeline.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .generate import write_pdf

ROOT = Path(__file__).resolve().parent
SOURCE_SHA256 = '46a6b92b77883d56df3be015895d93c85b2b2547121af0b8bc0a92f3fbefd17b'
DOI = '10.5281/zenodo.20829178'
DISCLOSURE = ('Real CHO cultivation measurements from Zenodo record 20829178 (CC BY 4.0), repackaged as fictional '
              'CMO deliveries. Company names, layouts, sample IDs and certificates are invented. Injected defects are artificial.')
PROCESS = [('pH', 'pH', 3), ('Dissolved O2 (% air sat)', 'do_saturation', 2), ('Temperature (°C)', 'temperature', 2),
           ('Agitation rate (rpm)', 'agitation', 1), ('Aeration rate total (sL/h)', 'aeration_standard', 2),
           ('Liquid volume (L)', 'volume', 3), ('Cumulative feed weight (g)', 'feed_added', 1)]
OFFLINE = [('VCD Mean (cells/mL)', 'viable_cell_density', 0), ('Viability Mean (%)', 'viability', 2),
           ('Avg. diameter Mean (µm)', 'cell_diameter', 2), ('Glucose (g/L)', 'glucose', 3), ('Lactate (g/L)', 'lactate', 3),
           ('Ammonium (mM)', 'ammonium', 2), ('Osmolality (mOsmol/kg)', 'osmolality', 0)]
UNITS = {'pH': 'pH', 'do_saturation': '% sat', 'temperature': 'K', 'agitation': 'rpm', 'aeration_standard': 'sL/h', 'volume': 'L',
         'feed_added': 'g', 'viable_cell_density': 'cells/mL', 'viability': '%', 'cell_diameter': 'µm', 'glucose': 'g/L',
         'lactate': 'g/L', 'ammonium': 'mM', 'osmolality': 'mOsm/kg'}
STEP_HOURS = 6


def text(value, digits):
    if value is None: return ''
    number = round(float(value), digits)
    return str(int(number)) if digits == 0 else f'{number:.{digits}f}'


def canonical(parameter, written):
    value = float(written)
    return value + 273.15 if parameter == 'temperature' else value


def read_experiment(archive, number):
    name = next(n for n in archive.namelist() if Path(n).name.startswith(f'EXP{number:03d}_') and n.endswith('.xlsx'))
    book = load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)
    runs = {}
    try:
        for reactor in ['R01', 'R02']:
            def rows(sheet):
                data = list(book[sheet].iter_rows(values_only=True))
                return data[0], data[1:]
            head, body = rows(f'{reactor}_inline_ProcessParameters')
            _, feed = rows(f'{reactor}_online_Feed')
            # Snapshot the first sample at or after every 6-hour mark, in file order.
            process, mark = [], 0.0
            for index, row in enumerate(body, 2):
                if row[0] is None or float(row[0]) < mark: continue
                process.append({'time': float(row[0]), 'source_row': index, **dict(zip(head, row))})
                mark = (int(float(row[0]) // STEP_HOURS) + 1) * STEP_HOURS
            feed_times = [(float(r[0]), r[1]) for r in feed if r[0] is not None]
            for snapshot in process:
                # Cumulative feed at the same instant: last reading not after it.
                prior = [w for t, w in feed_times if t <= snapshot['time']]
                snapshot['Cumulative feed weight (g)'] = prior[-1] if prior else None
            cells_head, cells = rows(f'{reactor}_offline_CellAnalysis')
            medium_head, medium = rows(f'{reactor}_offline_MediumAnalysis')
            offline = {}
            for table, header in [(cells, cells_head), (medium, medium_head)]:
                for row in table:
                    if row[0] is None: continue
                    offline.setdefault(round(float(row[0]), 2), {}).update(dict(zip(header, row)))
            runs[f'CHO-{number:03d}-{reactor}'] = {'process': process, 'offline': [dict(time=t, **v) for t, v in sorted(offline.items())],
                                                   'mode': 'fed-batch' if 'fed-batch' in name else 'batch', 'source': Path(name).name}
    finally:
        book.close()
    return runs


class Package:
    def __init__(self, root, ident):
        self.root = root / ident; self.raw = self.root / 'raw'; self.raw.mkdir(parents=True, exist_ok=False)
        (self.root / 'evaluation').mkdir()
        self.fields, self.defects, self.documents, self.transformations = [], [], {}, []

    def field(self, batch, time, parameter, written, source, origin=None, present=True):
        self.fields.append({'batch': batch, 'time': time, 'parameter': parameter, 'present': present and written != '',
                            'value': canonical(parameter, written) if written != '' else None, 'unit': UNITS[parameter],
                            'source': {'path': source['path'], 'sheet': source.get('sheet'), 'row': source.get('row'),
                                       'column': source.get('column'), 'cell': None, 'page': source.get('page')},
                            'original': origin or {}})

    def finish(self, item):
        files = [{'path': p.name, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(self.raw.iterdir())]
        key = {'package_id': item['id'], 'dataset': 'cho', 'condition': item['condition'], 'disclosure': DISCLOSURE,
               'fields': self.fields, 'defects': self.defects, 'documents': self.documents, 'relationships': [],
               'files': files, 'transformations': self.transformations, 'source': {'doi': DOI, 'sha256': SOURCE_SHA256}}
        (self.root / 'evaluation' / 'answer_key.json').write_text(json.dumps(key, indent=1, ensure_ascii=False) + '\n')
        return {**item, 'files': [{'path': f['path']} for f in files]}


def process_workbook(package, runs, path, columns, *, batch_header='Batch', notes=True, sheet='Process data'):
    book = Workbook(); ws = book.active; ws.title = sheet
    headers = [batch_header, 'Time (h)', *[c for c, _p, _d in columns]]
    ws.append(headers)
    row = 2
    for batch, run in runs.items():
        for snap in run['process']:
            time = f"{snap['time']:.2f}"
            out = [batch, time]
            for header, parameter, digits in columns:
                source = next(c for c, p, _d in PROCESS if p == parameter)
                written = text(snap.get(source), digits)
                out.append(float(written) if written != '' else None)
                package.field(batch, time, parameter, written, {'path': path, 'sheet': sheet, 'row': row, 'column': header},
                              {'workbook': run['source'], 'sheet': batch[-3:] + ('_online_Feed' if parameter == 'feed_added' else '_inline_ProcessParameters'),
                               'row': snap['source_row'], 'value': snap.get(source)})
            ws.append(out); row += 1
    if notes:
        info = book.create_sheet('Notes'); info.append(['Comment'])
        info.append(['Six-hour snapshots exported from the site historian.'])
        info.append(['Repackaged public data; see the delivery disclosure.'])
    book.save(package.raw / path)
    package.documents[path] = 'instrument_export'


def lims_csv(package, runs, path, *, batch_header='Batch', extra=(), duplicate=None):
    headers = [batch_header, 'Sample ID', 'Time (h)', *[c for c, _p, _d in OFFLINE], *[e for e, _v in extra]]
    rows = []
    for batch, run in runs.items():
        for index, sample in enumerate(run['offline'], 1):
            rows.append((batch, index, sample, f'S-{batch[4:]}-{index:02d}'))
    with (package.raw / path).open('w', newline='', encoding='utf-8') as stream:
        writer = csv.writer(stream); writer.writerow(headers)
        line = 2
        for batch, index, sample, sample_id in rows:
            time = f"{sample['time']:.2f}"
            values = []
            for header, parameter, digits in OFFLINE:
                written = text(sample.get(header), digits)
                values.append(written)
                package.field(batch, time, parameter, written, {'path': path, 'row': line, 'column': header},
                              {'sheet': batch[-3:] + '_offline', 'time': sample['time'], 'value': sample.get(header)})
            writer.writerow([batch, sample_id, time, *values, *[v for _e, v in extra]]); line += 1
            if duplicate and (batch, index) == duplicate[:2]:
                # Injected: the same sample re-reported with a different glucose result.
                changed = list(values); position = [p for _h, p, _d in OFFLINE].index('glucose')
                changed[position] = text(float(values[position]) + duplicate[2], 3)
                writer.writerow([batch, sample_id, time, *changed, *[v for _e, v in extra]])
                package.defects.append({'rule_code': 'DUPLICATE_CONFLICT', 'path': path, 'parameter': 'glucose', 'batch': batch, 'time': time,
                                        'injected': True})
                package.transformations.append(f'Injected conflicting duplicate glucose for {batch} at {time} h (row {line}).')
                line += 1
    package.documents[path] = 'lab_results'


def certificates(package, runs, company, *, skip=()):
    for batch, run in runs.items():
        if batch in skip:
            package.defects.append({'rule_code': 'BATCH_DOCUMENT_MISSING', 'path': None, 'batch': batch, 'doc_type': 'coa', 'injected': True})
            package.transformations.append(f'Withheld the certificate of analysis for {batch}.')
            continue
        # Report at harvest, taken as the peak viable-cell-density sample.
        final = max((s for s in run['offline'] if s.get('Viability Mean (%)') is not None and s.get('VCD Mean (cells/mL)') is not None),
                    key=lambda s: float(s['VCD Mean (cells/mL)']))
        time = f"{final['time']:.2f}"
        viability = text(final['Viability Mean (%)'], 2); vcd = text(final['VCD Mean (cells/mL)'], 0)
        path = f'CoA_{batch}.pdf'
        write_pdf(package.raw / path, [f'{company} - Certificate of Analysis', 'Fictional company for demonstration; not a release document.',
                                       f'Source data: Zenodo {DOI} (CC BY 4.0)',
                                       f'Batch: {batch} | Hour: {time} | viability [%]: {viability}',
                                       f'Batch: {batch} | Hour: {time} | VCD [cells/mL]: {vcd}', 'Page 1 of 1'])
        package.field(batch, time, 'viability', viability, {'path': path, 'page': 1})
        package.field(batch, time, 'viable_cell_density', vcd, {'path': path, 'page': 1})
        package.documents[path] = 'coa'


def summary_export(package, runs, path):
    """Second CMO: one fed-batch export joining inline values at each lab sample."""
    headers = ['Lot', 'EFT (h)', 'pH', 'pO2 (% air sat)', 'Temp (°C)', 'Feed added (g)', 'VCD (cells/mL)', 'Viability (%)', 'Glucose (g/L)', 'Lactate (g/L)']
    plan = [('pH', 'pH', 'pH', 3), ('pO2 (% air sat)', 'do_saturation', 'Dissolved O2 (% air sat)', 2), ('Temp (°C)', 'temperature', 'Temperature (°C)', 2),
            ('Feed added (g)', 'feed_added', 'Cumulative feed weight (g)', 1), ('VCD (cells/mL)', 'viable_cell_density', 'VCD Mean (cells/mL)', 0),
            ('Viability (%)', 'viability', 'Viability Mean (%)', 2), ('Glucose (g/L)', 'glucose', 'Glucose (g/L)', 3), ('Lactate (g/L)', 'lactate', 'Lactate (g/L)', 3)]
    with (package.raw / path).open('w', newline='', encoding='utf-8') as stream:
        writer = csv.writer(stream); writer.writerow(headers); line = 2
        for batch, run in runs.items():
            for sample in run['offline']:
                inline = next((s for s in run['process'] if s['time'] >= sample['time']), run['process'][-1])
                merged = {**inline, **sample}
                time = f"{sample['time']:.2f}"; values = []
                for header, parameter, source, digits in plan:
                    written = text(merged.get(source), digits); values.append(written)
                    package.field(batch, time, parameter, written, {'path': path, 'row': line, 'column': header})
                writer.writerow([batch, time, *values]); line += 1
    package.documents[path] = 'instrument_export'
    package.transformations.append('Inline values joined to each offline sample at the first inline sample at or after it.')


def build(archive_path, output):
    digest = hashlib.sha256(Path(archive_path).read_bytes()).hexdigest()
    if digest != SOURCE_SHA256: raise ValueError(f'CHO archive SHA-256 mismatch: {digest}')
    output = Path(output)
    with zipfile.ZipFile(archive_path) as archive:
        experiments = {n: read_experiment(archive, n) for n in range(1, 13)}
    def runs(*numbers): return {b: r for n in numbers for b, r in experiments[n].items()}
    company = 'Northfield Biologics'; items = []
    base_columns = [(h, p, d) for h, p, d in PROCESS if p != 'feed_added']

    for condition, numbers, month in [('calibration', (1, 2, 3), 'September'), ('returning', (4, 5, 6), 'October')]:
        package = Package(output, f'cho-{condition}')
        selected = runs(*numbers)
        process_workbook(package, selected, f'Northfield_{month}_process_trends.xlsx', base_columns)
        lims_csv(package, selected, f'LIMS_offline_results_{month}.csv')
        certificates(package, selected, company)
        items.append(package.finish({'id': f'cho-{condition}', 'dataset': 'cho', 'condition': condition, 'cmo': 'cho-northfield',
                                     'name': f'{company} · {month} runs', 'synthetic': False, 'public_data': True, 'disclosure': DISCLOSURE}))

    package = Package(output, 'cho-changed'); selected = runs(7, 8, 9)
    renamed = {'Temperature (°C)': 'Temp (°C)', 'Dissolved O2 (% air sat)': 'DO (% air sat)'}
    order = ['Temperature (°C)', 'pH', 'Dissolved O2 (% air sat)', 'Agitation rate (rpm)', 'Liquid volume (L)', 'Aeration rate total (sL/h)', 'Cumulative feed weight (g)']
    columns = [(renamed.get(h, h), p, d) for h in order for hh, p, d in PROCESS if hh == h]
    process_workbook(package, selected, 'Northfield_November_process_trends.xlsx', columns)
    package.defects.append({'rule_code': 'LAYOUT_CHANGED', 'path': 'Northfield_November_process_trends.xlsx', 'injected': True})
    package.transformations.append('Process export: renamed Temperature and Dissolved O2, reordered columns, added cumulative feed.')
    batches = list(selected)
    lims_csv(package, selected, 'LIMS_offline_results_November.csv', batch_header='Lot', extra=[('Analyst', 'QC-2')], duplicate=(batches[1], 3, 0.25))
    package.defects.append({'rule_code': 'LAYOUT_CHANGED', 'path': 'LIMS_offline_results_November.csv', 'injected': True})
    certificates(package, selected, company, skip={batches[4]})
    items.append(package.finish({'id': 'cho-changed', 'dataset': 'cho', 'condition': 'changed', 'cmo': 'cho-northfield',
                                 'name': f'{company} · November runs', 'synthetic': False, 'public_data': True, 'disclosure': DISCLOSURE}))

    package = Package(output, 'cho-unseen'); selected = runs(10, 11, 12)
    summary_export(package, selected, 'harborview_fedbatch_summary.csv')
    items.append(package.finish({'id': 'cho-unseen', 'dataset': 'cho', 'condition': 'unseen', 'cmo': 'cho-harborview',
                                 'name': 'Harborview CDMO · fed-batch runs', 'synthetic': False, 'public_data': True, 'disclosure': DISCLOSURE}))

    catalog_path = output / 'catalog.json'
    catalog = [c for c in (json.loads(catalog_path.read_text()) if catalog_path.exists() else []) if c.get('dataset') != 'cho']
    catalog_path.write_text(json.dumps(catalog + items, indent=2, ensure_ascii=False) + '\n')
    return items


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--archive', type=Path, default=ROOT.parents[1] / '.cache/cho-source/Dataset.zip')
    parser.add_argument('--output', type=Path, default=ROOT / 'demo')
    args = parser.parse_args()
    for item in build(args.archive, args.output):
        print(item['id'], len(item['files']), 'files')


if __name__ == '__main__':
    main()
