"""Generate packages and isolated evaluator keys. No ingestion imports this module."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import random
import re
import zipfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
ARCHIVE_SHA = 'bc434c6cf96444fee3c915d434ac84e2a9c6492be4daebcfa4aed9cdff181f00'
DOWNLOAD = 'https://data.mendeley.com/public-files/datasets/pdnjz7zz5x/files/e7007dff-4061-494f-a9f5-16c20f457d1d/file_downloaded'
FDA = 'https://www.fda.gov/inspections-compliance-enforcement-and-criminal-investigations/warning-letters/sanofi-690604-01152025'
# Source column positions are checked against exact headers before use.
FIELDS = {
    'pH': (16, 'pH(pH:pH)', 'pH'),
    'temperature': (17, 'Temperature(T:K)', 'K'),
    'dissolved_oxygen': (12, 'Dissolved oxygen concentration(DO2:mg/L)', 'mg/L'),
    'substrate': (11, 'Substrate concentration(S:g/L)', 'g/L'),
    'product': (13, 'Penicillin concentration(P:g/L)', 'g/L'),
    'sugar_feed': (3, 'Sugar feed rate(Fs:L/h)', 'L/h'),
    'aeration': (1, 'Aeration rate(Fg:L/h)', 'L/h'),
    'agitation': (2, 'Agitator RPM(RPM:RPM)', 'rpm'),
    'acid_feed': (4, 'Acid flow rate(Fa:L/h)', 'L/h'),
    'base_feed': (5, 'Base flow rate(Fb:L/h)', 'L/h'),
    'volume': (14, 'Vessel Volume(V:L)', 'L'),
    'raman_1800': (637, '1800', 'unspecified'),
    'raman_1000': (1437, '1000', 'unspecified'),
    'raman_500': (1937, '500', 'unspecified'),
}


def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False)+'\n')


def sha(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def freeze_indpensim(archive: Path, output: Path, samples=4):
    from benchmarks.common.artifact_store import asset_path
    with asset_path(archive) as local:
        return _freeze_indpensim(local, output, samples)


def _freeze_indpensim(archive: Path, output: Path, samples=4):
    """Scan every batch, retaining original decimal strings and physical CSV rows."""
    if sha(archive) != ARCHIVE_SHA:
        raise ValueError('IndPenSim archive differs from publisher checksum')
    if not 1 <= samples <= 1000:
        raise ValueError('samples must be 1..1000')
    member = 'Mendeley_data/100_Batches_IndPenSim_V3.csv'
    # Fixed-seed reservoir per batch samples the whole trajectory, not only startup.
    counts, selected, rng = {}, {}, {}
    with zipfile.ZipFile(archive) as zipped:
        reader = csv.reader(io.TextIOWrapper(zipped.open(member), encoding='utf-8-sig', newline=''))
        headers = next(reader)
        # Publisher CSV has two unquoted commas in the PAT label. Repair ONLY
        # that verified header span, never shift/drop numerical data columns.
        if headers[33:36] != ['1- No Raman spec',' 1-Raman spec recorded','2-PAT control(PAT_ref:PAT ref)']:
            raise ValueError('Unexpected PAT header; manual source inspection required')
        headers[33:36] = [','.join(headers[33:36])]
        for parameter, (index, header, _) in FIELDS.items():
            if headers[index] != header:
                raise ValueError(f'Source schema drift: {parameter}: {headers[index]} != {header}')
        for row_index, row in enumerate(reader, 2):
            if not row:
                continue
            if len(row) != len(headers):
                raise ValueError(f'Source width mismatch at row {row_index}')
            batch = int(float(row[34]))
            if batch not in counts:
                counts[batch], selected[batch], rng[batch] = 0, [], random.Random(20261020+batch)
            counts[batch] += 1
            record = {'batch': f'IP-{batch:03}', 'time': row[0], 'original_row': row_index,
                      'original_member': member, 'fields': {}}
            for parameter, (index, header, unit) in FIELDS.items():
                raw = row[index]
                if raw and math.isfinite(float(raw)):
                    record['fields'][parameter] = {'raw': raw, 'value': float(raw), 'unit': unit,
                                                    'original_header': header, 'original_column': index+1}
            if len(selected[batch]) < samples:
                selected[batch].append(record)
            else:
                position = rng[batch].randrange(counts[batch])
                if position < samples:
                    selected[batch][position] = record
    if set(counts) != set(range(1, 101)):
        raise ValueError(f'Expected all 100 batches, found {len(counts)}')
    frozen = {'source': 'IndPenSim v2, Stephen Goldrick', 'doi': '10.17632/pdnjz7zz5x.2',
              'license': 'CC BY 4.0', 'archive_sha256': ARCHIVE_SHA, 'sampling': 'per-batch seeded reservoir',
              'samples_per_batch': samples, 'source_rows_per_batch': counts,
              'header_repair':'Joined the three unquoted PAT label fragments at physical header fields 34..36; numerical rows untouched.',
              'records': [r for b in sorted(selected) for r in sorted(selected[b], key=lambda x:float(x['time']))]}
    dump(output, frozen)
    return frozen


def header(parameter, unit, layout):
    aliases = {'pH':'Acidity', 'temperature':'Temp', 'dissolved_oxygen':'DO2',
               'substrate':'Residual sugar', 'product':'Penicillin', 'sugar_feed':'Feed rate'}
    name = aliases.get(parameter, parameter.replace('_', ' ')) if layout == 'B' else parameter
    return f'{name} [{unit}]'


def export_value(parameter, field, layout):
    """Recorded unit/locale changes; canonical truth remains original source value."""
    if layout=='B' and parameter=='temperature': return repr(field['value']-273.15),'degC'
    if layout=='B' and parameter=='sugar_feed': return repr(field['value']/0.06),'mL/min'
    if layout=='B' and parameter=='substrate': return field['raw'],'mg/mL'
    if layout=='B' and parameter=='pH': return field['raw'].replace('.',','),field['unit']
    return field['raw'],field['unit']


def write_xlsx(path, sheets):
    from openpyxl import Workbook
    book = Workbook(); book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = 'C2'
    book.properties.created = book.properties.modified = datetime(2026, 1, 1)
    buffer = io.BytesIO(); book.save(buffer)
    # openpyxl embeds wall-clock ZIP member dates; canonicalize archive metadata.
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0)); info.compress_type = zipfile.ZIP_DEFLATED
            content=source.read(name)
            if name=='docProps/core.xml':
                # Workbook.save overwrites modified, even if explicitly set.
                content=re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)',
                               rb'\g<1>2026-01-01T00:00:00Z\2',content)
            target.writestr(info, content)


def write_pdf(path, lines, scanned=False):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    canvas = Canvas(str(path), pagesize=(900, 640), invariant=1)
    pages = [lines[i:i+28] for i in range(0, len(lines), 28)] or [[]]
    for page in pages:
        if scanned:
            from PIL import Image, ImageDraw, ImageFont
            image = Image.new('RGB', (1800, 1280), 'white')
            draw = ImageDraw.Draw(image)
            font_path = '/System/Library/Fonts/Monaco.ttf'
            if not Path(font_path).exists(): font_path = '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'
            font = ImageFont.truetype(font_path, 20) if Path(font_path).exists() else ImageFont.load_default(size=20)
            for i, line in enumerate(page): draw.text((50, 40+i*40), line, fill='black', font=font)
            # Raster-only page with a small deterministic rotation; no hidden text layer.
            image = image.rotate(0.35, expand=False, fillcolor='white')
            canvas.drawImage(ImageReader(image), 0, 0, width=900, height=640)
        else:
            canvas.setFont('Courier', 10)
            for i, line in enumerate(page): canvas.drawString(25, 610-i*20, line)
        canvas.showPage()
    canvas.save()


def build_package(output, dataset, condition, records, *, seed=20261020):
    """The answer key lives outside raw/. The public delivery list contains no answers."""
    package_id = f'{dataset}-{condition}'
    base = output/package_id; raw = base/'raw'; raw.mkdir(parents=True, exist_ok=True)
    if any(raw.iterdir()):
        raise ValueError(f'{raw} must be empty; choose a fresh output directory')
    layout = 'B' if condition == 'unseen' else 'A'
    rng = random.Random(seed)
    truth = {'package_id':package_id, 'dataset':dataset, 'condition':condition, 'fields':[],
             'documents':{}, 'defects':[], 'relationships':[], 'transformations':[]}
    synthetic = dataset == 'sanofi'
    disclosure = ('SYNTHETIC RECONSTRUCTION: invented operational records; NOT actual Sanofi data.' if synthetic else
                  'IndPenSim v2 simulated measurements; reformatted and sampled; injected defects are artificial.')
    truth['disclosure'] = disclosure
    parameters = sorted({p for r in records for p in r['fields']})
    units = {p:next(export_value(p,r['fields'][p],layout)[1] for r in records if p in r['fields']) for p in parameters}
    if layout=='B':
        truth['transformations'].append({'operation':'unit_and_locale_conversion',
            'temperature':'K to degC, subtract 273.15','sugar_feed':'L/h to mL/min, divide by 0.06',
            'substrate':'g/L to mg/mL, same numeric value','pH':'decimal comma',
            'truth':'Canonical numerical truth remains the original source value and unit.'})
    headers = ['batch_id', 'elapsed_h']+[header(p,units[p],layout) for p in parameters]
    if layout == 'B': headers = ['Lot number', 'Process hours']+headers[2:]
    if condition == 'changed':
        headers[2], headers[-1] = headers[-1], 'Unqualified reading [unknown]'
        parameters[0], parameters[-1] = parameters[-1], parameters[0]
        truth['transformations'].append({'operation':'rename_and_reorder_column','field':parameters[-1]})
    rows = [headers]; process = 'bioreactor_process.xlsx' if layout == 'A' else 'fermentation_export.csv'
    sheet = 'Process' if layout == 'A' else None
    for index, record in enumerate(records):
        values = [record['batch'], record['time']]
        for column, parameter in enumerate(parameters, 3):
            field = record['fields'].get(parameter)
            value = export_value(parameter,field,layout)[0] if field else ''
            missing = condition == 'changed' and index == 0 and column == 3
            if missing:
                value = ''
                truth['defects'].append({'rule_code':'MISSING_FIELD','path':process,'row':index+2,'parameter':parameter})
            values.append(value)
            if field:
                from openpyxl.utils import get_column_letter
                truth['fields'].append({'batch':record['batch'],'time':record['time'],'parameter':parameter,
                    'value':field['value'],'unit':field['unit'], 'present':not missing,
                    'source':{'path':process,'sheet':sheet,'row':index+2,'column':headers[column-1],
                              'cell':f'{get_column_letter(column)}{index+2}' if sheet else None},
                    'original':{k:v for k,v in field.items() if k.startswith('original_')} |
                               {k:v for k,v in record.items() if k.startswith('original_')}})
        rows.append(values)
    if condition == 'changed':
        duplicate = list(rows[2]); numeric = next(i for i in range(2,len(duplicate)) if duplicate[i] != '')
        duplicate[numeric] = str(float(duplicate[numeric])+1.5); rows.append(duplicate)
        truth['defects'] += [{'rule_code':'DUPLICATE_CONFLICT','path':process,'row':len(rows),'parameter':parameters[numeric-2]},
                             {'rule_code':'LAYOUT_CHANGED','path':process}]
        # Expected extraction includes the conflicting source, never treats it as original ground truth.
        for field in list(truth['fields']):
            if field['source']['row'] == 3 and field['present']:
                copied = json.loads(json.dumps(field)); copied['source']['row'] = len(rows)
                if sheet:
                    copied['source']['cell'] = ''.join(c for c in copied['source']['cell'] if c.isalpha())+str(len(rows))
                if field['parameter'] == parameters[numeric-2]:
                    copied['value'] = float(duplicate[numeric]); copied['injected_conflict'] = True
                truth['fields'].append(copied)
        truth['transformations'].append({'operation':'duplicate_conflicting_row','source_row':3,'new_row':len(rows),'offset':1.5})
    if sheet:
        write_xlsx(raw/process, {sheet:rows, 'Delivery notes':[['Disclosure'],[disclosure]]})
    else:
        with (raw/process).open('w', newline='') as stream: csv.writer(stream).writerows(rows)
    truth['documents'][process] = 'instrument_export'
    # Native and scanned PDFs share a simple printable record form. Layout B changes separators.
    kinds = ['batch_record'] if not synthetic else ['batch_record','environmental_monitoring','microbiology',
                      'coa','maintenance','deviation_report','investigation_report']
    for doc_index, kind in enumerate(kinds):
        path = f'{kind}.pdf'
        scanned = synthetic and kind == 'microbiology'
        lines = [disclosure, f'Document type: {kind}', f'Form: {layout}', 'Page 1 of 1']
        if synthetic: lines.append('Public themes: contamination, pressure/leaks, deficient investigations (FDA 2025).')
        for i, record in enumerate(records[:2]):
            parameter = ('product' if not synthetic else {
                'batch_record':'pH','environmental_monitoring':'environment_cfu','microbiology':'bioburden',
                'coa':'product','maintenance':'pressure','deviation_report':'pressure',
                'investigation_report':'bioburden'}[kind])
            field = record['fields'][parameter]
            batch = record['batch']
            if condition == 'changed' and kind == 'coa' and i == 0:
                batch += '-TYPO'; truth['defects'].append({'rule_code':'BATCH_MISMATCH','path':path})
            # Structured, human-readable form; no coordinates or answer labels encoded in sources.
            sep = ' ; ' if layout == 'B' else ' | '
            rendered,rendered_unit=export_value(parameter,field,layout)
            text = sep.join([f'Batch: {batch}',f'Hour: {record["time"]}',
                             f'{header(parameter, rendered_unit, layout)}: {rendered}'])
            lines.append(text)
            truth['fields'].append({'batch':batch, 'time':record['time'], 'parameter':parameter,
                'value':field['value'],'unit':field['unit'],'present':True,
                'source':{'path':path,'page':1,'span':f'line:{len(lines)}'},
                'original':{k:v for k,v in record.items() if k.startswith('original_')}})
            truth['relationships'].append({'left':path,'right':process,'batch':batch,'relation':'batch_evidence'})
        if condition == 'changed' and kind == 'investigation_report':
            lines[3] = 'Page 1 of 2'
            lines.append('CAPA effectiveness: [not supplied]')
            truth['defects'].append({'rule_code':'MISSING_PAGE','path':path})
        write_pdf(raw/path,lines,scanned=scanned)
        truth['documents'][path] = kind
        truth['transformations'].append({'operation':'rasterize_rotate','path':path,'degrees':0.35} if scanned else
                                         {'operation':'render_native_pdf','path':path})
    truth['transformations'] += [{'operation':'layout','style':layout,'seed':seed},
                                 {'operation':'preserve_original_decimal_strings','calendar_timestamps':'not invented; elapsed process hours'}]
    files = [{'path':p.name, 'sha256':sha(p)} for p in sorted(raw.iterdir())]
    truth['files'] = files
    dump(base/'evaluation'/'answer_key.json',truth)
    public = {'id':package_id,'dataset':dataset,'condition':condition,'cmo':f'{dataset}-demo',
              'disclosure':disclosure,'files':files,'synthetic':synthetic,
              'scope':'sampled benchmark; not complete process qualification'}
    dump(base/'delivery.json', public)
    return public


def generate(output: Path, archive: Path, samples=4):
    output.mkdir(parents=True, exist_ok=True)
    frozen = freeze_indpensim(archive,output/'reference'/'indpensim.json',samples)
    manifests=[]
    conditions=['calibration','returning','unseen','changed']
    # Modulo split spreads each of the 4 control regimes across all conditions.
    for index, condition in enumerate(conditions):
        records=[r for r in frozen['records'] if (int(r['batch'].split('-')[1])-1)%4 == index]
        manifests.append(build_package(output,'indpensim',condition,records))
    rng = random.Random(20261020)
    facts = {'source':FDA,'public_facts':[
        {'section':'1','fact':'FDA documented recurring microbiological contamination and inadequate investigations.'},
        {'section':'2','fact':'FDA documented high-pressure and leakage events and departures from validated procedures.'},
        {'section':'3','fact':'FDA questioned equipment suitability and contamination controls.'},
        {'section':'4','fact':'FDA documented approximately 84 open, past-due investigations as of June 21, 2024.'}],
        'synthetic_details':'All package batch identifiers, dates, measurements, equipment, results and relationships are invented.'}
    dump(output/'reference'/'sanofi-public-facts.json',facts)
    for index, condition in enumerate(conditions):
        records=[]
        for batch in range(2):
            for hour in [0,12,24,36]:
                fields={p:{'raw':str(v),'value':v,'unit':unit} for p,v,unit in [
                    ('pH',round(rng.uniform(6.7,7.2),3),'pH'),('temperature',round(rng.uniform(308,311),3),'K'),
                    ('dissolved_oxygen',round(rng.uniform(3,7),3),'mg/L'),('product',round(rng.uniform(1,4),3),'g/L'),
                    ('pressure',round(rng.uniform(1.1,2.3),3),'bar'),('bioburden',float(rng.randrange(0,30)),'CFU/mL'),
                    ('environment_cfu',float(rng.randrange(0,12)),'CFU/m3')]}
                records.append({'batch':f'SYN-{index+1}{batch+1:02}', 'time':str(hour),'fields':fields})
        manifests.append(build_package(output,'sanofi',condition,records))
    dump(output/'catalog.json', manifests)
    return manifests


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--archive',type=Path,default=ROOT/'raw'/'100_Batches_IndPenSim.zip')
    parser.add_argument('--output',type=Path,default=ROOT/'generated')
    parser.add_argument('--samples-per-batch',type=int,default=4)
    args=parser.parse_args()
    print(json.dumps(generate(args.output,args.archive,args.samples_per_batch),indent=2))


if __name__ == '__main__': main()
