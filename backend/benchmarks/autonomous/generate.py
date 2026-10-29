"""Fictional manufacturing corpus, with physically separate evaluator labels."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import random
import re
import textwrap
import zipfile
import shutil
import tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

VERSION = 'fictional-manufacturing-v2'
CUTOFF = '2024-07-01T00:00:00+00:00'
CATEGORIES = {'batch': 150, 'historian': 200, 'laboratory': 100, 'deviation': 100,
              'maintenance': 100, 'procedure': 100, 'distractor': 250}
FORMATS = ('.txt', '.csv', '.xlsx', '.json', '.docx', '.pdf', '.parquet')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save(path, value):
    Path(path).write_text(json.dumps(value, indent=2, sort_keys=True) + '\n')


def write_record(path, rows, fmt):
    """Serializers only; extraction remains in manufacturing_data.ingest."""
    if fmt == '.csv':
        with path.open('w', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader(); writer.writerows(rows)
    elif fmt == '.xlsx':
        from openpyxl import Workbook
        book = Workbook(); sheet = book.active; sheet.title = 'Export'
        sheet.append(list(rows[0]))
        for row in rows: sheet.append(list(row.values()))
        book.save(path)
        # XLSX container timestamps are not data; freeze them for reproducible bytes.
        with zipfile.ZipFile(path) as archive:
            members = [(m.filename, archive.read(m)) for m in archive.infolist()]
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, body in members:
                if name == 'docProps/core.xml':
                    body = re.sub(rb'<dcterms:modified[^>]*>[^<]*</dcterms:modified>',
                                  b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:type="dcterms:W3CDTF">2024-01-01T00:00:00Z</dcterms:modified>', body)
                    body = re.sub(rb'<dcterms:created[^>]*>[^<]*</dcterms:created>',
                                  b'<dcterms:created xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:type="dcterms:W3CDTF">2024-01-01T00:00:00Z</dcterms:created>', body)
                item = zipfile.ZipInfo(name, (2024, 1, 1, 0, 0, 0)); item.compress_type = zipfile.ZIP_DEFLATED
                archive.writestr(item, body)
    elif fmt == '.parquet':
        import polars as pl
        pl.DataFrame(rows).write_parquet(path)
    elif fmt == '.json':
        save(path, rows)
    elif fmt == '.docx':
        from docx import Document
        doc = Document()
        doc.core_properties.created = datetime(2024, 1, 1)
        doc.core_properties.modified = datetime(2024, 1, 1)
        for row in rows:
            for key, value in row.items(): doc.add_paragraph(f'{key}: {value}')
        doc.save(path)
        with zipfile.ZipFile(path) as archive:
            members = [(m.filename, archive.read(m)) for m in archive.infolist()]
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, body in members:
                item = zipfile.ZipInfo(name, (2024, 1, 1, 0, 0, 0)); item.compress_type = zipfile.ZIP_DEFLATED
                archive.writestr(item, body)
    elif fmt in {'.pdf', '.scan.pdf'}:
        from reportlab.pdfgen.canvas import Canvas
        canvas = Canvas(str(path), invariant=1)
        lines = [line for key, value in rows[0].items() for line in textwrap.wrap(f'{key}: {value}', 93)]
        if fmt == '.scan.pdf':
            from PIL import Image, ImageDraw, ImageFont
            from reportlab.lib.utils import ImageReader
            image = Image.new('RGB', (1700, 2200), 'white'); draw = ImageDraw.Draw(image)
            font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 25)
            for i, line in enumerate(lines): draw.text((60, 60 + i * 36), line, fill='black', font=font)
            canvas.drawImage(ImageReader(image), 0, 0, width=595, height=770)
        else:
            for i, line in enumerate(lines): canvas.drawString(35, 810 - i * 13, line)
        canvas.save()
    else:
        path.write_text('\n\n'.join('\n'.join(f'{k}: {v}' for k, v in r.items()) for r in rows))


def generate(destination, count=1000, seed=417, scenario_seed=718):
    destination = Path(destination)
    if count < 100: raise ValueError('at least 100 files are required')
    if destination.exists() and any(destination.iterdir()): raise ValueError('destination must be empty')
    raw = destination / 'raw'; raw.mkdir(parents=True)
    hidden = destination / 'evaluator'; hidden.mkdir(mode=0o700)
    rng = random.Random(seed); scenarios = random.Random(scenario_seed)
    identifiers = lambda prefix, n: [f'{prefix}-{scenarios.getrandbits(24):06X}' for _ in range(n)]
    questions = []; special = []
    families = {
        'microbiological': [
            ('support', 'QC confirmed matching isolate fingerprints in three rejected cultures; pressure runs remained within limits.'),
            ('support', 'Shared sampling manifold showed a seal defect; positive cultures followed sampling on this manifold.'),
            ('counter', 'The same media lot was used in three successful comparator batches; supplier attribution is unsupported.'),
            ('support', 'The seal was replaced before three subsequent negative culture results; temporal association requires a controlled test.'),
            ('future', 'Later sequencing and controlled sampling study completed; unavailable to the historical investigation.')],
        'transfer_pressure': [
            ('support', 'Three transfers exceeded 1.8 bar after filter differential pressure rose; transfer connections leaked.'),
            ('support', 'Maintenance found a partially blocked vent filter shared by these transfers; pressure calibration passed.'),
            ('counter', 'Engineering tests found the proposed bag defect did not reproduce leakage at validated pressure.'),
            ('support', 'Comparator transfers with the same bag lot and clean vent filters stayed below 1.2 bar without leakage.'),
            ('future', 'Later controlled vent restriction experiment confirmed a mechanism; this evidence was not yet available.')],
        'procedure_change': [
            ('support', 'Operator notes describe substituting a transfer liner following an informal pre-job briefing.'),
            ('support', 'The approved procedure specified the original liner and required change-control approval before substitution.'),
            ('counter', 'Deviation summary says no procedural change occurred; the contemporaneous operator notes contradict that statement.'),
            ('support', 'The change register contains no approved substitution; records associate three runs with the changed liner.'),
            ('future', 'A later quality review approved a revised procedure; approval did not exist at the investigation cutoff.')],
        'inconclusive': [
            ('support', 'Three runs had atypical particle counts; samples were discarded before chemical identification.'),
            ('support', 'Maintenance records omit component serial numbers; the proposed equipment association is unresolved.'),
            ('counter', 'A suspected calibration failure was not reproduced; retained standards passed an independent check.'),
            ('support', 'No retained samples or complete transfer logs are available; request new sampling and component traceability.'),
            ('future', 'Later samples identified particle composition; these results cannot support an earlier conclusion.')]}
    contributors = ['sampling manifold seal', 'blocked vent filter', 'unapproved liner substitution', None]
    for (family, notes), contributor in zip(families.items(), contributors):
        batches = identifiers('B', 6); campaign = identifiers('CP', 1)[0]; equipment = identifiers('EQ', 1)[0]
        question = {'id': family, 'objective': f'Investigate manufacturing failures in campaign {campaign}. Find supporting and contradictory evidence, affected batches, competing explanations, missing evidence and engineering follow-up. Do not assert a verified root cause.',
                    'campaign': campaign, 'affected_batches': batches[:3], 'negative_controls': batches[3:],
                    'contributors': [contributor] if contributor else [], 'inconclusive': contributor is None,
                    'critical': [], 'next_actions': ['controlled test', 'QA review', 'retrieve missing records']}
        questions.append(question)
        for kind, note in notes:
            special.append((question, kind, note, batches, campaign, equipment))
    positions = dict(zip(rng.sample(range(count), len(special)), special))
    categories = [category for category, n in CATEGORIES.items() for _ in range(round(n * count / 1000))]
    categories = (categories + ['distractor'] * count)[:count]; rng.shuffle(categories)
    files = []; expected = []; ordinary = []
    for i in range(count):
        fmt = '.scan.pdf' if i % 97 == 0 else FORMATS[i % len(FORMATS)]
        rel = f'{rng.randrange(8):02d}/{rng.getrandbits(80):020x}{fmt}'
        path = raw / rel; path.parent.mkdir(exist_ok=True)
        batch = f'B-{rng.getrandbits(24):06X}'; campaign = f'CP-{rng.getrandbits(24):06X}'
        equipment = f'EQ-{rng.getrandbits(24):06X}'
        note = 'Routine reconciliation completed; calibration, sterility, transfer and pressure reviewed without confirmed discrepancy.'
        available = '2024-06-25T00:00:00+00:00'; affected = batch; comparator = ''
        if i in positions:
            q, kind, note, batches, campaign, equipment = positions[i]
            batch = batches[0]; affected = ', '.join(batches[:3]); comparator = ', '.join(batches[3:])
            if kind == 'future': available = '2024-07-08T00:00:00+00:00'
            q['critical'].append({'path': rel, 'kind': kind, 'span': note})
        else: ordinary.append(i)
        if i not in positions and rng.random() < .65:
            # Campaign/equipment matches alone are deliberately insufficient:
            # routine and misleading records share the target's real context.
            q = rng.choice(questions)
            campaign = q['campaign']
            batch = rng.choice(q['negative_controls'])
            affected = batch
            note = rng.choice([
                note,
                'Night shift and media supplier changed during this campaign; successful comparator runs share both factors. No causal test is recorded.',
                'Pressure alarm in a training log was simulated; this entry does not describe a production failure.',
                'Meeting minutes discuss previous sterility events at an unrelated site; the local campaign is mentioned only for scheduling.',
                'Draft investigation references a possible component defect; identification and causality remain unconfirmed.'])
        record = {'disclosure': 'FICTIONAL benchmark factory. Not Sanofi operational data.',
                  'record_type': categories[i], 'batch_id': batch, 'campaign_id': campaign,
                  'equipment_id': equipment, 'event_time': '2024-06-20T00:00:00+00:00',
                  'available_at': available, 'time(h)': 0.0, 'temperature(degC)': round(36.9 + rng.uniform(-.15, .15), 3),
                  'pressure(bar)': 1.9 if 'exceeded 1.8' in note else 1.1,
                  'associated_batches': affected, 'comparator_batches': comparator, 'record_note': note}
        rows = [record]
        if categories[i] == 'historian' and fmt in {'.csv', '.xlsx', '.parquet'}:
            rows = [dict(record, **{'time(h)': j / 2, 'temperature(degC)': round(36.9 + .04 * __import__('math').sin(j / 3), 3),
                                   'pressure(bar)': round((1.6 if 'exceeded 1.8' in note else 1.08) + .015 * j, 3)}) for j in range(24)]
        write_record(path, rows, fmt)
        stamp = datetime.fromisoformat(available).timestamp(); os.utime(path, (stamp, stamp))
        files.append({'path': rel, 'format': fmt, 'category': categories[i], 'sha256': digest(path),
                      'bytes': path.stat().st_size, 'available_at': available})
        expected.append({'path': rel, 'fields': rows[0]})
    # Duplicate ordinary records only: do not multiply causal support.
    duplicates = []
    for fmt in FORMATS:
        group = [i for i in ordinary if files[i]['format'] == fmt]
        for a, b in zip(group[:max(1, len(group)//10)], group[len(group)//2:]):
            source = raw / files[a]['path']; target = raw / files[b]['path']
            target.write_bytes(source.read_bytes()); os.utime(target, (source.stat().st_mtime,) * 2)
            files[b].update(sha256=digest(target), bytes=target.stat().st_size, available_at=files[a]['available_at'])
            expected[b]['fields'] = expected[a]['fields']; duplicates.append([files[a]['path'], files[b]['path']])
    corrupted, unsupported = ordinary[-2:]
    for index, extension, body, status in [(corrupted, '.xlsx', b'broken spreadsheet', 'error'),
                                           (unsupported, '.bin', b'unsupported proprietary export', 'preserved_unsupported')]:
        old = raw / files[index]['path']; new = old.with_suffix(extension); old.unlink(); new.write_bytes(body)
        stamp = datetime.fromisoformat(files[index]['available_at']).timestamp(); os.utime(new, (stamp, stamp))
        files[index].update(path=str(new.relative_to(raw)), format=extension, sha256=digest(new), bytes=len(body), expected_status=status)
    truth = {'version': VERSION, 'questions': questions, 'files': files, 'duplicates': duplicates,
             'extraction_samples': [e for i, e in enumerate(expected) if i not in {corrupted, unsupported}][:35]}
    save(hidden / 'truth.json', truth); (hidden / 'truth.json').chmod(0o600)
    manifest = {'version': VERSION, 'seed': seed, 'scenario_seed': scenario_seed, 'files': count,
                'formats': dict(Counter(f['format'] for f in files)), 'categories': dict(Counter(categories)),
                'bytes': sum(f['bytes'] for f in files), 'cutoff': CUTOFF,
                'corpus_sha256': hashlib.sha256(json.dumps(files, sort_keys=True).encode()).hexdigest(),
                'generator_sha256': digest(__file__),
                'disclosure': 'All operational evidence is fictional. FDA context is a separate public source.'}
    save(destination / 'manifest.json', manifest)
    return destination / 'manifest.json'


def extend_corpus(base, destination, count, seed=912):
    """Controlled distractor growth: original source bytes and critical spans stay fixed."""
    base = Path(base); destination = Path(destination)
    original = json.loads((base / 'evaluator/truth.json').read_text())
    if count <= len(original['files']): raise ValueError('extended size must exceed base size')
    if destination.exists() and any(destination.iterdir()): raise ValueError('destination must be empty')
    raw = destination / 'raw'; raw.mkdir(parents=True)
    hidden = destination / 'evaluator'; hidden.mkdir(mode=0o700)
    for source in original['files']:
        target = raw / source['path']; target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(base / 'raw' / source['path'], target)
    files = list(original['files'])
    with tempfile.TemporaryDirectory(prefix='fictional-distractors-') as temp:
        noise = Path(temp) / 'noise'
        generate(noise, count + 100, seed=seed, scenario_seed=json.loads((base / 'manifest.json').read_text())['scenario_seed'])
        truth = json.loads((noise / 'evaluator/truth.json').read_text())
        reserved = {e['path'] for q in truth['questions'] for e in q['critical']}
        candidates = [f for f in truth['files'] if f['path'] not in reserved and not f.get('expected_status')]
        random.Random(seed).shuffle(candidates)
        for source in candidates[:count - len(files)]:
            rel = 'additional/' + source['path']; target = raw / rel; target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(noise / 'raw' / source['path'], target)
            files.append({**source, 'path': rel})
    if len(files) != count: raise ValueError('not enough distractor sources')
    original['files'] = files; original['version'] = VERSION
    save(hidden / 'truth.json', original); (hidden / 'truth.json').chmod(0o600)
    manifest = {**json.loads((base / 'manifest.json').read_text()), 'files': count, 'scaling_seed': seed,
                'scaling_design': 'Byte-identical base corpus plus distractors; no additional critical or future evidence.',
                'base_corpus_sha256': json.loads((base / 'manifest.json').read_text())['corpus_sha256'],
                'bytes': sum(f['bytes'] for f in files), 'formats': dict(Counter(f['format'] for f in files)),
                'categories': dict(Counter(f['category'] for f in files)),
                'corpus_sha256': hashlib.sha256(json.dumps(files, sort_keys=True).encode()).hexdigest(),
                'generator_sha256': digest(__file__)}
    save(destination / 'manifest.json', manifest)
    return destination / 'manifest.json'


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('output'); parser.add_argument('--count', type=int, default=1000)
    parser.add_argument('--seed', type=int, default=417); parser.add_argument('--scenario-seed', type=int, default=718)
    parser.add_argument('--base', help='Preserve this corpus and add only distractors for controlled scaling')
    args = parser.parse_args()
    from contextlib import ExitStack
    from benchmarks.common.artifact_store import asset_path, publish_generated
    with ExitStack() as stack:
        base = stack.enter_context(asset_path(args.base)) if args.base else None
        remote = str(args.output).startswith('dagshub:')
        output = Path(stack.enter_context(tempfile.TemporaryDirectory(prefix='biodemo-generate-'))) / 'data' if remote else Path(args.output)
        result = extend_corpus(base, output, args.count, args.seed) if base else generate(output, args.count, args.seed, args.scenario_seed)
        print(json.dumps(publish_generated(output,str(args.output).split(':',1)[1])) if remote else result)
