"""Independent source readers check the existing ingested store, not its own output.

Checks systematic first-row samples across structured formats. OCR and scientific
interpretation are explicitly outside this automated round-trip measurement.
"""
import argparse
import csv
import json
import math
import re
from collections import Counter
from pathlib import Path

from src.manufacturing_data import ManufacturingStore
from .generate import save


def check(corpus, workspace, output, maximum=24):
    corpus = Path(corpus); store = ManufacturingStore(Path(workspace) / 'state/store')
    sources = [s for s in store.visible_sources(as_of='2024-07-01T00:00:00Z', scopes=('evaluation',))
               if s['format'] in {'.csv', '.xlsx', '.parquet'} and s['extraction_status'] == 'indexed'][:maximum]
    wanted = [s['id'] for s in sources]
    rows = store.select('SELECT * FROM observations WHERE source_id IN (' + ','.join('?' for _ in wanted) +
                        ") AND ((extraction_method='parquet' AND source_row=1) OR (extraction_method<>'parquet' AND source_row=2))",
                        wanted, as_of='2024-07-01T00:00:00Z', scopes=('evaluation',))
    observed = {(r['source_id'], r['name']): r for r in rows}
    checks = []
    for source in sources:
        original = corpus / 'raw' / source['path']
        if source['format'] == '.csv':
            with original.open(newline='') as stream: expected = next(csv.DictReader(stream))
        elif source['format'] == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(original, read_only=True, data_only=True)
            values = book.active.iter_rows(values_only=True); expected = dict(zip(next(values), next(values))); book.close()
        else:
            import pyarrow.parquet as pq
            expected = pq.read_table(original).slice(0, 1).to_pylist()[0]
        for name, value in expected.items():
            if value in ('', None): continue
            row = observed.get((source['id'], name)); equal = False
            if row:
                equal = str(value) == row['value_text']
                if not equal:
                    try: equal = float(value) == row['value_number']
                    except (ValueError, TypeError): pass
            expected_unit = re.search(r'\(([^()]*)\)$', name)
            ref = row['source_ref'] if row else {}
            if isinstance(ref, str): ref = json.loads(ref)
            checks.append({'path': source['path'], 'field': name, 'value_matches': equal,
                           'unit_matches': (row.get('unit') == expected_unit[1]) if row and expected_unit else None,
                           'provenance_matches': bool(row and ref.get('path') == source['path'] and ref.get('column') == name and ref.get('row') == (1 if source['format'] == '.parquet' else 2))})
    units = [c['unit_matches'] for c in checks if c['unit_matches'] is not None]
    result = {'method': 'Systematic first-row sample; csv/openpyxl/pyarrow source readers vs existing ingested store.',
              'files': len(sources), 'formats': dict(Counter(s['format'] for s in sources)), 'fields': len(checks),
              'field_accuracy': sum(c['value_matches'] for c in checks) / len(checks),
              'unit_accuracy': sum(units) / len(units) if units else None,
              'provenance_accuracy': sum(c['provenance_matches'] for c in checks) / len(checks),
              'checks': checks, 'not_evaluated': ['OCR transcription accuracy', 'complete trajectories', 'scientific interpretation']}
    save(output, result); print(json.dumps({k: v for k, v in result.items() if k != 'checks'})); return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('corpus'); parser.add_argument('workspace'); parser.add_argument('output')
    args = parser.parse_args(); check(args.corpus, args.workspace, args.output)
