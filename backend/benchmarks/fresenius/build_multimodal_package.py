from __future__ import annotations

import csv
import hashlib
import importlib.metadata
import json
import io
import platform
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from .temporal_replay import FROZEN_AGENT_DATA, state_at
from .source_formats import SCAN_SEED, add_discrete_process_charts, add_scanned_lab_report


ROOT = Path(__file__).resolve().parent
DESTINATIONS = {
    "batches": ("batch_records", "batch_records.csv"),
    "lab_results": ("laboratory", "lims_export.csv"),
    "process_events": ("manufacturing", "historian_export.csv"),
    "environmental_monitoring": ("laboratory", "environmental_monitoring.csv"),
    "deviations": ("quality", "deviations.csv"),
    "maintenance": ("maintenance", "work_orders.csv"),
    "complaints": ("quality", "complaints.csv"),
    "operator_events": ("operators", "operator_events.csv"),
    "incident_memory": ("quality", "prior_incidents.csv"),
}


def _json_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, separators=(",", ":"))
    return str(value)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _normalize_xlsx_zip(path: Path) -> None:
    """Remove wall-clock ZIP entry timestamps from otherwise deterministic exports."""
    original = path.read_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original), "r") as source, zipfile.ZipFile(output, "w") as target:
        for member in source.infolist():
            entry = zipfile.ZipInfo(member.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = member.compress_type
            entry.external_attr = member.external_attr
            entry.create_system = member.create_system
            body = source.read(member.filename)
            if member.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:modified[^>]*>)[^<]+",
                              rb"\g<1>2000-01-01T00:00:00Z", body)
            target.writestr(entry, body)
    path.write_bytes(output.getvalue())


def build_package(output: Path = ROOT / "raw_package", *, as_of: str | None = None,
                  map_output: Path | None = None) -> dict[str, Any]:
    """Render only fields visible at as_of; every source row has a hidden map."""
    as_of = as_of or "9999-12-31T23:59:59Z"
    rows = state_at(as_of)
    output.mkdir(parents=True, exist_ok=True)
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[row["_table"]].append(row)
    mapping: list[dict[str, Any]] = []
    files = []
    for table, (directory, filename) in DESTINATIONS.items():
        source_rows = grouped[table]
        if not source_rows:
            continue
        path = output / directory / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        columns = sorted({key for row in source_rows for key in row if key not in {"_table", "field_available_at"}})
        with path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns)
            writer.writeheader()
            for index, row in enumerate(source_rows, 2):
                writer.writerow({field: _json_cell(row.get(field)) for field in columns})
                for column in columns:
                    if column in row:
                        mapping.append({"source_file": str(path.relative_to(output)), "row": index,
                                        "column": column, "canonical_record_id": row["_evidence_id"],
                                        "canonical_field": column})
        files.append({"path": str(path.relative_to(output)), "sha256": _sha(path), "records": len(source_rows)})

    # A second lab export preserves the same observations but uses local time,
    # aliased columns and percent-of-limit display for ratio measurements.
    lab_rows = grouped["lab_results"]
    if lab_rows:
        path = output / "laboratory" / "lims_review_local.csv"
        fields = ["Record", "Lot / batch", "Released local (UTC-07)", "Test",
                  "Result (% of limit)", "Limit basis", "Categorical result", "Review status"]
        aliases = {"Record": "_evidence_id", "Lot / batch": "batch_id",
                   "Released local (UTC-07)": "available_at", "Test": "test_name",
                   "Result (% of limit)": "numerical_result", "Limit basis": "unit",
                   "Categorical result": "categorical_result", "Review status": "status"}
        with path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            for index, row in enumerate(lab_rows, 2):
                timestamp = datetime.fromisoformat(row["available_at"].replace("Z", "+00:00"))
                local = timestamp.astimezone(timezone(timedelta(hours=-7))).isoformat()
                value = row.get("numerical_result")
                displayed = float(value)*100 if isinstance(value, (int, float)) else ""
                values = {"Record": row["_evidence_id"], "Lot / batch": row.get("batch_id", ""),
                          "Released local (UTC-07)": local, "Test": row.get("test_name", ""),
                          "Result (% of limit)": displayed,
                          "Limit basis": row.get("unit", ""),
                          "Categorical result": row.get("categorical_result", ""),
                          "Review status": row.get("status", "")}
                writer.writerow(values)
                for alias, canonical in aliases.items():
                    mapping.append({"source_file": str(path.relative_to(output)), "row": index,
                                    "column": alias, "source_record_id_column": "Record",
                                    "canonical_record_id": row["_evidence_id"],
                                    "canonical_field": canonical,
                                    "conversion": "UTC to fixed UTC-07 instant" if alias == "Released local (UTC-07)"
                                                  else "ratio multiplied by 100" if alias == "Result (% of limit)" else None})
        files.append({"path": str(path.relative_to(output)), "sha256": _sha(path),
                      "records": len(lab_rows), "format": "aliased_local_time_csv"})

    # A second, independently parseable export format tests duplicate linking.
    from openpyxl import Workbook
    for table, filename in (("lab_results", "microbiology_results.xlsx"),
                            ("maintenance", "maintenance_work_orders.xlsx")):
        source_rows = grouped[table]
        if not source_rows:
            continue
        directory = "laboratory" if table == "lab_results" else "maintenance"
        path = output / directory / filename
        columns = sorted({key for row in source_rows for key in row if key not in {"_table", "field_available_at"}})
        book = Workbook()
        book.properties.created = datetime(2000, 1, 1)
        book.properties.modified = datetime(2000, 1, 1)
        sheet = book.active
        sheet.title = "Results" if table == "lab_results" else "Work orders"
        sheet.append(columns)
        for index, row in enumerate(source_rows, 2):
            sheet.append([_json_cell(row.get(field)) for field in columns])
            for column_index, column in enumerate(columns, 1):
                if column in row:
                    from openpyxl.utils import get_column_letter
                    mapping.append({"source_file": str(path.relative_to(output)), "sheet": sheet.title,
                                    "cell": f"{get_column_letter(column_index)}{index}",
                                    "canonical_record_id": row["_evidence_id"], "canonical_field": column})
        book.save(path)
        _normalize_xlsx_zip(path)
        files.append({"path": str(path.relative_to(output)), "sha256": _sha(path), "records": len(source_rows)})

    # A readable report gives the package a second format without modifying observations.
    report = output / "quality" / "source_review.pdf"
    report.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(report), pagesize=letter, invariant=1)
    pdf.setTitle("Synthetic factory records - source review")
    page = 1
    y = 750
    pdf.setFont("Helvetica", 9)
    pdf.drawString(40, y, "SYNTHETIC FACTORY RECORDS - simulated source review")
    y -= 22
    for row in grouped["deviations"] + grouped["complaints"]:
        if y < 55:
            pdf.showPage(); page += 1; y = 750; pdf.setFont("Helvetica", 9)
        line = f"{row['_evidence_id']} | {row.get('batch_id', row.get('lot', ''))} | {row.get('available_at', '')}"
        pdf.drawString(40, y, line[:105]); y -= 14
        mapping.append({"source_file": str(report.relative_to(output)), "page": page,
                        "section": "source review", "source_span": line,
                        "canonical_record_id": row["_evidence_id"], "canonical_field": "summary"})
    pdf.save()
    files.append({"path": str(report.relative_to(output)), "sha256": _sha(report),
                  "records": len(grouped["deviations"]) + len(grouped["complaints"])})
    if grouped["lab_results"]:
        scan, count = add_scanned_lab_report(output, grouped["lab_results"], mapping)
        files.append({"path": str(scan.relative_to(output)), "sha256": _sha(scan),
                      "records": count, "format": "image_only_pdf", "seed": SCAN_SEED})
    if grouped["process_events"]:
        chart, count = add_discrete_process_charts(output, grouped["process_events"], mapping)
        files.append({"path": str(chart.relative_to(output)), "sha256": _sha(chart),
                      "records": count, "format": "discrete_event_chart"})
    manifest = {"version": "fresenius-source-package-v2.1", "source_version": "v1.1-frozen",
                "as_of": as_of, "synthetic_notice": "All factory records are reconstructed synthetic data, not Fresenius internal files.",
                "source_sha256": _sha(FROZEN_AGENT_DATA), "files": files,
                "record_count": len(rows), "generator": "benchmarks.fresenius.build_multimodal_package.v2.1",
                "scan_artifact_seed": SCAN_SEED,
                "python_version": platform.python_version(),
                "generator_dependencies": {name: importlib.metadata.version(name)
                                           for name in ("Pillow", "reportlab", "openpyxl")},
                "local_time_rule": "fixed UTC-07 display; same instant as canonical UTC",
                "ratio_rule": "numerical ratio displayed as percent of its stated limit",
                "chart_method": "discrete source event points only; no interpolated trajectory"}
    (output / "package_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    if map_output:
        map_output.parent.mkdir(parents=True, exist_ok=True)
        map_output.write_text(json.dumps({"as_of": as_of, "mappings": mapping}, indent=2), encoding="utf-8")
    return manifest
