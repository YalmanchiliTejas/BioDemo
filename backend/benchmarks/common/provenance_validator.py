from __future__ import annotations

import csv
import json
from pathlib import Path


def validate_mapping(package: Path, hidden_map: Path) -> dict:
    """Check source pointers against generated files without exposing the map to agents."""
    mapping = json.loads(hidden_map.read_text(encoding="utf-8"))["mappings"]
    missing_files = sorted({item["source_file"] for item in mapping
                            if not (package / item["source_file"]).is_file()})
    bad_csv_rows = []
    by_file: dict[str, list[dict]] = {}
    for item in mapping:
        if item["source_file"].endswith(".csv"):
            by_file.setdefault(item["source_file"], []).append(item)
    for filename, pointers in by_file.items():
        with (package / filename).open(newline="", encoding="utf-8-sig") as stream:
            rows = list(csv.DictReader(stream))
        for pointer in pointers:
            row_index = pointer["row"]-2
            id_column = pointer.get("source_record_id_column", "_evidence_id")
            if (row_index < 0 or row_index >= len(rows)
                    or rows[row_index].get(id_column) != pointer["canonical_record_id"]
                    or pointer["column"] not in rows[row_index]):
                bad_csv_rows.append(pointer)
                if len(bad_csv_rows) >= 10:
                    break
    bad_xlsx_cells = []
    xlsx_files = {item["source_file"] for item in mapping if item["source_file"].endswith(".xlsx")}
    if xlsx_files:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        for filename in xlsx_files:
            book = load_workbook(package / filename, read_only=True, data_only=True)
            headers = {sheet.title: next(sheet.iter_rows(max_row=1, values_only=True), ()) for sheet in book}
            for pointer in (item for item in mapping if item["source_file"] == filename):
                letter, _ = coordinate_from_string(pointer["cell"])
                column = column_index_from_string(letter)-1
                row_headers = headers[pointer["sheet"]]
                if column >= len(row_headers) or row_headers[column] != pointer["canonical_field"]:
                    bad_xlsx_cells.append(pointer)
                    if len(bad_xlsx_cells) >= 10:
                        break
            book.close()
    bad_pdf_spans = []
    bad_pdf_regions = []
    pdf_files = {item["source_file"] for item in mapping if item["source_file"].endswith(".pdf")}
    if pdf_files:
        from pypdf import PdfReader
        for filename in pdf_files:
            reader = PdfReader(str(package / filename))
            text_cache = {}
            for pointer in (item for item in mapping if item["source_file"] == filename):
                page_number = pointer["page"]
                if page_number < 1 or page_number > len(reader.pages):
                    bad_pdf_regions.append(pointer)
                    continue
                page = reader.pages[page_number-1]
                region = pointer.get("image_region") or pointer.get("point_region")
                if region and not (0 <= region[0] < region[2] <= float(page.mediabox.width)
                                   and 0 <= region[1] < region[3] <= float(page.mediabox.height)):
                    bad_pdf_regions.append(pointer)
                if pointer.get("image_region"):
                    if not page.images:
                        bad_pdf_regions.append(pointer)
                else:
                    if page_number not in text_cache:
                        text_cache[page_number] = page.extract_text() or ""
                    if pointer["canonical_record_id"] not in text_cache[page_number]:
                        bad_pdf_spans.append(pointer)
    return {"mapping_count": len(mapping), "missing_files": missing_files,
            "bad_csv_rows": bad_csv_rows, "bad_xlsx_cells": bad_xlsx_cells,
            "bad_pdf_spans": bad_pdf_spans, "bad_pdf_regions": bad_pdf_regions,
            "pass": not any((missing_files, bad_csv_rows, bad_xlsx_cells, bad_pdf_spans, bad_pdf_regions))}
