"""Batch original-source verification: one hash/read per source and sheet."""
from __future__ import annotations

import hashlib
import re
import json
from datetime import date,datetime
from collections import defaultdict
from pathlib import Path


def verify_originals(store,rows):
    grouped=defaultdict(list)
    for row in rows:grouped[row['source_id']].append(row)
    results={}
    with store.connect() as con:
        sources={sid:dict(con.execute('SELECT * FROM sources WHERE id=?',(sid,)).fetchone()) for sid in grouped}
    for sid,items in grouped.items():
        source=sources[sid]
        path=store.originals/source['sha256'][:2]/(source['sha256']+Path(source['path']).suffix.lower())
        integrity=hashlib.sha256(path.read_bytes()).hexdigest()==source['sha256']
        observed={}
        if source['format'] in {'.xlsx','.xlsm'}:
            from openpyxl import load_workbook
            wanted=defaultdict(lambda:defaultdict(list))
            for item in items:
                ref=item['source_ref']
                if ref.get('row') is not None:wanted[ref.get('sheet')][ref['row']].append(item)
            workbook=load_workbook(path,read_only=True,data_only=True)
            try:
                for name,targets in wanted.items():
                    if name not in workbook.sheetnames:continue
                    sheet=workbook[name]
                    headers=[str(v) if v is not None else f'column_{i}' for i,v in enumerate(next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),()),1)]
                    for index,values in enumerate(sheet.iter_rows(min_row=2,max_row=max(targets),values_only=True),2):
                        for item in targets.get(index,[]):
                            ref=item['source_ref'];column=ref.get('column')
                            if column in headers:
                                observed[item['id']]=values[headers.index(column)]
            finally:workbook.close()
        elif source['format'] in {'.csv','.tsv','.parquet'}:
            from .ingest import _rows
            wanted=defaultdict(list)
            for item in items:
                if item['source_ref'].get('row') is not None:wanted[item['source_ref']['row']].append(item)
            if wanted:
                for sheet,index,headers,values in _rows(path):
                    for item in wanted.get(index,[]):
                        if item['source_ref'].get('sheet')==sheet:
                            column=item['source_ref'].get('column')
                            if column in headers:observed[item['id']]=values[headers.index(column)]
                    if index>=max(wanted):break
        elif source['format'] == '.docx':
            from docx import Document
            doc = Document(path)
            for item in items:
                span = item['source_ref'].get('span', '')
                paragraph = re.fullmatch(r'paragraph:(\d+)', span)
                cell = re.fullmatch(r'table:(\d+):row:(\d+):cell:(\d+)', span)
                try:
                    if paragraph: observed[item['id']] = doc.paragraphs[int(paragraph[1])-1].text
                    elif cell: observed[item['id']] = doc.tables[int(cell[1])-1].cell(int(cell[2])-1, int(cell[3])-1).text
                except IndexError:
                    pass
        elif any(item['extraction_method']=='pypdf' for item in items):
            from .ingest import _document_observations
            wanted={item['id'] for item in items}
            for candidate in _document_observations(sid,source['path'],path,source['access_scope'],source['available_at'],'pypdf'):
                if candidate.observation_id in wanted:observed[candidate.observation_id]=candidate.value_text
        elif any(item['extraction_method']=='docling' for item in items) and source['format']=='.pdf':
            # Independent native-text agreement can verify a localized Docling
            # element. A raster page has no independent text and stays pending.
            from pypdf import PdfReader
            reader=PdfReader(path)
            page_text={}
            for item in items:
                page=item['source_ref'].get('page')
                if not isinstance(page,int) or not 1<=page<=len(reader.pages): continue
                if page not in page_text:
                    page_text[page]=re.sub(r'\s+',' ',reader.pages[page-1].extract_text() or '').strip()
                candidate=re.sub(r'\s+',' ',item['value_text'] or '').strip()
                if candidate and candidate in page_text[page]: observed[item['id']]=item['value_text']
        elif source['format'] == '.json' and any(str(item['source_ref'].get('span', '')).startswith('record:') for item in items):
            records = json.loads(path.read_text(encoding='utf-8'), parse_float=str, parse_int=str)
            for item in items:
                ref = item['source_ref']; index = ref.get('row')
                if isinstance(index, int) and 1 <= index <= len(records) and ref.get('column') in records[index-1]:
                    observed[item['id']] = records[index-1][ref['column']]
        elif source['format'] in {'.json','.txt','.log','.md'}:
            # These observations retain physical source-line numbers, including
            # blank lines. Do not reserialize JSON: that changes the coordinates.
            lines=path.read_text(encoding='utf-8',errors='replace').splitlines()
            for item in items:
                ref=item['source_ref'];match=re.fullmatch(r'line:([1-9][0-9]*)',str(ref.get('span','')))
                if item['extraction_method']=='text' and ref.get('page') is None and match:
                    index=int(match[1])-1
                    if index<len(lines):observed[item['id']]=lines[index]
        for item in items:
            value=observed.get(item['id']);text=str(value) if item['id'] in observed else None
            if item['id'].startswith('b1:') and isinstance(value,bool):text=text.lower()
            matches=text==item['value_text'] if text is not None else None
            if matches is False and isinstance(value,(date,datetime)):
                try:
                    parsed=datetime.fromisoformat(item['value_text']) if isinstance(value,datetime) else date.fromisoformat(item['value_text'])
                    matches=parsed==value
                except (ValueError,TypeError):pass
            if matches is False and item['value_number'] is not None:
                try:matches=float(text)==item['value_number']
                except (ValueError,TypeError):pass
            results[item['id']]={'observation_id':item['id'],'source_integrity':integrity,
                                  'value_matches':matches,'observed':text,'validation_status':item['validation_status'],
                                  'source_ref':item['source_ref'],'evidence_ids':[item['id']],
                                  'verification_status':'failed' if not integrity or matches is False else 'automatically_verified' if matches is True else 'human_review_pending',
                                  'verification_method':item['extraction_method'] if text is not None else None,
                                  'review_reason':None if text is not None else 'Original image requires visual review; OCR output is not independent confirmation.' if item['extraction_method'] in {'tesseract','docling','chandra','paddleocr-vl'} else 'Source location is not supported by the automatic verifier.'}
            if matches is None:
                # Read only exact source/text/location-bound human receipts.
                # The automatic value result remains unknown even after review.
                from .source_review import fingerprint
                binding={'observation_id':item['id'],'source_sha256':source['sha256'],
                         'source_ref':item['source_ref'],'extracted_value_sha256':fingerprint(item['value_text']),
                         'access_scope':source['access_scope']}
                receipts=[]
                for file in (store.root/'source-reviews'/fingerprint(item['id'])).glob('*.json'):
                    receipt=json.loads(file.read_text())
                    if receipt.get('packet_sha256')==fingerprint(binding):receipts.append(receipt)
                if receipts:
                    latest=max(receipts,key=lambda r:(r['recorded_at'],r['review_id']))
                    results[item['id']]['human_source_review']=latest
                    results[item['id']]['verification_status']='human_verified' if latest['decision']=='matches' else 'human_rejected' if latest['decision']=='does_not_match' else 'human_review_pending'
    return results
