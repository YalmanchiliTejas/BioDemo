from __future__ import annotations

import csv
import hashlib
import json
import re
import tempfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field

from .models import ToolRequest, ToolResponse
from .store import ManufacturingStore

class QueryRequest(ToolRequest):
    sql: str

class SearchRequest(ToolRequest):
    query: str = Field(max_length=4096)
    mode: Literal["lexical","semantic","hybrid"] = "lexical"
    evidence_type: Literal['auto','exact','structured','narrative'] = 'auto'
    observation_id: str | None = None
    source_id: str | None = None
    entity_id: str | None = None
    metric: str | None = None
    value_text: str | None = None
    sheet: str | None = None
    source_row: int | None = Field(default=None,ge=1)
    time_min: float | None = Field(default=None,allow_inf_nan=False)
    time_max: float | None = Field(default=None,allow_inf_nan=False)
    time_basis: Literal['elapsed_hours','calendar_hours','unresolved'] | None = None

class EvidenceRequest(ToolRequest):
    observation_id: str

class EntityRequest(ToolRequest):
    entity_id: str

class UnitRequest(ToolRequest):
    observation_ids: list[str]
    target_unit: str

class RegionRequest(ToolRequest):
    source_id: str
    sheet: str | None = None
    page: int | None = None
    row: int | None = None
    column: str | None = None
    region: tuple[float,float,float,float] | None = None
    backend: str = "auto"

class MetricRequest(ToolRequest):
    metrics: list[dict] = Field(min_length=1,max_length=32)

class AlignmentRequest(ToolRequest):
    left: dict
    right: dict
    method: Literal['exact','nearest','linear'] = 'exact'
    tolerance_hours: float = Field(default=0,ge=0,allow_inf_nan=False)
    allow_time_reordering: bool = False

class ManufacturingTools:
    """The main RLM calls these functions; this class does no investigation planning."""
    def __init__(self, store: ManufacturingStore, *, as_of: str | None = None,
                 scopes: tuple[str,...] = ("default",)):
        self.store = store
        self.as_of = as_of
        self.scopes = scopes
        self._analyses = {}
        self._retrieval = None

    def _visibility(self, arg: dict) -> dict:
        request = ToolRequest.model_validate({"as_of":self.as_of,"scopes":list(self.scopes),**arg})
        if self.as_of and request.as_of is None:
            raise PermissionError("cannot remove the session historical cutoff")
        if self.as_of and request.as_of:
            def instant(value: str) -> datetime:
                dt=datetime.fromisoformat(value.replace("Z","+00:00"))
                return dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt.astimezone(timezone.utc)
            if instant(request.as_of)>instant(self.as_of):
                raise PermissionError("cannot advance the session's historical cutoff")
        if not set(request.scopes).issubset(self.scopes):
            raise PermissionError("requested access scope is not granted")
        return request.model_dump()

    def _sources(self, arg: dict) -> list[dict]:
        v = self._visibility(arg)
        return self.store.visible_sources(as_of=v["as_of"], scopes=v["scopes"])

    def _obs(self, arg: dict, *, scan_all: bool = False, **filters) -> list[dict]:
        v = self._visibility(arg)
        return self.store.visible(as_of=v["as_of"],scopes=v["scopes"],limit=None if scan_all else v["limit"],
                                  offset=0 if scan_all else v["offset"],**filters)

    def inspect_data(self, **kwargs) -> dict:
        sources = self._sources(kwargs)
        visibility = self._visibility(kwargs)
        summary = self.store.summarize(as_of=visibility['as_of'],scopes=visibility['scopes'])
        by_source = Counter()
        variables = Counter()
        for item in summary:
            by_source[item['source_id']] += item['count']
            variables[item['name']] += item['count']
        by_schema=defaultdict(lambda:defaultdict(set))
        with self.store.connect() as con:
            for source in sources:
                for col in con.execute("SELECT sheet,name FROM source_columns WHERE source_id=? ORDER BY column_index",(source['id'],)):
                    by_schema[source['id']][col['sheet'] or '<file>'].add(col['name'])
        ids = self._sample_ids(visibility,summary)
        return {'sources':[{'id':s['id'],'path':s['path'],'sha256':s['sha256'],
                            'format':s['format'],'status':s['extraction_status'],
                            'observations':by_source[s['id']],
                            'sheets':{k:sorted(v) for k,v in by_schema[s['id']].items()}} for s in sources],
                'variables':dict(variables),'evidence_ids':ids,'evidence_id_count':sum(by_source.values())}

    def _sample_ids(self, visibility: dict, summary=None) -> list[str]:
        if summary is None:
            summary=self.store.summarize(as_of=visibility['as_of'],scopes=visibility['scopes'])
        # Round-robin across sources and series before taking additional anchors.
        # Includes ordinary endpoints as well as extrema; no anomaly-only filter.
        groups=defaultdict(list)
        for row in summary:
            groups[row['source_id']].append(row)
        ordered=[]
        for index in range(max((len(v) for v in groups.values()),default=0)):
            ordered.extend(v[index] for _,v in sorted(groups.items()) if index<len(v))
        ids=[]
        keys=('first_id','last_id','minimum_id','maximum_id','sample_id')
        for index,row in enumerate(ordered):
            anchors=list(dict.fromkeys(row[key] for key in keys if row.get(key)))
            if anchors:ids.append(anchors[index%len(anchors)])
        for key in keys:
            ids.extend(row[key] for row in ordered if row.get(key))
        return list(dict.fromkeys(ids))[:1000]

    def infer_schema(self, **kwargs) -> dict:
        v = self._visibility(kwargs)
        summary = self.store.summarize(as_of=v['as_of'],scopes=v['scopes'])
        grouped = defaultdict(list)
        for row in summary:
            grouped[row['name']].append(row)
        with self.store.connect() as con:
            for source in self._sources(kwargs):
                for row in con.execute('SELECT name FROM source_columns WHERE source_id=?',(source['id'],)):
                    grouped[row[0]]
        mappings=[]
        for name, rows in sorted(grouped.items()):
            norm = re.sub(r'[^a-z0-9]+','_',name.lower()).strip('_')
            role = ('time' if any(x in norm for x in ('timestamp','time','date')) else
                    'entity_identifier' if any(x in norm for x in ('batch','lot','run','reactor','equipment','sample')) else
                    'measurement' if any(o['numeric_count'] for o in rows) else 'text')
            mappings.append({'source_name':name,'suggested_role':role,
                             'units':sorted({o['unit'] for o in rows if o['unit']}),
                             'status':'proposed','evidence_ids':[o['sample_id'] for o in rows][:5]})
        from .terminology import canonical
        for item in mappings:
            item['terminology']=[canonical(item['source_name'],unit) for unit in item['units'] or [None]]
        return {'mappings':mappings,'evidence_ids':self._sample_ids(v,summary),
                'evidence_id_count':sum(r['count'] for r in summary)}

    def reconcile_entities(self, **kwargs) -> dict:
        v = self._visibility(kwargs)
        rows = self.store.select('SELECT entity_id,list(DISTINCT source_id) AS source_ids,count(*) AS observation_count,'
                                 'min(id) AS sample_id FROM observations GROUP BY entity_id',
                                 as_of=v['as_of'],scopes=v['scopes'])
        return {'candidates':[{'entity_id':r['entity_id'],'source_ids':sorted(r['source_ids']),
                                'observation_count':r['observation_count'],'status':'candidate',
                                'evidence_ids':[r['sample_id']]} for r in rows if r['entity_id']],
                'unlinked_observations':sum(r['observation_count'] for r in rows if not r['entity_id']),
                'evidence_ids':self._sample_ids(v),
                'evidence_id_count':sum(r['observation_count'] for r in rows)}

    def normalize_measurements(self, **kwargs) -> dict:
        req = UnitRequest.model_validate({"as_of":self.as_of,"scopes":list(self.scopes),**kwargs})
        self._visibility(req.model_dump())
        import pint
        import polars as pl
        import pandera.polars as pa
        registry = pint.UnitRegistry()
        wanted=set(req.observation_ids)
        rows=self.store.get_many(sorted(wanted),as_of=req.as_of,scopes=req.scopes)
        output=[]
        for o in rows:
            item={"observation_id":o["id"],"original_value":o["value_number"],"original_unit":o["unit"]}
            try:
                if o["value_number"] is None or not o["unit"]: raise ValueError("numeric value or unit missing")
                value=registry.Quantity(o["value_number"],o["unit"]).to(req.target_unit).magnitude
                frame=pl.DataFrame({"value_number":[float(value)]})
                pa.DataFrameSchema({"value_number":pa.Column(float,nullable=False)}).validate(frame)
                item.update(value=float(value),unit=req.target_unit,status="normalized")
            except Exception as exc:
                item.update(status="unresolved",reason=str(exc))
            output.append(item)
        return {"measurements":output,"evidence_ids":[o["id"] for o in rows]}

    def query_data(self, **kwargs) -> dict:
        req=QueryRequest.model_validate({"as_of":self.as_of,"scopes":list(self.scopes),**kwargs})
        self._visibility(req.model_dump())
        sql=req.sql.strip()
        if not re.match(r"(?is)^select\s",sql) or ";" in sql or re.search(r"(?i)\b(read_csv|read_parquet|read_json|copy|attach|install|load|pragma|create|delete|insert|update|drop|alter|call|export|import)\b",sql):
            raise ValueError("only a single SELECT over the observations table is allowed")
        if not re.search(r"(?i)\bfrom\s+observations\b", sql):
            raise ValueError("unknown table; query the observations table, e.g. "
                             "SELECT id, name, value_number, unit FROM observations WHERE name = 'pH'")
        import duckdb
        with duckdb.connect(':memory:',config={'enable_external_access':False}) as parser:
            if parser.get_table_names(sql) != {'observations'} or re.search(r'(?i)\b(query|query_table|sqlite_scan|parquet_scan)\s*\(',sql):
                raise ValueError('query only the scoped observations table')
        with self.store.query_connection(as_of=req.as_of,scopes=req.scopes,numeric_only=True) as con:
            cursor=con.execute('SELECT * FROM ('+sql+') LIMIT ? OFFSET ?',[req.limit,req.offset])
            columns=[x[0] for x in cursor.description]
            values=[list(row) for row in cursor.fetchall()]
            citations=list(dict.fromkeys(str(row[columns.index(key)]) for row in values for key in ('id','time_evidence_id') if key in columns and row[columns.index(key)] is not None))
        source_ids={s['id'] for s in self._sources({'as_of':req.as_of,'scopes':req.scopes}) if s['extraction_status']=='indexed'}
        with self.store.connect() as metadata:
            shards=sum(metadata.execute('SELECT count(*) FROM sheets WHERE source_id=?',(sid,)).fetchone()[0] for sid in source_ids)
        return {'columns':columns,'rows':values,'evidence_ids':citations,
                'provenance_status':'row_ids' if citations else 'aggregate_lineage_unresolved',
                'source_count':len(source_ids),'parquet_shards':shards}

    def search_evidence(self, **kwargs) -> dict:
        req=SearchRequest.model_validate({"as_of":self.as_of,"scopes":list(self.scopes),**kwargs})
        self._visibility(req.model_dump())
        if req.time_min is not None and req.time_max is not None and req.time_min>req.time_max:
            raise ValueError('time_min must be <= time_max')
        from .retrieval import RetrievalEngine
        if self._retrieval is None:
            self._retrieval=RetrievalEngine(self.store)
        return self._retrieval.search(req)

    def verify_evidence(self, **kwargs) -> dict:
        req=EvidenceRequest.model_validate({"as_of":self.as_of,"scopes":list(self.scopes),**kwargs})
        matches=self._obs({"as_of":req.as_of,"scopes":req.scopes},observation_id=req.observation_id)
        if not matches: raise KeyError("evidence not visible")
        o=matches[0]
        return self.verify_many(matches)[req.observation_id]

    def verify_many(self, rows: list[dict]) -> dict:
        from .verification import verify_originals
        return verify_originals(self.store,rows)

    @staticmethod
    def _source_row(path: Path, sheet_name: str | None, row_number: int | None):
        if row_number is None:
            return None
        if path.suffix.lower() in {'.xlsx','.xlsm'}:
            from openpyxl import load_workbook
            workbook=load_workbook(path,read_only=True,data_only=True)
            try:
                sheets=[workbook[sheet_name]] if sheet_name in workbook.sheetnames else ([] if sheet_name else list(workbook))
                for sheet in sheets:
                    if row_number < 2 or (sheet.max_row is not None and row_number>sheet.max_row):
                        continue
                    header=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())
                    headers=[str(v) if v is not None else f'column_{i}' for i,v in enumerate(header,1)]
                    values=next(sheet.iter_rows(min_row=row_number,max_row=row_number,values_only=True),None)
                    if values is not None:
                        return sheet.title,dict(zip(headers,values))
            finally:
                workbook.close()
        else:
            from .ingest import _rows
            for sheet,row,headers,values in _rows(path):
                if sheet==sheet_name and row==row_number:
                    return sheet,dict(zip(headers,values))
        return None

    def extract_region(self, **kwargs) -> dict:
        req=RegionRequest.model_validate({"as_of":self.as_of,"scopes":list(self.scopes),**kwargs})
        source=next((s for s in self._sources(req.model_dump()) if s["id"]==req.source_id),None)
        if source is None: raise KeyError("source not visible")
        path=self.store.originals/source["sha256"][:2]/(source["sha256"]+Path(source["path"]).suffix.lower())
        if source["format"] in {".csv",".tsv",".xlsx",".xlsm",".parquet"}:
            found=self._source_row(path,req.sheet,req.row)
            if not found:
                return {'status':'region_not_found'}
            sheet,data=found
            # A visible source can also contain future rows or fields. Apply the
            # same cell visibility rule before returning an original region.
            visible=self.store.select(
                "SELECT name FROM observations WHERE source_id=? "
                "AND TRY_CAST(json_extract_string(source_ref,'$.row') AS BIGINT)=? "
                "AND json_extract_string(source_ref,'$.sheet') IS NOT DISTINCT FROM ?",
                [req.source_id,req.row,sheet],as_of=req.as_of,scopes=req.scopes)
            allowed={row['name'] for row in visible}
            data={name:value for name,value in data.items() if name in allowed}
            if req.column:
                data={req.column:data[req.column]} if req.column in data else {}
            if not data:
                return {'status':'region_not_found'}
            return {'source_id':req.source_id,'sheet':sheet,'row':req.row,'data':data,'status':'source_transcribed'}
        if source["format"]==".pdf":
            from pypdf import PdfReader
            pages=PdfReader(str(path)).pages
            if req.page is None or req.page<1 or req.page>len(pages): raise ValueError("valid page is required")
            text=pages[req.page-1].extract_text() or ""
            if text and req.region is None and req.backend=="auto":
                return {"source_id":req.source_id,"page":req.page,"text":text[:20000],
                        "status":"source_transcribed"}
            import pypdfium2
            image=pypdfium2.PdfDocument(str(path))[req.page-1].render(scale=2).to_pil()
        elif source["format"] in {".png",".jpg",".jpeg",".tif",".tiff"}:
            from PIL import Image
            image=Image.open(path)
        else:
            return {"source_id":req.source_id,"status":"unsupported_region_format"}
        if req.region is not None:
            image=image.crop(tuple(int(v) for v in req.region))
        from .ingest import _document_observations, _ocr_image
        if req.backend in {"auto","tesseract"}:
            text=_ocr_image(image)
            method="tesseract"
        elif req.backend in {"docling","chandra","paddleocr-vl"}:
            with tempfile.TemporaryDirectory(prefix="manufacturing-region-") as directory:
                crop=Path(directory)/"region.png"
                image.save(crop)
                text="\n".join(o.value_text or "" for o in _document_observations(
                    req.source_id,source["path"],crop,source["access_scope"],source["available_at"],req.backend))
            method=req.backend
        else:
            raise ValueError("unknown region extraction backend")
        return {"source_id":req.source_id,"page":req.page,"region":req.region,
                "text":text[:20000],"method":method,"status":"unverified" if text else "unresolved"}

    def get_process_context(self, **kwargs) -> dict:
        req=EntityRequest.model_validate({"as_of":self.as_of,"scopes":list(self.scopes),**kwargs})
        records=self._obs({"as_of":req.as_of,"scopes":req.scopes,"limit":req.limit},entity_id=req.entity_id)
        permitted={s["id"] for s in self._sources(req.model_dump()) if s["extraction_status"]=="indexed"}
        with self.store.connect() as con:
            if req.as_of:
                cursor=con.execute("SELECT * FROM relations WHERE (left_id=? OR right_id=?) AND "
                                   "(available_at IS NULL OR julianday(available_at)<=julianday(?))",
                                   (req.entity_id,req.entity_id,req.as_of))
            else:
                cursor=con.execute("SELECT * FROM relations WHERE left_id=? OR right_id=?",
                                   (req.entity_id,req.entity_id))
            links=[dict(row) for row in cursor if row["source_id"] in permitted]
        return {"entity_id":req.entity_id,"records":records,"source_ids":sorted({o["source_id"] for o in records}),
                "related_entities":sorted({x["right_id"] if x["left_id"]==req.entity_id else x["left_id"] for x in links}),
                "relations":links,"evidence_ids":[o["id"] for o in records]}

    def get_state_at(self, **kwargs) -> dict:
        req=ToolRequest.model_validate({"as_of":self.as_of,"scopes":list(self.scopes),**kwargs})
        if not req.as_of: raise ValueError("as_of required")
        records=self._obs(req.model_dump())
        structured=[s for s in self._sources(req.model_dump()) if s["format"] in
                    {".csv",".tsv",".xlsx",".xlsm",".parquet"}]
        warnings=[]
        with self.store.connect() as con:
            for source in structured:
                fields={row[0].lower() for row in con.execute("SELECT name FROM source_columns WHERE source_id=?",(source["id"],))}
                if "field_available_at" not in fields:
                    warnings.append(f"{source['path']}: field-level publication times absent; row cutoff cannot certify mutable fields")
        return {"as_of":req.as_of,"records":records,"evidence_ids":[o["id"] for o in records],
                "warnings":warnings,"temporal_integrity":"row_level_only" if warnings else "field_level_if_supplied"}

    def analysis(self, **kwargs):
        from .analysis import DeterministicAnalysis
        v=self._visibility(kwargs)
        key=(v['as_of'],tuple(sorted(v['scopes'])))
        if key not in self._analyses:
            self._analyses[key]=DeterministicAnalysis(self.store,as_of=v['as_of'],scopes=key[1])
        return self._analyses[key]

    def summarize_series(self, **kwargs) -> dict:
        v=self._visibility(kwargs)
        prepared=self.analysis(**kwargs).prepare()
        rows=prepared['series']
        for key in ('source_id','entity_id','name','sheet'):
            if key in kwargs: rows=[r for r in rows if r[key]==kwargs[key]]
        selected=rows[v['offset']:v['offset']+v['limit']]
        anomalies=[a for a in prepared['anomalies'] if all(key not in kwargs or a.get(key)==kwargs[key] for key in ('source_id','entity_id','name','sheet'))]
        return {'series':selected,'series_count':len(rows),'anomalies':anomalies[:v['limit']],
                'limitations':prepared['limitations'],
                'evidence_ids':list(dict.fromkeys(i for r in selected for i in r['evidence_ids']))}

    def calculate_metrics(self, **kwargs) -> dict:
        req=MetricRequest.model_validate({'as_of':self.as_of,'scopes':list(self.scopes),**kwargs})
        analysis=self.analysis(as_of=req.as_of,scopes=req.scopes)
        metrics=[analysis.calculate(spec) for spec in req.metrics]
        return {'metrics':metrics,'evidence_ids':list(dict.fromkeys(i for m in metrics for i in m['evidence_ids']))}

    def align_series(self, **kwargs) -> dict:
        req=AlignmentRequest.model_validate({'as_of':self.as_of,'scopes':list(self.scopes),**kwargs})
        return self.analysis(as_of=req.as_of,scopes=req.scopes).align(req.left,req.right,method=req.method,
                tolerance_hours=req.tolerance_hours,limit=req.limit,offset=req.offset,allow_time_reordering=req.allow_time_reordering)

    def handlers(self) -> dict[str,Any]:
        def typed(method):
            def call(**kwargs):
                return ToolResponse.model_validate(method(**kwargs)).model_dump(exclude_unset=True)
            return call
        return {name:typed(getattr(self,name)) for name in ("inspect_data","infer_schema","reconcile_entities",
                "normalize_measurements","query_data","search_evidence","verify_evidence",
                "extract_region","get_process_context","get_state_at","summarize_series","calculate_metrics","align_series")}
