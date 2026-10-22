from __future__ import annotations

import csv
import os
import time
from concurrent.futures import ThreadPoolExecutor
import hashlib
import importlib.util
import io
import json
import mimetypes
import re
import shutil
import stat
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from itertools import islice
from pathlib import Path
from typing import Any, Iterator

from .models import Observation, SourceRef
from .store import ManufacturingStore

SUPPORTED = {".csv", ".tsv", ".xlsx", ".xlsm", ".parquet", ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".txt", ".log", ".md", ".json"}
PIPELINE_VERSION = "0.6.0-verified-columnar"
TIME_WORDS = ("time", "timestamp", "date", "recorded", "collected", "sampled", "occurred")
AVAIL_WORDS = ("available", "released", "approved", "reviewed", "reported")
ENTITY_WORDS = {"batch": "batch", "lot": "batch", "run": "run", "reactor": "equipment", "equipment": "equipment", "sample": "sample", "site": "site", "campaign": "campaign", "process": "process"}
STRUCTURED = {".csv", ".tsv", ".xlsx", ".xlsm", ".parquet"}
PLAIN_TEXT = {".txt", ".log", ".md", ".json"}
IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}


def _routing_capabilities() -> dict[str, bool]:
    """Invalidate the incremental cache when an optional backend becomes available."""
    def installed(name: str) -> bool:
        try:
            return importlib.util.find_spec(name) is not None
        except (ImportError, ValueError):
            return False
    return {"docling": installed("docling"), "chandra_cli": bool(shutil.which("docling")),
            "paddleocr_vl": installed("paddleocr"), "tesseract": bool(shutil.which("tesseract"))}


def _route_file(rel: str, path: Path, requested: str) -> dict:
    suffix = path.suffix.lower()
    if suffix in STRUCTURED:
        return {"kind": "structured", "reason": "tabular file", "candidates": ["structured"]}
    if suffix in PLAIN_TEXT:
        return {"kind": "text", "reason": "plain text or log", "candidates": ["text"]}
    if suffix not in SUPPORTED:
        return {"kind": "unsupported", "reason": "unrecognized extension", "candidates": ["preserve"]}
    if requested != "auto":
        return {"kind": "document", "reason": "explicit document backend", "candidates": [requested]}
    if suffix in IMAGES:
        return {"kind": "image", "reason": "image or scanned source; handwriting may be present",
                "candidates": ["chandra", "paddleocr-vl", "pypdf"]}
    # A PDF with little extractable text or a detected form benefits from a VLM.
    # Filename markers are weak hints; the selection remains auditable.
    label = re.sub(r"[^a-z0-9]+", " ", rel.lower())
    name_hint = bool(re.search(r"handwrit|\bscanned?\b|\bforms?\b", label))
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        form_fields = bool(reader.get_fields())
        text_lengths = [len((page.extract_text() or "").strip()) for page in reader.pages]
        scan_like = not text_lengths or any(length < 40 for length in text_lengths)
    except Exception:
        form_fields, scan_like = False, False
        unreadable = True
    else:
        unreadable = False
    if name_hint or form_fields or scan_like:
        reason = ("handwriting/scan filename hint" if name_hint else
                  "PDF form fields detected" if form_fields else "page has little extractable text")
        return {"kind": "scan_or_form_pdf", "reason": reason,
                "candidates": ["chandra", "paddleocr-vl", "docling", "pypdf"]}
    return {"kind": "text_pdf", "reason": "PDF has extractable text" if not unreadable else
            "PDF text inspection failed; document parser will retry",
            "candidates": ["docling", "pypdf"]}

def _iso(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).isoformat() if value.tzinfo else value.isoformat()
    s = str(value).strip()
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return dt.astimezone(timezone.utc).isoformat() if dt.tzinfo else dt.isoformat()
    except ValueError:
        return None

def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None or value == "":
        return None
    try:
        n = float(value)
        return n if n == n and abs(n) != float("inf") else None
    except (ValueError, TypeError):
        return None

def _unit(header: str) -> str | None:
    found = re.findall(r"\(([^()]*)\)|\[([^][]*)\]", header)
    return next((a or b for a,b in found[-1:]), None)

def _key(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", header.lower()).strip("_")

def _rows(path: Path, max_rows: int | None = None) -> Iterator[tuple[str | None, int, list[str], list[Any]]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        with path.open(newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f, delimiter="\t" if suffix == ".tsv" else ",")
            headers = [str(x) for x in next(reader, [])]
            for i, values in enumerate(islice(reader,max_rows) if max_rows is not None else reader, 2):
                yield None, i, headers, values
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in wb:
                it = sheet.iter_rows(values_only=True)
                headers = [str(x) if x is not None else f"column_{i}" for i,x in enumerate(next(it, []), 1)]
                for i, values in enumerate(islice(it,max_rows) if max_rows is not None else it, 2):
                    yield sheet.title, i, headers, list(values)
        finally:
            wb.close()
    elif suffix == ".parquet":
        import polars as pl
        frame = pl.scan_parquet(path)
        headers = frame.collect_schema().names()
        if max_rows is not None:
            frame = frame.head(max_rows)
        for i, values in enumerate(frame.collect(engine="streaming").iter_rows(), 1):
            yield None, i, headers, list(values)

def _headers(path: Path) -> Iterator[tuple[str | None,list[str]]]:
    suffix=path.suffix.lower()
    if suffix in {".csv",".tsv"}:
        with path.open(newline="",encoding="utf-8-sig",errors="replace") as handle:
            yield None,[str(x) for x in next(csv.reader(handle,delimiter="\t" if suffix==".tsv" else ","),[])]
    elif suffix in {".xlsx",".xlsm"}:
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=True)
        try:
            for sheet in wb:
                head=next(sheet.iter_rows(max_row=1,values_only=True),())
                yield sheet.title,[str(x) if x is not None else f"column_{i}" for i,x in enumerate(head,1)]
        finally: wb.close()
    elif suffix==".parquet":
        import polars as pl
        yield None,pl.scan_parquet(path).collect_schema().names()

def _entity_from_context(rel: str, sheet: str | None, record: dict[str, Any]) -> tuple[str | None, str | None]:
    for header, value in record.items():
        key = _key(header)
        if value is None or str(value).strip() == "":
            continue
        for marker, kind in ENTITY_WORDS.items():
            if marker in key and (key.endswith("id") or key == marker or key.endswith("number")):
                return kind, str(value).strip()
    # Generic source identifiers are candidates only; they are never certified links.
    file_match = re.search(r"(?<![A-Za-z0-9])([A-Za-z]{1,6}[-_ ]?\d{1,8})(?![A-Za-z0-9])", Path(rel).stem)
    sheet_match = re.search(r"(?<![A-Za-z0-9])([A-Za-z]{1,6}[-_ ]?\d{1,8})(?![A-Za-z0-9])", sheet or "")
    if file_match and sheet_match:
        left = re.sub(r"[ _]", "", file_match.group(1)).upper()
        right = re.sub(r"[ _]", "", sheet_match.group(1)).upper()
        if left != right:
            return "run", f"{left}/{right}"
    for label in (sheet or "", Path(rel).stem):
        match = re.search(r"(?<![A-Za-z0-9])([A-Za-z]{1,6}[-_ ]?\d{1,8})(?![A-Za-z0-9])", label)
        if match:
            token = re.sub(r"[ _]", "", match.group(1)).upper()
            return ("run" if label == sheet else "batch"), token
    return None, None

def _related_entities(rel: str, sheet: str | None, record: dict[str,Any]) -> list[tuple[str,str]]:
    found=[]
    for header,value in record.items():
        key=_key(header)
        if value is None or str(value).strip()=="": continue
        for marker,kind in ENTITY_WORDS.items():
            if marker in key and (key.endswith("id") or key==marker or key.endswith("number")):
                found.append((kind,str(value).strip()))
                break
    file_match=re.search(r"(?<![A-Za-z0-9])([A-Za-z]{1,6}[-_ ]?\d{1,8})(?![A-Za-z0-9])",Path(rel).stem)
    sheet_match=re.search(r"(?<![A-Za-z0-9])([A-Za-z]{1,6}[-_ ]?\d{1,8})(?![A-Za-z0-9])",sheet or "")
    if file_match and sheet_match:
        found.append(("batch",re.sub(r"[ _]","",file_match.group(1)).upper()))
        found.append(("run",re.sub(r"[ _]","",file_match.group(1)).upper()+"/"+
                      re.sub(r"[ _]","",sheet_match.group(1)).upper()))
    return list(dict.fromkeys(found))

def _make_observations(source_id: str, rel: str, path: Path, scope: str, source_available: str | None,
                       *, max_rows: int | None = None) -> Iterator[Observation]:
    for sheet, row_index, headers, values in _rows(path,max_rows=max_rows):
        record = dict(zip(headers, values))
        field_times: dict[str,str] = {}
        for key,value in record.items():
            if _key(key) == "field_available_at" and isinstance(value,str):
                try:
                    parsed = json.loads(value)
                    if isinstance(parsed,dict): field_times = parsed
                except json.JSONDecodeError:
                    pass
        kind, entity_id = _entity_from_context(rel, sheet, record)
        related_entities = _related_entities(rel,sheet,record)
        occurred = available = None
        for header, value in record.items():
            key = _key(header)
            if any(word in key for word in AVAIL_WORDS) and _iso(value):
                available = _iso(value)
            elif any(word in key for word in TIME_WORDS) and _iso(value):
                occurred = _iso(value)
        available = available or source_available
        for col_index, (header, value) in enumerate(zip(headers, values), 1):
            if value is None or value == "":
                continue
            from openpyxl.utils import get_column_letter
            ref = SourceRef(source_id=source_id, path=rel, sheet=sheet, row=row_index,
                            column=header, cell=f"{get_column_letter(col_index)}{row_index}" if sheet else None)
            ident = hashlib.sha256(f"{source_id}|{sheet}|{row_index}|{col_index}".encode()).hexdigest()[:24]
            number = _number(value)
            field_available = _iso(field_times.get(header))
            effective_available = max(filter(None,(available,field_available)), default=None)
            yield Observation(observation_id=ident, source=ref, name=header,
                              value_text=str(value), value_number=number, unit=_unit(header),
                              occurred_at=occurred, available_at=effective_available, entity_id=entity_id, entity_kind=kind,
                              related_entities=related_entities,
                              access_scope=scope, extraction_method=path.suffix.lower().lstrip("."),
                              confidence=1.0 if kind and any(k in _key(header) for k in ENTITY_WORDS) else None)

def _document_observations(source_id: str, rel: str, path: Path, scope: str,
                           available: str | None, backend: str) -> Iterator[Observation]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".log", ".md", ".json"}:
        pages = [(None, path.read_text(encoding="utf-8", errors="replace"), "text")]
    elif backend == "docling":
        from docling.document_converter import DocumentConverter
        artifacts = os.getenv('DOCLING_ARTIFACTS_PATH')
        if artifacts:
            from docling.document_converter import PdfFormatOption
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import PdfPipelineOptions, TesseractCliOcrOptions
            options = PdfPipelineOptions(artifacts_path=Path(artifacts), enable_remote_services=False,
                                         ocr_options=TesseractCliOcrOptions(lang=['eng']))
            converter = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)})
        else:
            converter = DocumentConverter()
        doc = converter.convert(str(path)).document
        # Retain Docling's physical provenance instead of flattening the entire
        # document into page-less Markdown. OCR remains unverified.
        if hasattr(doc, 'iterate_items'):
            for item_index, (item, _level) in enumerate(doc.iterate_items(), 1):
                body = getattr(item, 'text', None)
                if not body and hasattr(item, 'export_to_markdown'):
                    body = item.export_to_markdown(doc=doc)
                if not body or not body.strip():
                    continue
                provenance = list(getattr(item, 'prov', []) or [])
                # Multi-page table text cannot be attributed wholesale to its
                # first page. Preserve the text but require source review.
                localized = provenance[0] if len(provenance) == 1 else None
                page = getattr(localized, 'page_no', None)
                box = getattr(localized, 'bbox', None)
                region = (float(box.l), float(box.t), float(box.r), float(box.b)) if box else None
                ref = SourceRef(source_id=source_id, path=rel, page=page,
                                span=f'docling:item:{item_index}', region=region)
                ident = hashlib.sha256(f'{source_id}|docling|{item_index}'.encode()).hexdigest()[:24]
                yield Observation(observation_id=ident, source=ref, name='document_text',
                                  value_text=body, available_at=available, access_scope=scope,
                                  extraction_method='docling', validation_status='unverified')
            return
        pages = [(None, doc.export_to_markdown(), "docling")]
    elif backend == "chandra":
        if not shutil.which("docling"):
            raise RuntimeError("Docling CLI is required for Chandra VLM extraction")
        with tempfile.TemporaryDirectory(prefix="manufacturing-chandra-") as directory:
            command=["docling","convert",str(path),"--pipeline","vlm","--vlm-model",
                     "chandra_ocr2","--to","md","--output",directory]
            subprocess.run(command,check=True,capture_output=True,text=True,timeout=900)
            output=Path(directory)/(path.stem+".md")
            if not output.exists():
                raise RuntimeError("Chandra VLM completed without Markdown output")
            pages=[(None,output.read_text(),"chandra")]
    elif backend == "paddleocr-vl":
        try:
            from paddleocr import PaddleOCRVL
        except ImportError as exc:
            raise RuntimeError("Install paddleocr and its inference dependencies for PaddleOCR-VL") from exc
        with tempfile.TemporaryDirectory(prefix="manufacturing-paddleocr-") as directory:
            pipeline=PaddleOCRVL()
            for result in pipeline.predict(input=str(path)):
                result.save_to_markdown(save_path=directory)
            pages=[(i,text.read_text(),"paddleocr-vl") for i,text in
                   enumerate(sorted(Path(directory).glob("*.md")),1)]
    elif suffix == ".pdf":
        from pypdf import PdfReader
        pages = [(i, page.extract_text() or "", "pypdf") for i,page in enumerate(PdfReader(str(path)).pages, 1)]
        if any(not body.strip() for _,body,_ in pages):
            import pypdfium2
            pdf = pypdfium2.PdfDocument(str(path))
            pages = [(i, body if body.strip() else _ocr_image(pdf[i-1].render(scale=2).to_pil()),
                      method if body.strip() else "tesseract") for i,body,method in pages]
    elif suffix in {".png",".jpg",".jpeg",".tif",".tiff"}:
        from PIL import Image
        pages = [(1,_ocr_image(Image.open(path)),"tesseract")]
    else:
        pages = []
    for page, body, method in pages:
        if not body.strip():
            continue
        for line_no, line in enumerate(body.splitlines(), 1):
            if not line.strip():
                continue
            ident = hashlib.sha256(f"{source_id}|{page}|{line_no}".encode()).hexdigest()[:24]
            ref = SourceRef(source_id=source_id, path=rel, page=page, span=f"line:{line_no}")
            yield Observation(observation_id=ident, source=ref, name="document_text", value_text=line,
                              available_at=available, access_scope=scope, extraction_method=method,
                              validation_status="unverified" if method in {"docling","tesseract","chandra","paddleocr-vl"} else "source_transcribed")

def _ocr_image(image) -> str:
    if not shutil.which("tesseract"):
        return ""
    with tempfile.NamedTemporaryFile(suffix=".png") as tmp:
        image.save(tmp.name)
        process=subprocess.run(["tesseract",tmp.name,"stdout"],capture_output=True,text=True,check=False,timeout=120)
        return process.stdout if process.returncode==0 else ""

def _members(package: Path) -> Iterator[tuple[str, bytes, str | None, int]]:
    if package.is_dir():
        for path in sorted(package.rglob("*")):
            if path.is_symlink():
                raise ValueError(f"symlink in package: {path}")
            if path.is_file():
                relative = path.relative_to(package).as_posix()
                yield relative, path.read_bytes(), datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(), stat.S_IMODE(path.stat().st_mode)
    elif package.suffix.lower() == ".zip":
        with zipfile.ZipFile(package) as archive:
            for item in archive.infolist():
                posix = Path(item.filename)
                if item.is_dir():
                    continue
                if posix.is_absolute() or ".." in posix.parts or item.external_attr >> 16 & 0o170000 == 0o120000:
                    raise ValueError(f"unsafe ZIP member: {item.filename}")
                if item.file_size > 2_000_000_000:
                    raise ValueError("ZIP member exceeds 2 GB safety limit")
                yield item.filename, archive.read(item), _iso(datetime(*item.date_time)), (item.external_attr >> 16) & 0o777
    else:
        raise ValueError("input must be a folder or ZIP")

def ingest_package(package: str | Path, store: ManufacturingStore, *, access_scope: str = "default",
                   document_backend: str = "auto", max_rows_per_file: int | None = None,
                   workers: int | None = None) -> dict:
    """Copy immutable source bytes and incrementally index changed files. Never alters input."""
    package = Path(package).resolve()
    if document_backend not in {"auto", "docling", "pypdf", "chandra", "paddleocr-vl"}:
        raise ValueError("document_backend must be auto, pypdf, docling, chandra, or paddleocr-vl")
    if max_rows_per_file is not None and max_rows_per_file < 1:
        raise ValueError("max_rows_per_file must be positive")
    workers = workers if workers is not None else min(4, os.cpu_count() or 1)
    if workers < 1 or workers > 32:
        raise ValueError("workers must be between 1 and 32")
    started = time.monotonic()
    capabilities = _routing_capabilities() if document_backend == "auto" else {}
    result = {"files_seen": 0, "files_indexed": 0, "files_preserved_unsupported":0,
              "observations": 0, "unchanged": 0, "errors": [], "routes": [], "workers": workers}
    from .columnar import prepare_source
    pending = []

    def finish(job):
        sid, rel, route, future = job
        try:
            prepared = future.result()
            store.commit_columnar(sid, prepared)
            route.update(selected_backend="structured", reader="fastexcel-polars", storage="parquet_rows",
                         extraction_methods=[Path(rel).suffix.lower().lstrip('.')],
                         ingestion_seconds=prepared['seconds'], sheets=len(prepared['sheets']),
                         rows=sum(s['row_count'] for s in prepared['sheets']))
            with store.connect() as con:
                con.execute("UPDATE sources SET extraction_status='indexed',routing_json=? WHERE id=?",
                            (json.dumps(route,sort_keys=True),sid))
            result['files_indexed'] += 1
            result['observations'] += prepared['observations']
        except Exception as exc:
            with store.connect() as con:
                con.execute("UPDATE sources SET extraction_status='error',error=? WHERE id=?",(str(exc),sid))
            result['errors'].append({'path':rel,'error':str(exc)})
        result['routes'].append({'path':rel,**route})

    with ThreadPoolExecutor(max_workers=workers,thread_name_prefix='workbook') as executor:
        for rel, content, modified, permissions in _members(package):
            result["files_seen"] += 1
            safe_availability=(datetime.fromtimestamp(package.stat().st_mtime,timezone.utc).isoformat()
                               if package.suffix.lower()==".zip" else modified)
            digest = store.digest(content)
            sid = hashlib.sha256(f"{package}|{rel}|{access_scope}".encode()).hexdigest()[:24]
            parse_path = store.originals / digest[:2] / (digest + Path(rel).suffix.lower())
            parse_path.parent.mkdir(exist_ok=True)
            if not parse_path.exists():
                parse_path.write_bytes(content)
                parse_path.chmod(0o400)
            route = _route_file(rel, parse_path, document_backend)
            cache_signature = (f"{PIPELINE_VERSION}|requested={document_backend}|"
                               f"route={route['kind']}|capabilities={json.dumps(capabilities,sort_keys=True)}|"
                               f"rows={max_rows_per_file if max_rows_per_file is not None else 'all'}")
            with store.connect() as con:
                prior = con.execute("SELECT sha256,extraction_status,pipeline_version,routing_json FROM sources WHERE id=?",
                                    (sid,)).fetchone()
            with store.connect() as con:
                partitions_exist = all(Path(row[0]).is_file() and row[1] for row in con.execute(
                    "SELECT parquet_path,parquet_sha256 FROM sheets WHERE source_id=?",(sid,)))
            if partitions_exist and prior and prior["sha256"] == digest and prior["extraction_status"] in {"indexed","preserved_unsupported"} and prior["pipeline_version"] == cache_signature:
                result["unchanged"] += 1
                result["routes"].append({"path": rel, **(json.loads(prior["routing_json"]) if prior["routing_json"] else route),
                                          "unchanged": True})
                continue
            with store.connect() as con:
                con.execute("DELETE FROM observations WHERE source_id=?", (sid,))
                con.execute("DELETE FROM records WHERE source_id=?", (sid,))
                con.execute("DELETE FROM relations WHERE source_id=?", (sid,))
                con.execute("DELETE FROM source_columns WHERE source_id=?", (sid,))
                original_path=f"{package}!/{rel}" if package.suffix.lower()==".zip" else str(package/rel)
                values=(sid,rel,original_path,digest,len(content),modified,
                        None if Path(rel).suffix.lower() in STRUCTURED else safe_availability,
                        access_scope,Path(rel).suffix.lower(),"pending",None,cache_signature,permissions,
                        json.dumps(route,sort_keys=True))
                if prior:
                    con.execute("UPDATE sources SET path=?,original_path=?,sha256=?,size_bytes=?,modified_at=?,available_at=?,access_scope=?,format=?,extraction_status=?,error=?,pipeline_version=?,original_mode=?,routing_json=? WHERE id=?",
                                (*values[1:],sid))
                else:
                    con.execute("INSERT INTO sources VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",values)
            try:
                if route["kind"] == "unsupported":
                    route["selected_backend"] = "preserve"
                    with store.connect() as con:
                        con.execute("UPDATE sources SET extraction_status='preserved_unsupported',routing_json=? WHERE id=?",
                                    (json.dumps(route,sort_keys=True),sid))
                    result["files_preserved_unsupported"]+=1
                    result["routes"].append({"path":rel,**route})
                    continue
                if route["kind"] == "structured":
                    future = executor.submit(prepare_source,sid,rel,parse_path,access_scope,
                                             safe_availability,store.measurements,max_rows_per_file)
                    pending.append((sid,rel,route,future))
                    # Bound live workbook buffers rather than enqueueing the complete package.
                    if len(pending) >= workers:
                        finish(pending.pop(0))
                    continue
                attempts = []
                for backend in route["candidates"]:
                    try:
                        observations = _document_observations(
                            sid,rel,parse_path,access_scope,safe_availability,backend)
                        count = store.add_observations(observations)
                        if route["kind"] not in {"text"} and count == 0:
                            raise RuntimeError("document extraction produced no text")
                        route["selected_backend"] = backend
                        break
                    except Exception as exc:
                        attempts.append({"backend":backend,"error":f"{type(exc).__name__}: {exc}"})
                else:
                    route["attempts"] = attempts
                    raise RuntimeError("all document backends failed")
                if attempts:
                    route["attempts"] = attempts
                    route["fallback_used"] = True
                store.export_numerical(sid)
                with store.connect() as con:
                    route["extraction_methods"] = [row[0] for row in con.execute(
                        "SELECT DISTINCT extraction_method FROM observations WHERE source_id=? ORDER BY extraction_method",
                        (sid,))]
                if "tesseract" in route["extraction_methods"]:
                    route["ocr_used"] = True
                with store.connect() as con:
                    con.execute("UPDATE sources SET extraction_status='indexed',routing_json=? WHERE id=?",
                                (json.dumps(route,sort_keys=True),sid))
                result["files_indexed"] += 1; result["observations"] += count
                result["routes"].append({"path":rel,**route})
            except Exception as exc:
                with store.connect() as con:
                    con.execute("UPDATE sources SET extraction_status='error',error=?,routing_json=? WHERE id=?",
                                (str(exc),json.dumps(route,sort_keys=True),sid))
                result["errors"].append({"path":rel,"error":str(exc)})
                result["routes"].append({"path":rel,**route})
        for job in pending:
            finish(job)
    result['runtime_seconds'] = round(time.monotonic()-started,3)
    return result
