"""CMO delivery API mounted in the existing application; no arbitrary server paths."""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import os
import re
import threading
import time
import uuid
import zipfile
from typing import Literal
from contextvars import ContextVar
from pathlib import Path

from fastapi import APIRouter, Depends, Header, HTTPException, Request
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field

from . import quality
from .evidence import restore
from .persistence import current_tenant, set_tenant, tenant_scope
from .pipeline import CONFIGURATIONS, IngestionService, digest, hosted_inference_approved, utc
from .schema import VOCAB

ALLOWED_FILES={'.pdf','.csv','.xlsx','.xlsm','.tsv','.txt','.log','.docx','.json','.png','.jpg','.jpeg','.tif','.tiff'}
DELIVERY_LIMIT=256*1024*1024
FILE_LIMIT=64*1024*1024


class RunInput(BaseModel):
    package_id: str
    configuration: str = 'auto'
    # Sample packages only: run under a separate CMO workspace (for example one
    # per demo visitor) so confirmed formats from other sessions do not apply.
    cmo: str | None = Field(default=None,min_length=1,max_length=120)


class DeliveryInput(BaseModel):
    cmo: str = Field(min_length=1,max_length=120)
    name: str = Field(default='',max_length=200)


# Routine actions get a descriptive default reason; a typed note is optional.
# Rejections, corrections and dismissing missing information need a reason.
DEFAULT_REASONS={'accept':'Accepted after checking the source','template':'Confirmed column mapping',
                 'acknowledge':'Read; not process data','reprocess':'Apply confirmed mappings'}
NEEDS_REASON={'EXPECTED_DOCUMENT_MISSING','EXPECTED_PARAMETER_MISSING','BATCH_DOCUMENT_MISSING','MISSING_PAGE','MISSING_FIELD',
              'PARSE_ERROR','DUPLICATE_CONFLICT','DUPLICATE_RECORD','BATCH_MISMATCH','TIME_ORDER'}


def reason_for(kind, reason):
    reason=(reason or '').strip()
    if reason: return reason
    if kind not in DEFAULT_REASONS: raise HTTPException(400,'A reason is required for this decision')
    return DEFAULT_REASONS[kind]


class TemplateInput(BaseModel):
    fingerprint: str
    mapping: dict
    actor: str = Field(min_length=1)
    reason: str = ''
    origin: Literal['human','study_protocol'] = 'human'


class UploadInput(BaseModel):
    filename: str = Field(min_length=1,max_length=255)
    size: int = Field(ge=0)


class ReviewInput(BaseModel):
    measurement_id: str
    binding: str
    decision: str
    actor: str = Field(min_length=1)
    reason: str = ''
    value: str | None = None
    parameter: str | None = None
    unit: str | None = None
    batch: str | None = None
    time: str | None = None


class BulkItem(BaseModel):
    measurement_id: str
    binding: str


class BulkReviewInput(BaseModel):
    items: list[BulkItem] = Field(min_length=1,max_length=20000)
    decision: Literal['accept','reject']
    actor: str = Field(min_length=1)
    reason: str = ''


class DispositionInput(BaseModel):
    issue_id: str
    actor: str = Field(min_length=1)
    reason: str = ''
    ignore_in_future: bool = False
    also: list[str] = Field(default_factory=list,max_length=5000)


class ReprocessInput(BaseModel):
    actor: str = Field(min_length=1)
    reason: str = ''


class ActivityInput(BaseModel):
    session_id: str
    screen: str
    action: str


def safe_member(name):
    path=Path(name)
    if not name or path.is_absolute() or '..' in path.parts or '\\' in name or path.name.startswith('.'):
        raise ValueError('Unsafe file path')
    if path.suffix.lower() not in ALLOWED_FILES: raise ValueError(f'Unsupported file type: {path.suffix or "none"}')
    return path


def check_token(supplied):
    import secrets
    expected=os.getenv('CMO_API_TOKEN')
    if expected and (not supplied or not secrets.compare_digest(expected,supplied)):
        raise HTTPException(401,'CMO API token required')


def on_lambda():
    return bool(os.getenv('AWS_LAMBDA_FUNCTION_NAME'))


def router(repository: Path, tenant_dependency=None, tasks=None, permission_check=None):
    """``tenant_dependency`` resolves the caller's platform identity when mounted in
    the platform API; standalone deployments fall back to the X-Tenant-Id header.

    ``tasks`` runs work after the response (``src.lifesci.tasks.TaskRunner``): a
    thread on a long-lived server, an asynchronous invocation on AWS Lambda.
    Run progress is always read from the database, never from process memory.

    Every record query is scoped to the bound tenant (row-level security on
    PostgreSQL), so another customer's runs are simply not found. The
    dependency is async so the tenant it binds is inherited by the endpoint."""
    from .settings import load_local_settings
    load_local_settings()
    request_actor = ContextVar('cmo_request_actor', default=None)
    def recorded_actor(proposed):
        return request_actor.get() or proposed
    if tenant_dependency is None:
        async def authorize(x_cmo_token: str | None = Header(default=None),x_tenant_id: str | None = Header(default=None)):
            check_token(x_cmo_token); set_tenant(x_tenant_id)
    else:
        async def authorize(request: Request, x_cmo_token: str | None = Header(default=None),ctx=Depends(tenant_dependency)):
            check_token(x_cmo_token)
            # A signed-in account is not yet a member: someone must have assigned it a role.
            if not ctx.roles: raise HTTPException(403,'Your account has no role in this workspace yet. Ask an administrator to assign one.')
            if ctx.site_ids:
                raise HTTPException(403, 'Delivery records currently have organization-wide scope; site-restricted accounts cannot access this workspace')
            if permission_check:
                path = request.url.path
                action = 'read' if request.method == 'GET' else 'propose'
                if path.endswith('/export'): action = 'export'
                if '/reviews' in path or path.endswith('/issues'): action = 'approve'
                permission_check(ctx, action)
            request_actor.set(ctx.actor_id)
            set_tenant(ctx.tenant_id)
    api=APIRouter(prefix='/api/cmo-ingestion',dependencies=[Depends(authorize)])
    root=Path(os.getenv('CMO_INGESTION_ROOT',str(repository/'.cache'/'cmo-ingestion')))
    packages=Path(os.getenv('CMO_PACKAGES_ROOT',str(repository/'benchmarks'/'cmo_ingestion'/'generated')))
    demo=Path(os.getenv('CMO_DEMO_ROOT',str(repository/'benchmarks'/'cmo_ingestion'/'demo')))
    service=IngestionService(root)
    activity={}; lock=threading.RLock()
    if tasks is None:
        from ...lifesci.tasks import TaskRunner
        tasks=TaskRunner()
    # A run still queued or running after this long has lost its worker (Lambda
    # stops a task at 15 minutes). Unset on a long-lived server, where a restart
    # is what interrupts work.
    run_timeout=float(os.getenv('CMO_RUN_TIMEOUT_SECONDS') or (960 if on_lambda() else 0))

    if not on_lambda():
        # Worker threads do not survive a restart; say so instead of showing a
        # run that is "running" forever. On Lambda other instances may still be
        # working, so staleness is judged by age instead (see interrupt_if_stale).
        for tenant in service.repo.tenants():
            with tenant_scope(tenant):
                for stale in service.repo.run_headers():
                    if stale['status'] in {'queued','running'}:
                        service.repo.update_run(stale['id'],status='interrupted',
                            error='The service restarted during processing. Start the run again; no partial results were accepted.')

    def interrupt_if_stale(header):
        if not run_timeout or header['status'] not in {'queued','running'}: return header
        from datetime import datetime,timezone
        age=(datetime.now(timezone.utc)-datetime.fromisoformat(header['started_at'])).total_seconds()
        if age<=run_timeout: return header
        error='Processing stopped before it finished (time limit reached). Start the run again; no partial results were accepted.'
        service.repo.update_run(header['id'],status='interrupted',error=error)
        return {**header,'status':'interrupted','error':error}

    def guard(function):
        try: return function()
        except KeyError as exc: raise HTTPException(404,str(exc).strip("'")) from exc
        except FileNotFoundError as exc: raise HTTPException(410,str(exc)) from exc
        except (ValueError,PermissionError) as exc: raise HTTPException(400,str(exc)) from exc

    def catalog():
        entries=[]
        for folder in [demo,packages]:
            if (folder/'catalog.json').exists():
                entries.extend({**entry,'kind':'sample','root':str(folder)} for entry in json.loads((folder/'catalog.json').read_text())
                               if not any(e['id']==entry['id'] for e in entries))
        uploads=root/'uploads'
        for file in sorted(uploads.glob('*/delivery.json')):
            entries.append({'kind':'upload',**json.loads(file.read_text())})
        entries.extend({'kind':'delivery',**item} for item in service.repo.deliveries())
        return entries

    def public_entry(item):
        return {k:v for k,v in item.items() if k!='root'}

    def delivery_folder(ident):
        if not ident.startswith('delivery-') or Path(ident).name!=ident: raise ValueError('Invalid delivery ID')
        return root/'deliveries'/ident

    def find(package_id):
        item=next((i for i in catalog() if i['id']==package_id),None)
        if not item: raise KeyError('Unknown package')
        if item.get('kind')=='delivery':
            folder=delivery_folder(package_id)
            for file in item['files']:
                restore(service.evidence,f"deliveries/{package_id}/{file['path']}",folder/file['path'],file['sha256'])
            return item,folder
        if package_id.startswith('upload-'): path=root/'uploads'/package_id/'raw.zip'
        else:
            if Path(package_id).name!=package_id: raise ValueError('Invalid package ID')
            path=Path(item.get('root',packages))/package_id/'raw'
        return item,path

    def public(run):
        run=dict(run)
        # Private store paths and the per-element ledger are not public API fields.
        run.pop('store',None); elements=run.pop('elements',None)
        run['measurements']=[{**m,'review_binding':digest(m)} for m in run.get('measurements',[])]
        if run.get('status')=='complete':
            run['review_groups']=quality.review_groups(run)
            run['guarantee']=quality.guarantee({**run,'elements':elements})
        return run

    def approved_for(item):
        # Public datasets may use hosted extraction; customer files need the deployment approval.
        return bool(item.get('public_data')) or (bool(hosted_inference_approved()) and not item.get('synthetic'))

    def launch(item,path,configuration,*,supersedes=None,cmo=None):
        if configuration not in CONFIGURATIONS: raise HTTPException(400,'Unknown configuration')
        run_id=uuid.uuid4().hex
        service.queue(run_id,cmo or item['cmo'],configuration,synthetic=item.get('synthetic',False),approved=approved_for(item),
                      package_ref=item['id'],details={'disclosure':item.get('disclosure'),'delivery_name':item.get('name')})
        tasks.submit('cmo_run',{'tenant':current_tenant(),'run_id':run_id,'package_id':item['id'],'configuration':configuration,
                                'cmo':cmo,'supersedes':supersedes})
        return {'id':run_id,'status':'queued'}

    def run_job(payload):
        """Task handler: process one queued run. Safe to receive twice (only the first claim proceeds)."""
        with tenant_scope(payload['tenant']):
            run_id=payload['run_id']
            if not service.repo.claim_queued(run_id): return None
            try:
                item,path=find(payload['package_id'])
                result=service.run(path,payload['cmo'] or item['cmo'],payload['configuration'],synthetic=item.get('synthetic',False),
                                   approved=approved_for(item),run_id=run_id,package_ref=item['id'])
                result.update(package_id=item['id'],disclosure=item.get('disclosure'),delivery_name=item.get('name'))
                if item.get('public_data') and result.get('hosted_inference'):
                    result['hosted_inference']['basis']='public dataset (CC BY 4.0); hosted extraction permitted'
                service.save(result)
                if payload['supersedes'] and result['status']=='complete': result=service.carry_over(payload['supersedes'],result)
                return result['status']
            except Exception as exc:
                header=service.repo.run_header(run_id)
                if header and header['status'] in {'queued','running'}:
                    service.repo.update_run(run_id,status='failed',error=f'{type(exc).__name__}: {exc}')
                raise

    tasks.register('cmo_run',run_job)

    @api.get('/catalog')
    def list_packages():
        approved=hosted_inference_approved()
        return {'packages':[public_entry(i) for i in catalog()],'configurations':CONFIGURATIONS,
                'availability':{c:service.availability(c,synthetic=True) for c in CONFIGURATIONS},
                'customer_availability':{c:service.availability(c,approved=approved) for c in ['clairic_decisions','decisions']},
                'hosted_inference_approved':approved,
                'storage':{'records':service.repo.backend,'originals':service.evidence.backend},
                'cmos':sorted({r['cmo'] for r in service.repo.run_headers()}|{i['cmo'] for i in catalog() if i.get('kind')=='delivery'}),
                'upload':{'mode':upload_mode(),'max_bytes':FILE_LIMIT},
                'deployment':'Set CMO_API_TOKEN before exposing beyond loopback.'}

    def upload_mode():
        # Browsers send files straight to private storage when it is configured;
        # AWS Lambda accepts at most 6 MB per request.
        if os.getenv('CMO_UPLOAD_MODE','').strip()=='direct' or service.evidence.backend!='supabase': return 'direct'
        return 'storage'

    @api.get('/vocabulary')
    def vocabulary():
        return {'parameters':[{'id':p,'unit':v['unit'],'aliases':v['aliases'],'range':list(v['range'])} for p,v in VOCAB.items()]}

    @api.post('/deliveries')
    def create_delivery(value:DeliveryInput):
        ident='delivery-'+uuid.uuid4().hex[:16]
        item={'id':ident,'cmo':value.cmo.strip(),'name':value.name.strip() or f'{value.cmo.strip()} delivery','created':utc(),
              'dataset':'customer','condition':'delivery','synthetic':False,'files':[],
              'disclosure':'Customer delivery. Synthetic status and provenance are as stated by the uploader.'}
        delivery_folder(ident).mkdir(parents=True,mode=0o700)
        service.repo.save_delivery(item); service.audit('delivery_created',delivery=item)
        return item

    def delivery(delivery_id):
        item=next((i for i in service.repo.deliveries() if i['id']==delivery_id),None)
        if not item: raise HTTPException(404,'Unknown delivery')
        return item

    def store_files(delivery_id,filename,data):
        """Validate, unpack, hash and durably store one uploaded file or ZIP, then append it to the delivery."""
        delivery(delivery_id)
        folder=guard(lambda:delivery_folder(delivery_id))
        if len(data)>FILE_LIMIT: raise HTTPException(413,'Per-file upload limit is 64 MiB')
        def unpack():
            name=Path(filename).name
            if Path(name).suffix.lower()!='.zip':
                return [(safe_member(name),bytes(data))]
            entries=[]
            with zipfile.ZipFile(io.BytesIO(bytes(data))) as archive:
                members=[m for m in archive.infolist() if not m.is_dir()]
                if len(members)>500 or sum(m.file_size for m in members)>128*1024*1024: raise ValueError('ZIP exceeds expansion limits')
                for member in members:
                    if member.flag_bits&1 or (member.external_attr>>16)&0o170000==0o120000: raise ValueError('Encrypted or linked ZIP member')
                    if Path(member.filename).name.startswith('.') or '__MACOSX' in member.filename: continue
                    entries.append((safe_member(member.filename),archive.read(member)))
            return entries
        try: entries=unpack()
        except (ValueError,zipfile.BadZipFile) as exc: raise HTTPException(400,str(exc)) from exc
        with lock:
            item=delivery(delivery_id)
            existing={f['path'] for f in item['files']}
            if sum(f['size'] for f in item['files'])+sum(len(b) for _,b in entries)>DELIVERY_LIMIT:
                raise HTTPException(413,'Delivery exceeds 256 MiB')
            added=[]
            for path,content in entries:
                rel=path.as_posix()
                if rel in existing: raise HTTPException(409,f'{rel} is already in this delivery; start a new delivery for a revised file')
                target=folder/rel; target.parent.mkdir(parents=True,exist_ok=True)
                target.write_bytes(content); target.chmod(0o400)
                sha=hashlib.sha256(content).hexdigest()
                service.evidence.put(f'deliveries/{delivery_id}/{rel}',target)
                added.append({'path':rel,'size':len(content),'sha256':sha}); existing.add(rel)
            item['files']=[*item['files'],*added]
            service.repo.save_delivery(item)
        service.audit('delivery_files_added',delivery_id=delivery_id,files=added)
        return item

    @api.post('/deliveries/{delivery_id}/files')
    async def add_file(delivery_id:str,request:Request,filename:str):
        delivery(delivery_id)
        data=bytearray()
        async for chunk in request.stream():
            data.extend(chunk)
            if len(data)>FILE_LIMIT: raise HTTPException(413,'Per-file upload limit is 64 MiB')
        return store_files(delivery_id,filename,bytes(data))

    # Large files go browser -> private storage -> backend. The pending object is
    # only a transport copy: it is validated and hashed like a direct upload,
    # stored under the delivery's normal key, then deleted.
    def pending_key(delivery_id,upload_id):
        if not re.fullmatch(r'[0-9a-f]{32}-[A-Za-z0-9_-]{1,400}',upload_id): raise HTTPException(400,'Invalid upload ID')
        delivery_folder(delivery_id)
        return f'uploads/{current_tenant()}/{delivery_id}/{upload_id}'

    def upload_filename(upload_id):
        encoded=upload_id.split('-',1)[1]
        try: return base64.urlsafe_b64decode(encoded+'='*(-len(encoded)%4)).decode()
        except (ValueError,UnicodeDecodeError) as exc: raise HTTPException(400,'Invalid upload ID') from exc

    @api.post('/deliveries/{delivery_id}/uploads')
    def begin_upload(delivery_id:str,value:UploadInput):
        delivery(delivery_id)
        if upload_mode()!='storage': raise HTTPException(409,'Upload directly to /files; no private storage is configured')
        if value.size>FILE_LIMIT: raise HTTPException(413,'Per-file upload limit is 64 MiB')
        name=Path(value.filename).name
        if Path(name).suffix.lower()!='.zip': guard(lambda:safe_member(name))
        upload_id=uuid.uuid4().hex+'-'+base64.urlsafe_b64encode(name.encode()).decode().rstrip('=')
        signed=service.evidence.signed_upload(pending_key(delivery_id,upload_id))
        return {'upload_id':upload_id,**signed}

    @api.post('/deliveries/{delivery_id}/uploads/{upload_id}/complete')
    def complete_upload(delivery_id:str,upload_id:str):
        key=pending_key(delivery_id,upload_id)
        local=root/'pending'/uuid.uuid4().hex
        try:
            guard(lambda:service.evidence.get(key,local))
            if local.stat().st_size>FILE_LIMIT: raise HTTPException(413,'Per-file upload limit is 64 MiB')
            item=store_files(delivery_id,upload_filename(upload_id),local.read_bytes())
        finally: local.unlink(missing_ok=True)
        service.evidence.delete(key)
        return item

    @api.post('/runs')
    def start(value:RunInput):
        item,path=guard(lambda:find(value.package_id))
        if item.get('kind')=='delivery' and not item['files']: raise HTTPException(400,'Add at least one file to the delivery first')
        if value.cmo and item.get('kind')!='sample': raise HTTPException(400,'Only sample packages can run under another CMO workspace')
        return launch(item,path,value.configuration,cmo=value.cmo.strip() if value.cmo else None)

    @api.get('/runs')
    def list_runs(cmo:str|None=None):
        result=[]
        for run in service.runs(cmo):
            result.append({k:run.get(k) for k in ['id','cmo','status','configuration','requested_configuration','started_at','completed_at',
                                                 'package_id','delivery_name','superseded_by','runtime_seconds']}
                          |{'summary':run.get('summary'),'documents':len(run.get('documents',[]))})
        return {'runs':list(reversed(result))}

    @api.get('/runs/{run_id}')
    def get_run(run_id:str):
        header=service.repo.run_header(run_id)
        if header is None: raise HTTPException(404,'Unknown run')
        interrupt_if_stale(header)
        run=guard(lambda:service.get(run_id))  # one read decides the view; the run may finish meanwhile
        if run['status'] in {'queued','running'}:
            return {'id':run_id,'status':run['status'],'stages':run['stages']}
        result=public(run)
        if result.get('guarantee') is not None:
            result['guarantee']['checks'].append({'id':'audit_chain','label':'The append-only audit chain verifies end to end',
                'passed':service.verify_audit(),'detail':'hash chain over every decision in this workspace'})
            result['guarantee']['holds']=all(c['passed'] for c in result['guarantee']['checks'])
        return result

    @api.post('/runs/{run_id}/templates')
    def confirm(run_id:str,value:TemplateInput):
        return guard(lambda:service.confirm_template(run_id,value.fingerprint,value.mapping,recorded_actor(value.actor),reason_for('template',value.reason),origin=value.origin))

    @api.post('/runs/{run_id}/reprocess')
    def reprocess(run_id:str,value:ReprocessInput):
        old=guard(lambda:service.get(run_id))
        if old.get('superseded_by'): raise HTTPException(409,f"Already reprocessed as run {old['superseded_by']}")
        item,path=guard(lambda:find(old.get('package_ref') or old.get('package_id') or ''))
        service.audit('reprocess_requested',run_id=run_id,actor=recorded_actor(value.actor),reason=reason_for('reprocess',value.reason))
        return launch(item,path,old.get('requested_configuration') or old['configuration'],supersedes=run_id,cmo=old['cmo'])

    @api.post('/runs/{run_id}/reviews')
    def review(run_id:str,value:ReviewInput):
        reason=reason_for(value.decision,value.reason)
        result=guard(lambda:service.review(run_id,value.measurement_id,value.decision,recorded_actor(value.actor),reason,
                        binding=value.binding,value=value.value,parameter=value.parameter,unit=value.unit,batch=value.batch,time_value=value.time))
        return {'summary':result['summary']}

    @api.post('/runs/{run_id}/reviews/bulk')
    def bulk_review(run_id:str,value:BulkReviewInput):
        reason=reason_for(value.decision,value.reason)
        return guard(lambda:service.review_many(run_id,[i.model_dump() for i in value.items],value.decision,recorded_actor(value.actor),reason))

    @api.post('/runs/{run_id}/issues')
    def disposition(run_id:str,value:DispositionInput):
        run=guard(lambda:service.get(run_id))
        issue=next((i for i in run['issues'] if i['id']==value.issue_id),None)
        kind='needs_reason' if issue and issue['rule_code'] in NEEDS_REASON else 'acknowledge'
        reason=reason_for(kind,value.reason)
        result=guard(lambda:service.resolve_issue(run_id,value.issue_id,recorded_actor(value.actor),reason,ignore_in_future=value.ignore_in_future,also=value.also))
        return {'summary':result['summary']}

    @api.get('/runs/{run_id}/sources/{source_id}')
    def source(run_id:str,source_id:str):
        path,doc=guard(lambda:service.original(run_id,source_id))
        return FileResponse(path,filename=Path(doc['path']).name,content_disposition_type='inline')

    @api.get('/runs/{run_id}/cells/{measurement_id}')
    def cell(run_id:str,measurement_id:str):
        run=guard(lambda:service.get(run_id)); m=next((m for m in run['measurements'] if m['id']==measurement_id),None)
        if not m: raise HTTPException(404,'Unknown measurement')
        path,doc=guard(lambda:service.original(run_id,m['source_id'],run=run)); ref=m['source']
        if path.suffix.lower() in {'.xlsx','.xlsm'}:
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=False)
            try:
                sheet=book[ref['sheet']]; row=ref['row']
                headers=[c.value for c in next(sheet.iter_rows(min_row=1,max_row=1))]
                values=[c.value for c in next(sheet.iter_rows(min_row=row,max_row=row))]
            finally: book.close()
        elif path.suffix.lower() in {'.csv','.tsv'}:
            with path.open(newline='',encoding='utf-8-sig') as stream:
                rows=csv.reader(stream,delimiter='\t' if path.suffix.lower()=='.tsv' else ',')
                headers=next(rows); values=next((r for i,r in enumerate(rows,2) if i==ref['row']),[])
        else: raise HTTPException(400,'Use the original PDF viewer for this measurement')
        return {'source':ref,'headers':headers,'values':[None if v is None else str(v) for v in values],'source_sha256':doc['sha256']}

    @api.get('/runs/{run_id}/pages/{source_id}/{page}')
    def page_image(run_id:str,source_id:str,page:int):
        path,_doc=guard(lambda:service.original(run_id,source_id))
        if path.suffix.lower()!='.pdf': raise HTTPException(400,'Source is not a PDF')
        import pypdfium2
        pdf=pypdfium2.PdfDocument(str(path))
        try:
            if not 1<=page<=len(pdf): raise HTTPException(400,'Page outside document')
            selected=pdf[page-1]
            try:
                width,height=selected.get_size()
                bitmap=selected.render(scale=1.5)
                try:
                    output=io.BytesIO(); bitmap.to_pil().save(output,format='PNG')
                    return Response(output.getvalue(),media_type='image/png',
                                    headers={'X-Page-Width':str(width),'X-Page-Height':str(height),'X-Page-Count':str(len(pdf))})
                finally: bitmap.close()
            finally: selected.close()
        finally: pdf.close()

    @api.get('/runs/{run_id}/export')
    def export(run_id:str,format:str='csv',wide:bool=False):
        content=guard(lambda:service.export(run_id,format=format,wide=wide))
        run=guard(lambda:service.get(run_id))
        return Response(content,media_type='text/csv' if format=='csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition':f'attachment; filename="verified-{run_id}.{format}"',
                     'X-Package-Complete':str(run['summary']['verified_package']).lower(),
                     'X-Open-Issues':str(run['summary']['open_issues']),
                     'X-Values-Pending':str(run['summary']['needs_review'])})

    def verified_payload(run_id:str,tenant:str|None=None):
        if tenant is not None:
            # Same-process callers (the platform API) pass their authenticated tenant.
            with tenant_scope(tenant): return verified_payload(run_id)
        run=guard(lambda:service.get(run_id))
        if run['status']!='complete' or not run.get('summary',{}).get('verified_package'):
            raise HTTPException(409,'Package still requires verification; downstream consumption is blocked')
        check=quality.guarantee(run)
        if not check['holds']: raise HTTPException(409,'Guarantee checks failed; downstream consumption is blocked')
        measurements=[m for m in run['measurements'] if m['status'] in {'auto_accepted','accepted','corrected'}]
        for sid in {m['source_id'] for m in measurements}: guard(lambda:service.original(run_id,sid,run=run))
        return {'schema_version':'verified-cmo-1','run_id':run_id,'cmo':run['cmo'],'verified':True,
                'package_id':run.get('package_id'),'delivery_name':run.get('delivery_name'),'completed_at':run.get('completed_at'),
                'superseded_by':run.get('superseded_by'),'disclosure':run.get('disclosure'),
                'coverage':{k:v for k,v in run['coverage'].items() if k!='by_document'},'guarantee':check,
                'measurements':measurements,'relationships':run.get('relationships',[])}

    @api.get('/runs/{run_id}/verified')
    def verified(run_id:str):
        return verified_payload(run_id)

    @api.post('/upload')
    async def upload(request:Request,dataset:str='sanofi'):
        if on_lambda(): raise HTTPException(400,'Sample ZIP uploads need a long-lived server. Create a delivery and add the files instead.')
        if dataset not in {'sanofi','indpensim'}: raise HTTPException(400,'Unknown dataset')
        data=bytearray()
        async for chunk in request.stream():
            data.extend(chunk)
            if len(data)>64*1024*1024: raise HTTPException(413,'Demo ZIP upload limit is 64 MiB')
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members=archive.infolist()
                if len(members)>500 or sum(i.file_size for i in members)>128*1024*1024:
                    raise ValueError('ZIP exceeds demo expansion limits')
                allowed={'.pdf','.csv','.xlsx','.xlsm','.tsv','.txt','.log','.png','.jpg','.jpeg','.tif','.tiff'}
                for member in members:
                    p=Path(member.filename)
                    if p.is_absolute() or '..' in p.parts or '\\' in member.filename: raise ValueError('Unsafe ZIP path')
                    if not member.is_dir() and p.suffix.lower() not in allowed: raise ValueError('Unsupported ZIP file type')
        except (ValueError,zipfile.BadZipFile) as exc: raise HTTPException(400,str(exc)) from exc
        ident='upload-'+uuid.uuid4().hex; folder=root/'uploads'/ident; folder.mkdir(parents=True)
        (folder/'raw.zip').write_bytes(data)
        # Uploaded data is not implicitly approved for hosted inference.
        item={'id':ident,'dataset':dataset,'condition':'uploaded','cmo':dataset+'-upload',
              'synthetic':False,'disclosure':'User upload. Provenance and synthetic status have not been independently established.',
              'files':[{'path':m.filename,'size':m.file_size} for m in members if not m.is_dir()]}
        (folder/'delivery.json').write_text(json.dumps(item))
        return item

    @api.post('/runs/{run_id}/activity')
    def record_activity(run_id:str,value:ActivityInput):
        guard(lambda:service.get(run_id))
        if value.screen not in {'review','templates'} or value.action not in {'start','heartbeat','stop'}:
            raise HTTPException(400,'Invalid activity event')
        if not value.session_id or len(value.session_id)>100: raise HTTPException(400,'Invalid session ID')
        now=time.monotonic(); ident=(run_id,value.session_id)
        with lock, service._lock:
            previous=activity.get(ident)
            # Heartbeats every 10s, visibility/focus + real input required in UI.
            # Gaps >15s never count as active work. No client-supplied durations.
            seconds=(now-previous) if previous is not None and now-previous<=15 and value.action!='start' else 0
            if value.action=='stop': activity.pop(ident,None)
            else: activity[ident]=now
            with service.repo.transaction():
                run=service.get(run_id)
                minutes=(run['human_minutes'] or 0)+seconds/60
                status='observed_interaction_partial; not a completed human study'
                service.repo.update_run(run_id,human_minutes=minutes,human_time_status=status)
                service.audit('human_activity',run_id=run_id,screen=value.screen,action_name=value.action,
                              session_id=value.session_id,credited_seconds=seconds)
        return {'human_minutes':minutes,'status':status}

    @api.get('/comparison')
    def comparison():
        report=Path(os.getenv('CMO_REPORT_PATH',str(repository/'results'/'cmo_ingestion'/'report.json')))
        return json.loads(report.read_text()) if report.exists() else {'results':[],'human_study':'No executed benchmark report is available. Run the evaluator first.'}

    @api.get('/audit')
    def audit(): return {'valid':service.verify_audit()}

    # Same-process consumers (the platform API) read verified datasets without an HTTP hop.
    api.verified_payload=verified_payload
    api.ingestion_service=service
    return api
