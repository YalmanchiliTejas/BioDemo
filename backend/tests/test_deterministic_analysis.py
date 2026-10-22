from __future__ import annotations

import math
import sqlite3

import pytest
from openpyxl import Workbook

from src.manufacturing_data import ManufacturingStore,ManufacturingTools,ingest_package
from src.prime_runtime.bridge import ManufacturingBridge
from src.state import BlackboardStore
from src.workflows import ManufacturingInvestigationWorkflow


@pytest.fixture
def dataset(tmp_path):
    source=tmp_path/'source';source.mkdir()
    for experiment in (1,2):
        book=Workbook();book.remove(book.active)
        for reactor in (1,2):
            sheet=book.create_sheet(f'R0{reactor}_offline')
            sheet.append(['Time (h)','VCD Mean (cells/mL)','TCD Mean (cells/mL)','VCD SEM (cells/mL)','Glucose (g/L)'])
            for t,v,g in [(0,1,10),(2,3,8),(7,6,5)]:sheet.append([t,v,10,.2 if experiment==1 else -.2,g])
            inline=book.create_sheet(f'R0{reactor}_inline')
            inline.append(['Time (h)','Signal (g/L)','Status'])
            for t in (0,1,3,7):inline.append([t,10+t,'Ok'])
        book.save(source/f'EXP00{experiment}.xlsx')
    (source/'replicates.csv').write_text('batch_id,Time (h),Replicate A (g/L),Replicate B (g/L),Replicate C (g/L)\nB1,0,1,2,3\n')
    store=ManufacturingStore(tmp_path/'store')
    assert not ingest_package(source,store)['errors']
    tools=ManufacturingTools(store,as_of='2030-01-01T00:00:00Z')
    rows=store.visible(limit=None)
    def ids(name,entity='EXP001/R01'):
        return [r['id'] for r in sorted(rows,key=lambda r:r['source_row']) if r['name']==name and r['entity_id']==entity]
    return store,tools,rows,ids


def test_source_times_reconstruct_offline_and_inline_without_row_proxies(dataset):
    store,tools,rows,ids=dataset
    selected=store.get_many(ids('VCD Mean (cells/mL)'))
    assert [r['time_value'] for r in selected]==[0,2,7]
    assert all(r['time_basis']=='elapsed_hours' for r in selected)
    assert all(r['time_evidence_id'].endswith(':1') for r in selected)
    args={'left':{'name':'VCD Mean (cells/mL)','entity_id':'EXP001/R01'},
          'right':{'name':'Signal (g/L)','entity_id':'EXP001/R01'}}
    exact=tools.align_series(**args)
    assert [r['status'] for r in exact['records']]==['exact','unmatched','exact']
    nearest=tools.align_series(**args,method='nearest',tolerance_hours=1)
    assert nearest['records'][1]['status']=='ambiguous_nearest_tie'
    linear=tools.align_series(**args,method='linear',tolerance_hours=1)
    assert linear['records'][1]['right_value']==12
    assert linear['records'][1]['weights']==[.5,.5]
    assert len(linear['records'][1]['right_ids'])==2
    assert linear['extrapolation'] is False and linear['source_measurements_modified'] is False
    assert store.get_many(ids('VCD Mean (cells/mL)'))==selected


def test_statistics_percentages_and_yields_are_recomputed(dataset):
    _,tools,_,ids=dataset
    replicate_ids=[ids(f'Replicate {letter} (g/L)','B1')[0] for letter in 'ABC']
    specs=[{'metric_id':'mean','operation':'mean','evidence_ids':replicate_ids,'claimed_value':2},
           {'metric_id':'sem','operation':'sem','evidence_ids':replicate_ids,'sampling_basis':'independent_replicates','claimed_value':1/math.sqrt(3)},
           {'metric_id':'viability','operation':'percentage','numerator_ids':[ids('VCD Mean (cells/mL)')[-1]],'denominator_ids':[ids('TCD Mean (cells/mL)')[-1]],'claimed_value':60},
           {'metric_id':'change','operation':'percent_change','numerator_ids':[replicate_ids[-1]],'denominator_ids':[replicate_ids[0]],'claimed_value':200},
           {'metric_id':'yield','operation':'yield','numerator_ids':ids('VCD Mean (cells/mL)')[::2],
            'denominator_ids':ids('Glucose (g/L)')[::2],'definition':'apparent VCD increase per glucose concentration decrease','claimed_value':1}]
    result=tools.calculate_metrics(metrics=specs)
    assert [m['value'] for m in result['metrics']]==pytest.approx([2,1/math.sqrt(3),60,200,1])
    assert all(m['verification']=='arithmetic_verified' and m['lineage_sha256'] for m in result['metrics'])
    assert result['metrics'][-1]['assumptions']
    with pytest.raises(ValueError,match='does not match'):
        tools.calculate_metrics(metrics=[{**specs[0],'claimed_value':200}])
    with pytest.raises(ValueError,match='SEM requires'):
        tools.calculate_metrics(metrics=[{'operation':'sem','selection':{'name':'VCD Mean (cells/mL)','entity_id':'EXP001/R01'},'sampling_basis':'independent_replicates'}])
    with pytest.raises(ValueError,match='reported means'):
        tools.calculate_metrics(metrics=[{'operation':'sem','evidence_ids':[ids('VCD Mean (cells/mL)')[0],ids('TCD Mean (cells/mL)')[0]],'sampling_basis':'independent_replicates'}])
    with pytest.raises(ValueError,match='duplicate'):
        tools.calculate_metrics(metrics=[{'operation':'mean','evidence_ids':[replicate_ids[0]]*2}])
    with pytest.raises(ValueError,match='exact start/end times'):
        tools.calculate_metrics(metrics=[{**specs[-1],'denominator_ids':ids('Glucose (g/L)')[:2]}])


def test_preparation_flags_invalid_source_sem_without_inventing_causation(dataset):
    _,tools,_,_=dataset
    prepared=tools.analysis().prepare()
    assert prepared['series_count']>0
    assert any('negative_reported_sem' in a['flags'] for a in prepared['anomalies'])
    assert all(s['time_basis']=='elapsed_hours' for s in prepared['series'])
    assert any(s['source_summary_status']=='reported_sem_not_recomputed' for s in prepared['series'])
    assert prepared['paired_peak_comparisons']
    assert all('design equivalence not established' in c['status'] for c in prepared['paired_peak_comparisons'])


def test_nested_hypothesis_citations_survive_and_scope_is_multi_experiment(dataset,tmp_path):
    store,tools,rows,ids=dataset
    boards=BlackboardStore(tmp_path/'boards')
    workflow=ManufacturingInvestigationWorkflow(tmp_path/'workflow',store,blackboards=boards)
    bridge=ManufacturingBridge(store=store,blackboards=boards,workflow=workflow,scope='default',incident_id='case',
                              objective='Investigate',as_of='2030-01-01T00:00:00Z',seed_ids=[r['id'] for r in rows])
    support=ids('VCD Mean (cells/mL)')[-1];counter=ids('VCD Mean (cells/mL)','EXP002/R01')[-1]
    hypothesis={'hypothesis_id':'H1','statement':'Competing explanation','supporting_evidence_ids':[support],
                'contradicting_evidence_ids':[counter]}
    with pytest.raises(ValueError,match='does not match'):
        bridge.submit_report({'assessment':'Bad calculation','hypotheses':[hypothesis],
                              'calculated_metrics':[{'operation':'mean','evidence_ids':[support],'claimed_value':999}]})
    assert boards.load('case') is None
    result=bridge.submit_report({'assessment':'Mixed evidence; mechanism unproven','severity':'CRITICAL',
                                 'hypotheses':[hypothesis],
                                 'observed_facts':[{'statement':'Measured source value','observation_id':support,'claimed_value':6}],
                                 'calculated_metrics':[{'operation':'mean','evidence_ids':[support,counter],'claimed_value':6}]})
    assert set(result['evidence_ids'])=={support,counter}
    assert result['scope']['type']=='multi_experiment'
    assert result['scope']['experiment_ids']==['EXP001','EXP002']
    assert result['scope']['batch_ids']==[]
    assert result['severity']=='UNKNOWN' and result['agent_proposed_severity']=='CRITICAL'
    assert result['triage']['policy_status']=='draft'
    assert result['causal_hypotheses'][0]['supporting_evidence_ids']==[support]
    assert result['causal_hypotheses'][0]['contradicting_evidence_ids']==[counter]
    assert result['verification']['numerical_metrics_verified']==1
    assert result['verification']['causality_verified'] is False
    assert boards.load('case').calculated_metrics


def test_hypothesis_requires_explicit_missing_evidence_and_keeps_absence_reason(dataset,tmp_path):
    store,_,rows,ids=dataset
    boards=BlackboardStore(tmp_path/'boards');workflow=ManufacturingInvestigationWorkflow(tmp_path/'workflow',store)
    seed=ids('VCD Mean (cells/mL)')[0]
    bridge=ManufacturingBridge(store=store,blackboards=boards,workflow=workflow,scope='default',incident_id='case',
                              objective='Investigate',as_of='2030-01-01T00:00:00Z',seed_ids=[seed])
    report={'assessment':'Unresolved','evidence_ids':[seed],
            'hypotheses':[{'hypothesis_id':'H1','statement':'Candidate','supporting_evidence_ids':[seed]}]}
    with pytest.raises(ValueError,match='empty contradicting'):bridge.submit_report(report)
    report['hypotheses'][0].update(contradicting_evidence_status='not_searched',contradicting_evidence_absence_reason='Tool budget ended before a counterevidence search.')
    result=bridge.submit_report(report)
    assert result['causal_hypotheses'][0]['contradicting_evidence_absence_reason'].startswith('Tool budget')
    assert result['verification']['hypotheses_without_contradiction_search']==['H1']


def test_citation_lookup_stays_below_database_parameter_limits(tmp_path,monkeypatch):
    source=tmp_path/'source';source.mkdir()
    (source/'data.csv').write_text('batch_id,value\n'+''.join(f'B1,{i}\n' for i in range(1200)))
    store=ManufacturingStore(tmp_path/'store');ingest_package(source,store)
    ids=[r['id'] for r in store.visible(limit=None)]
    original=store.connect
    def limited():
        con=original();con.setlimit(sqlite3.SQLITE_LIMIT_VARIABLE_NUMBER,600);return con
    monkeypatch.setattr(store,'connect',limited)
    assert [r['id'] for r in store.get_many(ids)]==ids


def test_status_expansion_ignores_routine_instrument_rows(tmp_path):
    source=tmp_path/'source';source.mkdir()
    (source/'sensor.csv').write_text('batch_id,Time (h),Status,value\n'+''.join(f'B1,{i},Ok,1\n' for i in range(35000))+'B1,35000,action,1\n')
    store=ManufacturingStore(tmp_path/'store');ingest_package(source,store)
    workflow=ManufacturingInvestigationWorkflow(tmp_path/'workflow',store)
    found=workflow._related_status_ids('2030-01-01T00:00:00Z',['B1'])
    assert len(found)==1
    row=store.get_many(found)[0]
    assert row['value_text']=='action'
    report=workflow.record_rollout('case','2030-01-01T00:00:00Z','Investigate',{'evidence_ids':[found[0]],'related_batches':['B1'],'causal_hypotheses':[]})
    assert len(report['evidence_ids'])==1
    assert report['triage']['level']=='L1'


def test_custom_policy_controls_severity_and_rejects_unreviewed_approval(dataset,tmp_path):
    from copy import deepcopy
    from src.workflows.triage import DEFAULT_POLICY
    store,_,_,_=dataset
    policy=deepcopy(DEFAULT_POLICY);policy.update(version='site-1',status='approved',reviewed_by=['qa-owner'])
    policy['severity_by_level']['L1']='LOW'
    workflow=ManufacturingInvestigationWorkflow(tmp_path/'workflow',store,triage_policy=policy)
    triage=workflow._triage([{'signal':'alert','entity_id':'B1','evidence_id':'test','validation_status':'source_transcribed'}])
    assert triage['severity']=='LOW' and triage['provisional'] is False and triage['policy']=='site-1'
    policy['reviewed_by']=[]
    with pytest.raises(ValueError,match='named reviewers'):
        ManufacturingInvestigationWorkflow(tmp_path/'bad',store,triage_policy=policy)


def test_excel_native_precision_and_numeric_cells_in_mixed_columns(tmp_path):
    source=tmp_path/'source';source.mkdir()
    book=Workbook();sheet=book.active
    sheet.append(['number','mixed','identifier'])
    sheet.append([1.23456789e-10,'text','00123'])
    sheet.append([12.7875995635986,13.68283661044016,'00456'])
    book.save(source/'data.xlsx')
    store=ManufacturingStore(tmp_path/'store');assert not ingest_package(source,store)['errors']
    rows=store.visible(limit=None)
    numbers=[r['value_number'] for r in rows if r['name']=='number']
    assert numbers==[1.23456789e-10,12.7875995635986]
    assert next(r['value_number'] for r in rows if r['name']=='mixed' and r['source_row']==3)==13.68283661044016
    assert {r['value_text'] for r in rows if r['name']=='identifier'}=={'00123','00456'}
    checked=ManufacturingTools(store).verify_many(rows)
    assert all(c['source_integrity'] and c['value_matches'] for c in checked.values())
    tiny=next(r['id'] for r in rows if r['name']=='number' and r['source_row']==2)
    with pytest.raises(ValueError,match='does not match'):
        ManufacturingTools(store).calculate_metrics(metrics=[{'operation':'mean','evidence_ids':[tiny],'claimed_value':0}])


def test_clock_resets_and_duplicate_times_are_not_silently_aligned(tmp_path):
    source=tmp_path/'source';source.mkdir()
    (source/'data.csv').write_text('batch_id,Time (h),left,right\nB1,0,1,10\nB1,2,2,20\nB1,1,3,30\n')
    store=ManufacturingStore(tmp_path/'store');ingest_package(source,store);tools=ManufacturingTools(store)
    args={'left':{'name':'left','entity_id':'B1'},'right':{'name':'right','entity_id':'B1'}}
    with pytest.raises(ValueError,match='moves backwards'):tools.align_series(**args)
    reordered=tools.align_series(**args,allow_time_reordering=True)
    assert reordered['source_order_warnings'][0]['backwards_steps']==1
    assert reordered['source_order_warnings'][0]['clock_reset_resolved'] is False
    assert any(any('backwards_source_time' in flag for flag in a['flags']) for a in tools.analysis().prepare()['anomalies'])
    (source/'data.csv').write_text('batch_id,Time (h),left,right\nB1,0,1,10\nB1,0,2,20\n')
    ingest_package(source,store)
    with pytest.raises(ValueError,match='duplicate source timestamps'):ManufacturingTools(store).align_series(**args)


def test_future_time_coordinate_and_unknown_time_units_cannot_enable_alignment(tmp_path):
    import csv,json
    source=tmp_path/'source';source.mkdir()
    with (source/'data.csv').open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(['batch_id','available_at','Time (h)','left','right','field_available_at'])
        writer.writerow(['B1','2024-01-01T00:00:00Z',0,1,2,json.dumps({'Time (h)':'2025-01-01T00:00:00Z'})])
    store=ManufacturingStore(tmp_path/'store');ingest_package(source,store)
    args={'left':{'name':'left','entity_id':'B1'},'right':{'name':'right','entity_id':'B1'}}
    with pytest.raises(ValueError,match='source time coordinates'):
        ManufacturingTools(store,as_of='2024-06-01T00:00:00Z').align_series(**args)
    (source/'data.csv').write_text('batch_id,Time,left,right\nB1,2020,1,2\n')
    ingest_package(source,store)
    with pytest.raises(ValueError,match='source time coordinates'):ManufacturingTools(store).align_series(**args)


def test_tiny_values_cannot_pass_as_zero_and_rounding_is_explicit(dataset):
    _,tools,_,ids=dataset
    result=tools.calculate_metrics(metrics=[{'operation':'mean','evidence_ids':ids('VCD Mean (cells/mL)'),
                                            'claimed_value':3.33,'display_decimals':2}])
    assert result['metrics'][0]['value']==pytest.approx(10/3)
    assert result['metrics'][0]['display_value']==3.33
    with pytest.raises(ValueError,match='does not match'):
        tools.calculate_metrics(metrics=[{'operation':'mean','evidence_ids':ids('VCD Mean (cells/mL)'),
                                         'claimed_value':3.33}])


def test_metric_verification_rejects_modified_parquet_population(dataset):
    import polars as pl
    store,tools,_,ids=dataset
    spec={'operation':'mean','selection':{'name':'VCD Mean (cells/mL)','entity_id':'EXP001/R01'}}
    assert tools.calculate_metrics(metrics=[spec])['metrics'][0]['value']==pytest.approx(10/3)
    selected=store.get_many(ids('VCD Mean (cells/mL)'))[0]
    with store.connect() as con:
        part=con.execute('SELECT parquet_path FROM sheets WHERE source_id=? AND sheet=?',
                         (selected['source_id'],selected['sheet'])).fetchone()[0]
    frame=pl.read_parquet(part).with_columns(pl.when(pl.col('__row')==3).then(pl.lit('3.5')).otherwise(pl.col('v2')).alias('v2'))
    frame.write_parquet(part)
    with pytest.raises(ValueError,match='ingestion checksum'):tools.calculate_metrics(metrics=[spec])


def test_mixed_date_number_and_text_column_keeps_date_meaning(tmp_path):
    from datetime import datetime
    source=tmp_path/'source';source.mkdir()
    book=Workbook();sheet=book.active;sheet.append(['mixed'])
    for value in [datetime(2024,1,1),'00123',13.68283661044016]:sheet.append([value])
    book.save(source/'mixed.xlsx')
    store=ManufacturingStore(tmp_path/'store');assert not ingest_package(source,store)['errors']
    rows=store.visible(limit=None)
    assert rows[0]['value_text'].startswith('2024-01-01') and rows[0]['value_number'] is None
    assert rows[1]['value_text']=='00123'
    assert rows[2]['value_number']==13.68283661044016
    assert all(c['value_matches'] for c in ManufacturingTools(store).verify_many(rows).values())
