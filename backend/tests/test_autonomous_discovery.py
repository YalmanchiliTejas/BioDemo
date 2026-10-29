import json
import tempfile
from pathlib import Path

from benchmarks.autonomous.generate import generate, digest
from benchmarks.autonomous.run import score
from src.manufacturing_data import ManufacturingStore, ManufacturingTools, ingest_package


def test_seeded_corpus_and_truth_boundary():
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        a = generate(root / 'a', 100); b = generate(root / 'b', 100)
        left = json.loads(a.read_text()); right = json.loads(b.read_text())
        assert left['corpus_sha256'] == right['corpus_sha256']
        truth = json.loads((root / 'a/evaluator/truth.json').read_text())
        assert len(list((root / 'a/raw').rglob('*.*'))) == 100
        assert len(truth['questions']) == 4
        for q in truth['questions']:
            assert len(q['critical']) == 5
            assert not set(q['affected_batches']) & set(q['negative_controls'])
        assert all('truth' not in p.name for p in (root / 'a/raw').rglob('*'))
        assert truth['duplicates']
        for first, second in truth['duplicates']:
            # Deliberately corrupted/unsupported records may replace a duplicate.
            if (root / 'a/raw' / second).exists():
                assert digest(root / 'a/raw' / first) == digest(root / 'a/raw' / second)


def test_docx_table_and_paragraph_provenance_and_explicit_failures():
    from docx import Document
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary); raw = root / 'raw'; raw.mkdir()
        doc = Document(); doc.add_paragraph('Pressure study contradicts bag attribution.')
        table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = 'B-A12345'; table.cell(0, 1).text = '1.2 bar'
        doc.save(raw / 'record.docx')
        (raw / 'broken.xlsx').write_bytes(b'invalid')
        (raw / 'vendor.bin').write_bytes(b'opaque')
        store = ManufacturingStore(root / 'store')
        report = ingest_package(raw, store, document_backend='pypdf')
        assert report['files_seen'] == 3 and report['files_indexed'] == 1
        assert len(report['errors']) == 1 and report['files_preserved_unsupported'] == 1
        rows = store.visible(limit=None)
        assert any(r['source_ref']['span'] == 'table:1:row:1:cell:2' and r['value_text'] == '1.2 bar' for r in rows)
        assert any(r['source_ref']['span'] == 'paragraph:1' for r in rows)
        current = ManufacturingTools(store)
        sid = next(r['source_id'] for r in rows)
        assert len(current.search_evidence(query='', source_id=sid, evidence_type='structured')['records']) == 3
        tools = ManufacturingTools(store, as_of='2020-01-01T00:00:00Z')
        assert tools.search_evidence(query='Pressure', evidence_type='narrative')['records'] == []


def test_scoring_does_not_equate_file_discovery_with_span_inspection():
    question = {'critical': [{'path': 'a.txt', 'kind': 'counter', 'span': 'contradictory observation'}], 'affected_batches': ['B-1']}
    truth = {'files': [{'path': 'a.txt', 'available_at': '2024-06-01'}, {'path': 'future.txt', 'available_at': '2024-08-01'}]}
    report = {'citations': [{'source': {'path': 'a.txt', 'span': 'line:1'}, 'text': 'campaign identity'},
                            {'source': {'path': 'future.txt'}, 'text': 'future'}], 'analysis': {}, 'ingestion': {}}
    metrics = score(report, question, truth)
    assert metrics['critical_file_recall'] == 1
    assert metrics['critical_span_recall'] == 0
    assert metrics['counterevidence_recall'] == 0
    assert metrics['future_information_citations'] == 1


def test_blank_query_never_infers_an_unmapped_metric():
    from src.manufacturing_data.terminology import query_matches
    assert not query_matches('', 'record_note')
    assert not query_matches('manufacturing failures', 'record_note')
    assert query_matches('dissolved oxygen', 'DO2', 'mg/L')


def test_source_diversity_and_fingerprint_are_stable_across_staging_paths():
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary); raw = root / 'raw'; raw.mkdir()
        (raw / 'many.csv').write_text('batch_id,note\n' + 'B-A12345,pressure excursion\n' * 30)
        (raw / 'other.txt').write_text('pressure excursion\n')
        outputs = []
        for i in range(2):
            package = root / f'package-{i}'; __import__('shutil').copytree(raw, package)
            store = ManufacturingStore(root / f'store-{i}'); ingest_package(package, store, document_backend='pypdf')
            tools = ManufacturingTools(store)
            result = tools.search_evidence(query='pressure', evidence_type='narrative', max_hits_per_source=1, limit=10)
            assert len(result['records']) == 2
            outputs.append(([r['source_ref']['path'] for r in result['records']], store.source_fingerprint()))
        assert outputs[0] == outputs[1]


def test_established_causal_claim_is_explicitly_an_error():
    q = {'critical': [], 'affected_batches': [], 'contributors': [], 'inconclusive': True}
    report = {'analysis': {'claims': [{'classification': 'verified_cause', 'statement': 'Cause established.'}]}}
    metrics = score(report, q, {'files': []})
    assert metrics['unsupported_causal_claims'] == 1
    assert not metrics['correct_abstention']


def test_json_records_keep_numeric_relationships_and_bad_json_is_an_error():
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary); raw = root / 'raw'; raw.mkdir()
        (raw / 'records.json').write_text(json.dumps([
            {'batch_id': 'B-A12345', 'time(h)': .2, 'pressure(bar)': 1.125, 'available_at': '2024-06-01T00:00:00Z'},
            {'batch_id': 'B-A12345', 'time(h)': .4, 'pressure(bar)': 1.9, 'available_at': '2024-08-01T00:00:00Z'}]))
        (raw / 'broken.json').write_text('{invalid')
        path = raw / 'records.json'
        path.write_text(path.read_text().replace('1.125', '1.125000000000000000000123'))
        store = ManufacturingStore(root / 'store'); result = ingest_package(raw, store, document_backend='pypdf')
        assert result['files_indexed'] == 1 and len(result['errors']) == 1
        tools = ManufacturingTools(store, as_of='2024-07-01T00:00:00Z')
        response = tools.query_data(sql="SELECT id,value_number,time_value FROM observations WHERE name='pressure(bar)'")
        assert len(response['rows']) == 1 and response['rows'][0][1:] == [1.125, .2]
        row = store.get_many([response['rows'][0][0]], as_of=tools.as_of)[0]
        assert row['value_text'] == '1.125000000000000000000123'
        assert row['source_ref']['span'] == 'record:1'
        assert tools.verify_evidence(observation_id=row['id'])['value_matches']
